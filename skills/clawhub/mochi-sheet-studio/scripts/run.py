"""Local, traceable business recipes. No spreadsheet formula evaluation."""
import csv,datetime,decimal,html,json,re,zipfile
from collections import defaultdict
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.chart import BarChart,Reference
from openpyxl.styles import Font,PatternFill,Alignment
from io_utils import require,read_file,source,write_json,cli
D=decimal.Decimal
LIMIT=50000

def text(v):
 if isinstance(v,(datetime.date,datetime.datetime)):return v.isoformat()[:10]
 return '' if v is None else str(v).strip()

def money(v):
 s=text(v);require(bool(re.fullmatch(r'[+-]?(?:\d+)(?:\.\d{1,4})?',s)), 'Amount must be an explicit decimal, without currency symbols/group separators')
 x=D(s);require(x.is_finite() and abs(x)<D('1000000000000'),'Amount outside supported range');return x

def read_rows(ws,spec):
 p=read_file(ws,spec.get('file'),25*1024*1024);ext=p.suffix.lower();formulas=[]
 if ext=='.csv':
  enc=spec.get('encoding','utf-8-sig');require(enc in ['utf-8-sig','utf-8','gb18030'],'encoding: utf-8-sig, utf-8, gb18030')
  delim=spec.get('delimiter',',');require(delim in [',',';','\t'],'delimiter: comma, semicolon or tab')
  with p.open(encoding=enc,newline='') as f:
   reader=csv.reader(f,delimiter=delim);headers=next(reader,[]);rows=[]
   for n,row in enumerate(reader,2):
    require(n<=LIMIT+1,'Source exceeds 50000 rows');require(len(row)<=200,'Source exceeds 200 columns')
    if any(text(v) for v in row):require(len(row)==len(headers),'CSV row '+str(n)+' has a different column count');rows.append((n,row))
 elif ext=='.xlsx':
  with zipfile.ZipFile(p) as z:require(sum(i.file_size for i in z.infolist())<=150*1024*1024,'Expanded workbook exceeds 150 MiB')
  book=load_workbook(p,read_only=True,data_only=False,keep_links=False)
  try:
   name=spec.get('sheet',book.sheetnames[0]);require(name in book.sheetnames,'Unknown worksheet');sheet=book[name]
   require(sheet.max_row<=LIMIT+1 and sheet.max_column<=200,'Sheet exceeds 50000 rows or 200 columns')
   rows=[];headers=[]
   for n,row in enumerate(sheet.iter_rows(),1):
    vals=[]
    for cell in row:
     v=cell.value
     if cell.data_type=='f':formulas.append({'cell':cell.coordinate,'formula':v});v={'formula':v,'cell':cell.coordinate}
     elif isinstance(v,(int,float)) and re.fullmatch('0{2,}',cell.number_format or ''):v=str(int(v)).zfill(len(cell.number_format))
     vals.append(v)
    if n==1:headers=vals
    elif any(text(v) for v in vals):rows.append((n,vals))
  finally:book.close()
 else:raise ValueError('Use CSV or XLSX. Convert XLS/XLSM to XLSX first.')
 headers=[text(v) for v in headers];require(1<=len(headers)<=200 and all(headers) and len(set(headers))==len(headers),'Headers must be nonempty and unique')
 return p,headers,rows,formulas

def put(sheet,rows):
 for row in rows:
  sheet.append(row)
  for cell in sheet[sheet.max_row]:
   if isinstance(cell.value,str):cell.data_type='s' # Never interpret source text as a formula.
 sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
 for c in sheet[1]:c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='284C69')
 for col in sheet.columns:
  letter=col[0].column_letter;sheet.column_dimensions[letter].width=min(55,max(14,max(len(str(c.value or '')) for c in list(col)[:250])+2))
 for row in sheet.iter_rows(min_row=2):
  for c in row:
   c.alignment=Alignment(vertical='top',wrap_text=True)

def process(ws,j,out):
 action=j.get('action','analyze');specs=j.get('sources');require(isinstance(specs,list) and 1<=len(specs)<=20,'Provide 1..20 sources')
 require(action in ['inspect','analyze','reconcile'],'action: inspect, analyze or reconcile');require(action!='reconcile' or len(specs)==2,'Reconciliation needs exactly two sources')
 recipe=j.get('recipe','sales');require(recipe in ['sales','expenses'],'recipe: sales or expenses')
 sources=[];loaded=[];inspection=[];total_rows=0
 period=j.get('period')
 if period is not None:require(isinstance(period,str) and re.fullmatch(r'\d{4}-\d{2}',period) is not None,'period must be YYYY-MM');datetime.date.fromisoformat(period+'-01')
 for i,spec in enumerate(specs):
  require(isinstance(spec,dict),'Each source needs file, sheet (XLSX) and mapping');p,headers,rows,formulas=read_rows(ws,spec);sources.append(source(ws,p));loaded.append((spec,p,headers,rows))
  total_rows+=len(rows);require(total_rows<=100000,'Combined sources exceed 100000 rows')
  inspection.append({'file':p.relative_to(ws).as_posix(),'headers':headers,'rows':len(rows),'formulas':formulas[:100],'formulaCount':len(formulas)})
 if action=='inspect':write_json(out/'inspection.json',inspection);return {'action':action,'inspection':inspection,'sources':sources}
 clean=[];issues=[];seen={};totals=defaultdict(lambda:D(0));prior=defaultdict(lambda:D(0));recon=[{},{}];ambiguous=[set(),set()]
 def issue(p,n,reason):issues.append([p.relative_to(ws).as_posix(),n,reason])
 for side,(spec,p,headers,rows) in enumerate(loaded):
  mapping=spec.get('mapping',{});required=['id','currency','amount']+(['date','category'] if action=='analyze' else [])
  require(all(mapping.get(k) in headers for k in required),'Map fields: '+', '.join(required));require(all(v in headers for v in mapping.values()),'Mapped header not found')
  indexes={k:headers.index(v) for k,v in mapping.items()}
  for n,raw in rows:
   values={k:raw[idx] if idx<len(raw) else None for k,idx in indexes.items()}
   try:
    require(not any(isinstance(v,dict) for v in values.values()),'Mapped field contains a formula; export calculated values first')
    rid=text(values['id']);require(bool(rid),'Missing identifier');cur=text(values['currency']).upper();require(re.fullmatch('[A-Z]{3}',cur) is not None,'Currency needs an explicit three-letter code');amount=money(values['amount'])
    if action=='reconcile':
     key=(rid,cur)
     if key in recon[side]:ambiguous[side].add(key);raise ValueError('Duplicate reconciliation identifier/currency')
     recon[side][key]={'amount':amount,'source':p.relative_to(ws).as_posix(),'row':n};continue
    date=datetime.date.fromisoformat(text(values['date']));category=text(values['category']);require(bool(category),'Missing category');kind=text(values.get('kind','sale' if recipe=='sales' else 'expense')).lower()
    require(kind in (['sale','refund'] if recipe=='sales' else ['expense','refund']),'kind must match the recipe');require(amount>=0,'Use positive amounts and kind=refund for reversals');net=-amount if kind=='refund' else amount
    key=(rid,cur,kind);content=(date.isoformat(),category,str(amount.normalize()));old=seen.get(key)
    if old is not None:require(old==content,'Conflicting duplicate identifier/currency/kind');issue(p,n,'Duplicate excluded');continue
    seen[key]=content
    period=j.get('period');require(period is None or re.fullmatch(r'\d{4}-\d{2}',period) is not None,'period must be YYYY-MM')
    month=date.strftime('%Y-%m');previous=None
    if period:
     first=datetime.date.fromisoformat(period+'-01');previous=(first-datetime.timedelta(days=1)).strftime('%Y-%m')
     if month not in [period,previous]:continue
    dest=prior if period and month==previous else totals;dest[(cur,category)]+=net
    clean.append([rid,date.isoformat(),cur,category,kind,str(amount),str(net),p.relative_to(ws).as_posix(),n,'previous' if dest is prior else 'current'])
   except (ValueError,decimal.InvalidOperation,TypeError) as e:issue(p,n,str(e))
 summary=[]
 if action=='reconcile':
  for key in sorted(set(recon[0])|set(recon[1])):
   left,right=recon[0].get(key),recon[1].get(key);delta=left['amount']-right['amount'] if left and right else None
   status='duplicate-review' if key in ambiguous[0]|ambiguous[1] else 'matched' if delta==0 else 'difference' if delta is not None else 'left-only' if left else 'right-only'
   summary.append([*key,str(left['amount']) if left else '',str(right['amount']) if right else '',str(delta) if delta is not None and status!='duplicate-review' else '',status,left['source'] if left else '',left['row'] if left else '',right['source'] if right else '',right['row'] if right else ''])
  heading=['ID','Currency','Left','Right','Difference','Status','Left source','Left row','Right source','Right row']
 else:
  for cur,cat in sorted(set(totals)|set(prior)):summary.append([cur,cat,str(totals[(cur,cat)]),str(prior[(cur,cat)]),str(totals[(cur,cat)]-prior[(cur,cat)])])
  heading=['Currency','Category','Current net','Previous net','Change']
 report={'action':action,'recipe':recipe,'period':j.get('period'),'summaryHeaders':heading,'summary':summary,'acceptedRows':len(clean),'issues':issues,'inspection':inspection,'sources':sources,'currencyConversion':False,'formulasEvaluated':False}
 write_json(out/'report.json',report);write_json(out/'recipe.json',{k:v for k,v in j.items() if k!='requestId'})
 book=Workbook();s=book.active;s.title='Summary';put(s,[heading,*summary])
 # Summary decimals are numeric only here, exact strings remain in report.json.
 numeric=[2,3,4] if action=='analyze' else [2,3,4]
 for row in s.iter_rows(min_row=2):
  for c in numeric:
   if row[c].value!='':row[c].value=float(D(row[c].value));row[c].number_format='#,##0.00;[Red](#,##0.00)'
 if action=='analyze' and summary and len({r[0] for r in summary})==1:
  chart=BarChart();chart.title='Net amounts by currency/category';chart.y_axis.title='Amount (separate currencies)';chart.add_data(Reference(s,min_col=3,max_col=3,min_row=1,max_row=len(summary)+1),titles_from_data=True);chart.set_categories(Reference(s,min_col=2,min_row=2,max_row=len(summary)+1));s.add_chart(chart,'G2')
 if action=='analyze':put(book.create_sheet('Detail'),[['ID','Date','Currency','Category','Kind','Amount','Net','Source','Source row','Period'],*clean])
 put(book.create_sheet('Issues'),[['Source','Row','Reason'],*issues]);put(book.create_sheet('Sources'),[['File','SHA256','Bytes'],*[[a['path'],a['sha256'],a['bytes']] for a in sources]])
 book.save(out/'report.xlsx');book.close()
 table=''.join('<tr>'+''.join('<td>'+html.escape(str(v))+'</td>' for v in row)+'</tr>' for row in [heading,*summary])
 (out/'preview.html').write_text('<!doctype html><meta charset="utf-8"><title>MochiLab Sheet Studio</title><style>body{font:16px system-ui;margin:40px;color:#193346}td{padding:10px;border-bottom:1px solid #ddd}tr:first-child{background:#e7f1f5;font-weight:bold}</style><h1>Business report</h1><p>'+str(len(issues))+' rows require review or were excluded.</p><table>'+table+'</table>',encoding='utf8')
 return {'action':action,'summary':summary,'issues':len(issues),'acceptedRows':len(clean),'sources':sources,'originalsModified':False}

if __name__=='__main__':cli('mochi-sheet-studio',process)
