import csv,hashlib,json,os,shutil,sqlite3,sys,tempfile,unittest,zipfile,uuid
from pathlib import Path
DOMAINS=Path(__file__).resolve().parents[1]/"scripts"
sys.path.insert(0,str(DOMAINS))
import run as sheet
from io_utils import transaction,digest
from openpyxl import Workbook,load_workbook

class Base(unittest.TestCase):
 def setUp(self):
  self.base=Path(os.environ.get('MOCHI_TEST_ROOT',tempfile.gettempdir()))/('case-'+uuid.uuid4().hex);self.base.mkdir(parents=True);self.ws=self.base/'workspace';self.ws.mkdir();self.seq=0
 def tearDown(self):pass # Retain isolated artifacts as functional-test evidence.
 def execute(self,module,job):
  self.seq+=1;return transaction(self.ws,'test-'+module.__name__,{'requestId':'r'+str(self.seq),**job},module.process)
 def output(self,r,name):return next(self.ws/f['path'] for f in r['files'] if f['path'].endswith('/'+name))
 def csv(self,name,rows,headers=None):
  p=self.ws/name
  with p.open('w',encoding='utf-8-sig',newline='') as f:csv.writer(f).writerows([headers or ['id','date','currency','category','amount','kind'],*rows])
  return p
 def spec(self,name):return {'file':name,'mapping':{k:k for k in ['id','date','currency','category','amount','kind']}}
 def row(self,rid='001',amount='9.90',currency='USD',date='2026-10-01',kind='sale',category='Software'):return [rid,date,currency,category,amount,kind]

class SheetTests(Base):
 def test_exact_decimal_and_refund(self):
  self.csv('sales.csv',[self.row(amount='.1'),self.row('002','0.10'),self.row('003','0.20'),self.row('004','0.05',kind='refund')]);r=self.execute(sheet,{'sources':[self.spec('sales.csv')]});self.assertEqual(r['summary'][0][2],'0.25');self.assertEqual(r['issues'],1)
 def test_currency_separation(self):
  self.csv('s.csv',[self.row(),self.row('2','20','CNY')]);r=self.execute(sheet,{'sources':[self.spec('s.csv')]});self.assertEqual(len(r['summary']),2);self.assertEqual([x[0] for x in r['summary']],['CNY','USD'])
 def test_id_zeros_and_trace(self):
  self.csv('s.csv',[self.row()]);r=self.execute(sheet,{'sources':[self.spec('s.csv')]});b=load_workbook(self.output(r,'report.xlsx'));self.assertEqual(b['Detail']['A2'].value,'001');self.assertEqual(b['Detail']['I2'].value,2);b.close()
 def test_repeat_import_no_double_count(self):
  self.csv('s.csv',[self.row()]);r=self.execute(sheet,{'sources':[self.spec('s.csv'),self.spec('s.csv')]});self.assertEqual(r['summary'][0][2],'9.90');self.assertEqual(r['issues'],1)
 def test_conflicting_duplicate_flagged(self):
  self.csv('s.csv',[self.row(),self.row(amount='20')]);r=self.execute(sheet,{'sources':[self.spec('s.csv')]});rep=json.loads(self.output(r,'report.json').read_text());self.assertIn('Conflicting',rep['issues'][0][2])
 def test_month_comparison(self):
  self.csv('s.csv',[self.row(),self.row('2','3',date='2026-09-30'),self.row('3','40',date='2026-08-01')]);r=self.execute(sheet,{'period':'2026-10','sources':[self.spec('s.csv')]});self.assertEqual(r['summary'][0][2:],['9.90','3','6.90'])
 def test_invalid_period(self):
  self.csv('s.csv',[self.row()]);self.assertRaises(ValueError,self.execute,sheet,{'period':'2026-99','sources':[self.spec('s.csv')]})
 def test_expense(self):
  self.csv('s.csv',[self.row(kind='expense'),self.row('2','1',kind='refund')]);r=self.execute(sheet,{'recipe':'expenses','sources':[self.spec('s.csv')]});self.assertEqual(r['summary'][0][2],'8.90')
 def test_invalid_rows_not_silent(self):
  self.csv('s.csv',[self.row(amount='NaN'),self.row('2','1,000'),self.row('3','-1'),self.row('4',currency='?'),self.row('5',date='no')]);r=self.execute(sheet,{'sources':[self.spec('s.csv')]});self.assertEqual(r['issues'],5);self.assertEqual(r['acceptedRows'],0)
 def test_formula_mapped_rejected(self):
  b=Workbook();s=b.active;s.append(['id','date','currency','category','amount','kind']);s.append(self.row());s['E2']='=1+2';b.save(self.ws/'s.xlsx');b.close();r=self.execute(sheet,{'sources':[self.spec('s.xlsx')]});self.assertEqual(r['issues'],1);self.assertEqual(r['acceptedRows'],0)
 def test_formula_inventory(self):
  b=Workbook();s=b.active;s.append(['id','amount']);s.append(['001','=1+2']);b.save(self.ws/'s.xlsx');b.close();r=self.execute(sheet,{'action':'inspect','sources':[{'file':'s.xlsx'}]});self.assertEqual(r['inspection'][0]['formulaCount'],1)
 def test_formatted_id(self):
  b=Workbook();s=b.active;s.append(['id','date','currency','category','amount','kind']);s.append([7,*self.row()[1:]]);s['A2'].number_format='0000';b.save(self.ws/'s.xlsx');b.close();r=self.execute(sheet,{'sources':[self.spec('s.xlsx')]});b=load_workbook(self.output(r,'report.xlsx'));self.assertEqual(b['Detail']['A2'].value,'0007');b.close()
 def test_formula_injection_literal(self):
  self.csv('s.csv',[self.row(category='=HYPERLINK("https://bad")')]);r=self.execute(sheet,{'sources':[self.spec('s.csv')]});b=load_workbook(self.output(r,'report.xlsx'));self.assertEqual(b['Summary']['B2'].data_type,'s');b.close()
 def test_missing_mapping(self):
  self.csv('s.csv',[self.row()]);self.assertRaises(ValueError,self.execute,sheet,{'sources':[{'file':'s.csv'}]})
 def test_malformed_csv(self):
  self.csv('s.csv',[['only']]);self.assertRaises(ValueError,self.execute,sheet,{'sources':[self.spec('s.csv')]})
 def test_reconciliation(self):
  self.csv('l.csv',[['001','USD','10'],['002','USD','4'],['003','CNY','20']],['id','currency','amount']);self.csv('r.csv',[['001','USD','10'],['002','USD','3'],['004','USD','5']],['id','currency','amount']);spec=lambda f:{'file':f,'mapping':{k:k for k in ['id','currency','amount']}};r=self.execute(sheet,{'action':'reconcile','sources':[spec('l.csv'),spec('r.csv')]});self.assertEqual([x[5] for x in r['summary']],['matched','difference','left-only','right-only'])
 def test_reconciliation_duplicate_not_matched(self):
  self.csv('l.csv',[['1','USD','10'],['1','USD','20']],['id','currency','amount']);self.csv('r.csv',[['1','USD','10']],['id','currency','amount']);spec=lambda f:{'file':f,'mapping':{k:k for k in ['id','currency','amount']}};r=self.execute(sheet,{'action':'reconcile','sources':[spec('l.csv'),spec('r.csv')]});self.assertEqual(r['summary'][0][5],'duplicate-review')
 def test_idempotency_and_source_change(self):
  self.csv('s.csv',[self.row()]);j={'requestId':'stable','sources':[self.spec('s.csv')]};r=self.execute(sheet,j);again=self.execute(sheet,j);self.assertTrue(again['replayed']);self.assertEqual(r['files'],again['files']);self.csv('s.csv',[self.row('002')]);self.assertRaises(ValueError,self.execute,sheet,j)
 def test_path_escape(self):self.assertRaises(ValueError,self.execute,sheet,{'sources':[{'file':'../secret.csv'}]})
 def test_originals_and_hashes(self):
  p=self.csv('s.csv',[self.row()]);old=p.read_bytes();r=self.execute(sheet,{'sources':[self.spec('s.csv')]});self.assertEqual(old,p.read_bytes())
  for f in r['files']:self.assertEqual(digest(self.ws/f['path']),f['sha256'])

if __name__=="__main__":unittest.main(verbosity=2)
