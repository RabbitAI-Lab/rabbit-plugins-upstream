# -*- coding: utf-8 -*-
"""
「千题千境 PPT skill」横版 6 页 —— v2（依据 sample3 用户提供的精致样图重建）。

v1 -> v2 修复清单：
1. slide2 卡片上方的白色横条 —— 换成 assets_v2 里「卡内裁剪」的头像素材；
2. slide3 卡片缺少顶部色条 —— 补 ROUND_2_SAME_RECTANGLE 单侧圆角色条 + 同色描边；
3. slide4「糖果撞色」标签压在色块上 —— 标签移到色块下方；风格药丸补小图标；
4. slide5/6 图标在彩色圆底上糊成白块 —— 图标增加 bgc 参数，内细节用底色反白；
5. slide1 气球串补上汇聚到标题药丸的细线。

原则：版式 / 卡片 / 药丸 / 流程圆 / CTA / 文字 / 图标 全部为 PowerPoint 原生可编辑对象，
仅 4 个插画（招手男孩、气球串、两张头像）为抠像后的独立 PNG 素材。
"""
import os

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn, nsdecls
from pptx.oxml import parse_xml

# Output/asset directories: override with env vars QIANTI_OUT / QIANTI_ASSETS.
# Defaults are relative to this script. Never overwrite existing artifacts.
_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.environ.get("QIANTI_ASSETS", os.path.join(_SCRIPT_DIR, "assets_v2"))
if not os.path.isdir(ASSETS):
    print("WARNING: asset dir not found: %s (set QIANTI_ASSETS if this deck needs image assets)" % ASSETS)
_OUTPUT_DIR = os.environ.get("QIANTI_OUT", os.path.join(_SCRIPT_DIR, "output"))
os.makedirs(_OUTPUT_DIR, exist_ok=True)


def safe_save_path(path):
    """Return a non-existing path: bump _v2/_v3... suffix instead of overwriting."""
    base, ext = os.path.splitext(path)
    cand, n = path, 1
    while os.path.exists(cand):
        n += 1
        cand = "%s_v%d%s" % (base, n, ext)
    return cand


OUTPUT = os.path.join(_OUTPUT_DIR, "千题千境_PPTskill_6页_v2.pptx")

# ----------------------------------------------------------------- palette
BG        = "FEFBF5"
ORANGE    = "FF9F45"
ORANGE_D  = "F0862A"
YELLOW    = "FFCD3C"
YELLOW_D  = "F0B410"
GREEN     = "88D8B0"
GREEN_D   = "4FB183"
BLUE      = "6FC2E8"
BLUE_D    = "3E9FD1"
PURPLE    = "B39DDB"
PURPLE_D  = "8B6FC4"
TEXT      = "2C3E50"
TEXT_LT   = "6B7A88"
NAVY      = "2B4E72"
PINK_BG   = "FEDDDE"
SKY_BG    = "CDEDFE"
TINT_B    = "E1F3FE"
TINT_Y    = "FEF5E1"
TINT_G    = "E9F5EB"
WHITE     = "FFFFFF"
LINE_GRAY = "EDE3D6"

# 卡片描边用的浅强调色
T_ORANGE = "FFE0C2"
T_YELLOW = "FFEFC4"
T_GREEN  = "D3EFE1"
T_PURPLE = "E2D9F6"

FONT = "微软雅黑"
SW, SH = 13.333, 7.5


# ----------------------------------------------------------------- helpers
def E(v):
    return Inches(v)


def no_shadow(shape):
    shape.shadow.inherit = False


def no_line(shape):
    shape.line.fill.background()


def solid(shape, hexc, alpha=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(hexc)
    if alpha is not None:
        srgb = shape._element.spPr.find(qn('a:solidFill')).find(qn('a:srgbClr'))
        srgb.append(parse_xml('<a:alpha %s val="%d"/>' % (nsdecls('a'), int(alpha * 1000))))


def set_font(run, name=FONT):
    rPr = run._r.get_or_add_rPr()
    for tag in ('a:latin', 'a:ea', 'a:cs'):
        el = rPr.find(qn(tag))
        if el is not None:
            rPr.remove(el)
    rPr.append(parse_xml('<a:latin %s typeface="%s"/>' % (nsdecls('a'), name)))
    rPr.append(parse_xml('<a:ea %s typeface="%s"/>' % (nsdecls('a'), name)))
    rPr.append(parse_xml('<a:cs %s typeface="%s"/>' % (nsdecls('a'), name)))


def bg(slide, hexc=BG):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(hexc)


def S(slide, kind, x, y, w, h, fill=None, line_c=None, lw=1.0, adj=None, adj2=None,
      rot=0, alpha=None):
    sh = slide.shapes.add_shape(kind, E(x), E(y), E(w), E(h))
    no_shadow(sh)
    try:
        if adj is not None:
            sh.adjustments[0] = adj
        if adj2 is not None and len(sh.adjustments) > 1:
            sh.adjustments[1] = adj2
    except Exception:
        pass
    if rot:
        sh.rotation = rot
    if fill is None:
        sh.fill.background()
    else:
        solid(sh, fill, alpha)
    if line_c is None:
        no_line(sh)
    else:
        sh.line.color.rgb = RGBColor.from_string(line_c)
        sh.line.width = Pt(lw)
    sh.text_frame.word_wrap = True
    return sh


def conn(slide, x1, y1, x2, y2, col, lw):
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
    cn.line.color.rgb = RGBColor.from_string(col)
    cn.line.width = Pt(lw)
    return cn


def box(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf


def para(tf, text, size, color, bold=True, align=PP_ALIGN.LEFT, first=False,
         space_after=0, line_spacing=None, font=FONT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    if space_after:
        p.space_after = Pt(space_after)
    if line_spacing:
        p.line_spacing = line_spacing
    pPr = p._p.get_or_add_pPr()
    pPr.append(parse_xml('<a:buNone %s/>' % nsdecls('a')))
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    set_font(r, font)
    return p


def pic(slide, name, x, y, bw, bh, halign='c', valign='m'):
    """把素材等比塞进 (x,y,bw,bh) 盒子。返回 (shape, px, py, w, h)。"""
    path = os.path.join(ASSETS, name + ".png")
    iw, ih = Image.open(path).size
    ar = iw / float(ih)
    if bw / bh > ar:
        h = bh
        w = h * ar
    else:
        w = bw
        h = w / ar
    px_ = x + (bw - w) / 2 if halign == 'c' else (x if halign == 'l' else x + bw - w)
    py_ = y + (bh - h) / 2 if valign == 'm' else (y if valign == 't' else y + bh - h)
    sh = slide.shapes.add_picture(path, E(px_), E(py_), E(w), E(h))
    return sh, px_, py_, w, h


# ----------------------------------------------------------------- icons
# 每个图标都接受 (col=主色, bgc=图标所在底板的颜色)，内细节用 bgc 反白，
# 这样同一套图标既能在白卡片上（bgc=WHITE）也能在彩色圆底上（bgc=圆色）使用。
def ic_magnifier(s, cx, cy, sz, col, bgc=WHITE):
    r = sz * 0.27
    cy0 = cy - sz * 0.10
    S(s, MSO_SHAPE.OVAL, cx - r, cy0 - r, 2 * r, 2 * r, fill=None, line_c=col, lw=sz * 7.5)
    conn(s, cx + r * 0.68, cy0 + r * 0.68, cx + sz * 0.46, cy + sz * 0.42, col, sz * 7.5)


def ic_doc(s, cx, cy, sz, col, bgc=WHITE):
    w, h = sz * 0.54, sz * 0.68
    S(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx - w / 2, cy - h / 2, w, h, fill=col, adj=0.18)
    for fy in (-0.17, -0.01, 0.15):
        S(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx - w * 0.30, cy + h * fy - sz * 0.023,
          w * 0.60, sz * 0.046, fill=bgc, adj=0.5)


def ic_eye(s, cx, cy, sz, col, bgc=WHITE):
    w, h = sz * 0.78, sz * 0.46
    S(s, MSO_SHAPE.OVAL, cx - w / 2, cy - h / 2, w, h, fill=None, line_c=col, lw=sz * 6.5)
    r = sz * 0.145
    S(s, MSO_SHAPE.OVAL, cx - r, cy - r, 2 * r, 2 * r, fill=col)


def ic_pencil(s, cx, cy, sz, col, bgc=WHITE):
    w = sz * 0.32
    bh = sz * 0.56
    top = cy - sz * 0.36
    S(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx - w / 2, top, w, bh, fill=col, adj=0.28)
    S(s, MSO_SHAPE.RECTANGLE, cx - w / 2, top + sz * 0.11, w, sz * 0.05, fill=bgc)
    S(s, MSO_SHAPE.ISOSCELES_TRIANGLE, cx - w / 2, top + bh, w, sz * 0.16,
      fill=col, rot=180)


def ic_mail(s, cx, cy, sz, col, bgc=WHITE):
    w, h = sz * 0.74, sz * 0.54
    S(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx - w / 2, cy - h / 2, w, h,
      fill=None, line_c=col, lw=sz * 5.0, adj=0.14)
    conn(s, cx - w * 0.44, cy - h * 0.34, cx, cy + h * 0.10, col, sz * 5.0)
    conn(s, cx, cy + h * 0.10, cx + w * 0.44, cy - h * 0.34, col, sz * 5.0)


def ic_gift(s, cx, cy, sz, col, bgc=WHITE):
    w, h = sz * 0.64, sz * 0.54
    S(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx - w / 2, cy - h * 0.16, w, h * 0.74,
      fill=col, adj=0.10)
    S(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx - w * 0.62, cy - h * 0.42, w * 1.24, h * 0.32,
      fill=col, adj=0.26)
    S(s, MSO_SHAPE.RECTANGLE, cx - sz * 0.045, cy - h * 0.42, sz * 0.09, h * 1.0, fill=bgc)


def ic_balloon(s, cx, cy, sz, col, bgc=WHITE):
    w, h = sz * 0.48, sz * 0.58
    top = cy - sz * 0.46
    S(s, MSO_SHAPE.OVAL, cx - w / 2, top, w, h, fill=col)
    S(s, MSO_SHAPE.ISOSCELES_TRIANGLE, cx - sz * 0.05, top + h - sz * 0.01,
      sz * 0.10, sz * 0.08, fill=col)
    conn(s, cx, top + h + sz * 0.07, cx + sz * 0.10, top + h + sz * 0.36, col, sz * 2.6)


def ic_briefcase(s, cx, cy, sz, col, bgc=WHITE):
    w, h = sz * 0.68, sz * 0.48
    S(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx - w / 2, cy - h * 0.26, w, h,
      fill=col, adj=0.16)
    S(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx - sz * 0.17, cy - h * 0.56, sz * 0.34, sz * 0.14,
      fill=None, line_c=col, lw=sz * 4.0, adj=0.4)
    S(s, MSO_SHAPE.RECTANGLE, cx - w * 0.30, cy - h * 0.26 + h * 0.02,
      w * 0.60, sz * 0.035, fill=bgc)


def ic_board(s, cx, cy, sz, col, bgc=WHITE):
    """教学：黑板 + 支架 + 两条粉笔线。"""
    w, h = sz * 0.74, sz * 0.52
    top = cy - h * 0.70
    S(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx - w / 2, top, w, h, fill=col, adj=0.14)
    for sgn in (-1, 1):
        S(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx + sgn * w * 0.26 - sz * 0.028,
          top + h, sz * 0.056, sz * 0.22, fill=col, adj=0.4)
    S(s, MSO_SHAPE.RECTANGLE, cx - w * 0.34, top + h * 0.28, w * 0.68, sz * 0.035, fill=bgc)
    S(s, MSO_SHAPE.RECTANGLE, cx - w * 0.34, top + h * 0.55, w * 0.42, sz * 0.035, fill=bgc)


def ic_leaf(s, cx, cy, sz, col, bgc=WHITE):
    S(s, MSO_SHAPE.TEAR, cx - sz * 0.29, cy - sz * 0.29, sz * 0.58, sz * 0.58,
      fill=col, rot=225)
    conn(s, cx + sz * 0.18, cy + sz * 0.18, cx + sz * 0.30, cy + sz * 0.30, col, sz * 3.0)


def ic_gear(s, cx, cy, sz, col, bgc=WHITE):
    S(s, MSO_SHAPE.GEAR_6, cx - sz * 0.36, cy - sz * 0.36, sz * 0.72, sz * 0.72, fill=col)
    S(s, MSO_SHAPE.OVAL, cx - sz * 0.13, cy - sz * 0.13, sz * 0.26, sz * 0.26, fill=bgc)


def ic_smile(s, cx, cy, sz, col, bgc=WHITE):
    S(s, MSO_SHAPE.SMILEY_FACE, cx - sz * 0.36, cy - sz * 0.36, sz * 0.72, sz * 0.72,
      fill=None, line_c=col, lw=sz * 5.0)


ICONS = {
    "magnifier": ic_magnifier, "doc": ic_doc, "eye": ic_eye, "pencil": ic_pencil,
    "mail": ic_mail, "gift": ic_gift, "balloon": ic_balloon,
    "briefcase": ic_briefcase, "board": ic_board, "leaf": ic_leaf,
    "gear": ic_gear, "smile": ic_smile,
}


def icon(s, name, cx, cy, sz, col, bgc=WHITE):
    ICONS[name](s, cx, cy, sz, col, bgc)


# ----------------------------------------------------------------- decorations
def deco(s, items):
    """items: (kind, cx, cy, size, color, rot)"""
    for kind, cx, cy, size, col, rot in items:
        if kind in ("star5", "star4", "heart", "sun"):
            k = {"star5": MSO_SHAPE.STAR_5_POINT, "star4": MSO_SHAPE.STAR_4_POINT,
                 "heart": MSO_SHAPE.HEART, "sun": MSO_SHAPE.SUN}[kind]
            S(s, k, cx - size / 2, cy - size / 2, size, size, fill=col, rot=rot)
        elif kind == "cloud":
            S(s, MSO_SHAPE.CLOUD, cx - size / 2, cy - size * 0.34,
              size, size * 0.68, fill=col)
        elif kind == "candy":
            w, h = size, size * 0.55
            S(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx - w * 0.34, cy - h / 2, w * 0.68, h,
              fill=col, adj=0.5, rot=rot)
            for sgn in (-1, 1):
                S(s, MSO_SHAPE.ISOSCELES_TRIANGLE, cx + sgn * w * 0.34 - w * 0.16,
                  cy - h / 2, w * 0.32, h, fill=col, rot=90 * sgn + rot)
        elif kind == "dot":
            S(s, MSO_SHAPE.OVAL, cx - size / 2, cy - size / 2, size, size, fill=col)


def corner_deco(s, seed, extra=()):
    """四角通用小装饰，seed 让每页略有差异。"""
    base = [
        ("star5", 0.45, 0.42, 0.34, YELLOW, -12),
        ("heart", 0.95, 0.34, 0.24, "F4A0B8", 8),
        ("cloud", 12.55, 0.62, 0.85, "DCEBFB", 0),
        ("star4", 12.70, 0.32, 0.26, PURPLE, 15),
        ("star5", 0.42, 7.12, 0.30, ORANGE, 18),
        ("heart", 0.86, 6.72, 0.20, "F4A0B8", -10),
        ("star4", 12.60, 7.05, 0.28, YELLOW, -8),
        ("cloud", 12.35, 6.55, 0.70, "DCEBFB", 0),
    ]
    if seed % 2 == 0:
        base[0] = ("star4", 0.45, 0.42, 0.30, ORANGE, -12)
    if seed % 3 == 0:
        base[4] = ("dot", 0.42, 7.12, 0.22, GREEN, 0)
    deco(s, base)
    if extra:
        deco(s, extra)


def balloon_strings(s, px, py, w, h, tie_x, tie_y):
    """从气球素材底部各气球的位置，拉细线汇聚到标题药丸顶端。"""
    import numpy as np
    a = np.asarray(Image.open(os.path.join(ASSETS, "p1_balloons.png")).convert("RGBA"))
    m = a[:, :, 3] > 96
    ih, iw = m.shape
    # 取底部若干行，找各气球的水平中心
    band = m[max(0, ih - 5):ih]
    cols = band.any(axis=0)
    runs, start = [], None
    for i, v in enumerate(cols):
        if v and start is None:
            start = i
        elif not v and start is not None:
            runs.append((start, i - 1))
            start = None
    if start is not None:
        runs.append((start, len(cols) - 1))
    for x0, x1 in runs:
        cx_img = (x0 + x1) / 2.0
        sx = px + cx_img / iw * w
        sy = py + (ih - 1) / float(ih) * h
        conn(s, sx, sy, tie_x, tie_y, "8A9AA8", 1.0)


# ----------------------------------------------------------------- pages
def page1(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    corner_deco(s, 1, extra=[("sun", 12.35, 0.60, 0.85, "FFCE5C", 0),
                             ("candy", 1.75, 6.95, 0.55, "F4A0B8", 20)])
    # 气球串 + 汇聚细线
    _, bx, by, bw, bh = pic(s, "p1_balloons", 5.05, 0.16, 3.25, 1.62, valign='t')
    balloon_strings(s, bx, by, bw, bh, 6.665, 2.02)
    # 药丸 + 标题
    S(s, MSO_SHAPE.ROUNDED_RECTANGLE, 3.92, 2.00, 5.49, 1.52, fill=ORANGE, adj=0.5)
    tf = box(s, 3.92, 2.22, 5.49, 1.08, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "千题千境", 64, WHITE, bold=True, align=PP_ALIGN.CENTER, first=True)
    # 副标题
    tf = box(s, 3.92, 3.66, 5.49, 0.85, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "PPT Skill", 40, BLUE, bold=True, align=PP_ALIGN.CENTER, first=True)
    # 一句话
    tf = box(s, 3.92, 4.60, 5.49, 0.55, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "不同题目，呈现不同的境界", 19, TEXT, bold=False,
         align=PP_ALIGN.CENTER, first=True)
    # 招手男孩
    pic(s, "p1_hero_boy", 0.42, 4.95, 1.95, 2.35, valign='b')
    # 落款
    tf = box(s, 9.9, 6.92, 2.9, 0.42, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "2026.09 · HR Ops", 12, TEXT_LT, bold=False,
         align=PP_ALIGN.RIGHT, first=True)


def page2(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    corner_deco(s, 2)
    tf = box(s, 0.75, 0.42, 6.5, 0.72, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "01 · 它是什么", 27, TEXT, bold=True, first=True)

    S(s, MSO_SHAPE.ROUNDED_RECTANGLE, 1.32, 1.52, 4.90, 4.02, fill=PINK_BG, adj=0.12)
    S(s, MSO_SHAPE.ROUNDED_RECTANGLE, 7.11, 1.52, 4.90, 4.02, fill=SKY_BG, adj=0.12)

    # 头像：素材已在卡片内部裁剪，不会再带出页面底色
    pic(s, "p2_avatar_left", 1.32, 1.72, 4.90, 1.72, halign='c', valign='t')
    pic(s, "p2_avatar_right", 7.11, 1.72, 4.90, 1.72, halign='c', valign='t')

    tf = box(s, 1.32, 3.52, 4.90, 0.68, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "千人千面", 29, TEXT, bold=True, align=PP_ALIGN.CENTER, first=True)
    tf = box(s, 7.11, 3.52, 4.90, 0.68, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "千题千境", 29, TEXT, bold=True, align=PP_ALIGN.CENTER, first=True)

    tf = box(s, 1.32, 4.26, 4.90, 1.06, anchor=MSO_ANCHOR.TOP)
    para(tf, "为每个用户", 15, TEXT, bold=False, align=PP_ALIGN.CENTER,
         first=True, line_spacing=1.42)
    para(tf, "定制专属体验", 15, TEXT, bold=False, align=PP_ALIGN.CENTER, line_spacing=1.42)
    tf = box(s, 7.11, 4.26, 4.90, 1.06, anchor=MSO_ANCHOR.TOP)
    para(tf, "为每种任务", 15, TEXT, bold=False, align=PP_ALIGN.CENTER,
         first=True, line_spacing=1.42)
    para(tf, "定制专属形态", 15, TEXT, bold=False, align=PP_ALIGN.CENTER, line_spacing=1.42)

    tf = box(s, 1.5, 5.78, 10.33, 0.52, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "同样的输入 → 形态各异的 PPT", 18, TEXT, bold=False,
         align=PP_ALIGN.CENTER, first=True)
    tf = box(s, 1.5, 6.30, 10.33, 0.52, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "不靠模板堆砌，靠方法生成", 18, TEXT, bold=False,
         align=PP_ALIGN.CENTER, first=True)


def page3(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    corner_deco(s, 3)
    tf = box(s, 0.75, 0.42, 6.5, 0.72, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "02 · 它能做什么", 27, TEXT, bold=True, first=True)

    cards = [
        (0.72, ORANGE, T_ORANGE, "magnifier", "01", "场景识别", "理解你的需求\n识别使用场景"),
        (3.82, YELLOW, T_YELLOW, "doc", "02", "结构建议", "推荐合适结构\n与内容框架"),
        (6.92, GREEN, T_GREEN, "eye", "03", "终稿预览", "生成完整效果\n提前预览"),
        (10.02, PURPLE, T_PURPLE, "pencil", "04", "可编辑成稿", "支持在线编辑\n轻松修改"),
    ]
    band_h = 0.36
    for x, col, tint, ic, num, title, desc in cards:
        # 卡片：圆角半径 0.07*2.60 ≈ 0.182，与色条 0.5*0.36 = 0.18 对齐
        S(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.42, 2.60, 4.68, fill=WHITE,
          line_c=tint, lw=1.4, adj=0.07)
        # 顶部色条：只圆上面两角
        S(s, MSO_SHAPE.ROUND_2_SAME_RECTANGLE, x, 1.42, 2.60, band_h,
          fill=col, adj=0.5, adj2=0)
        icon(s, ic, x + 1.30, 2.34, 0.98, col, WHITE)
        tf = box(s, x, 3.00, 2.60, 0.72, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, num, 30, col, bold=True, align=PP_ALIGN.CENTER, first=True)
        tf = box(s, x, 3.74, 2.60, 0.58, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, title, 19, TEXT, bold=True, align=PP_ALIGN.CENTER, first=True)
        tf = box(s, x + 0.22, 4.40, 2.16, 1.30, anchor=MSO_ANCHOR.TOP)
        for i, ln in enumerate(desc.split("\n")):
            para(tf, ln, 13, TEXT_LT, bold=False, align=PP_ALIGN.CENTER,
                 first=(i == 0), line_spacing=1.42)

    tf = box(s, 0.72, 6.55, 6.2, 0.50, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "✦  输入：需求文字 + 基本素材", 14, TEXT, bold=False, first=True)
    tf = box(s, 6.42, 6.55, 6.2, 0.50, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "▶  输出：5 步走完，拿走可改的 PPTX", 14, TEXT, bold=False,
         align=PP_ALIGN.RIGHT, first=True)


def page4(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    corner_deco(s, 4)
    tf = box(s, 0.75, 0.38, 6.5, 0.68, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "03 · 你说了算", 27, TEXT, bold=True, first=True)

    # ---- 区 1 风格选择（药丸 + 小图标）
    S(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.78, 1.18, 11.78, 1.66, fill=TINT_B, adj=0.18)
    tf = box(s, 1.10, 1.26, 6.0, 0.42, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "1. 风格选择", 16, TEXT, bold=True, first=True)
    tags = [("smile", "活泼", YELLOW_D, False), ("briefcase", "商务", NAVY, False),
            ("board", "教学", GREEN_D, False), ("gear", "科技", BLUE_D, False),
            ("leaf", "极简", GREEN_D, False), ("balloon", "童趣", WHITE, True)]
    tw, tg, tx, ty, th = 1.66, 0.22, 1.14, 1.82, 0.80
    for i, (ick, label, icol, active) in enumerate(tags):
        x = tx + i * (tw + tg)
        S(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, ty, tw, th,
          fill=(ORANGE if active else WHITE), adj=0.5)
        icon(s, ick, x + 0.32, ty + th / 2, 0.34, icol, ORANGE if active else WHITE)
        tf = box(s, x + 0.54, ty, tw - 0.60, th, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, label, 14, (WHITE if active else TEXT), bold=True, first=True)

    # ---- 区 2 元素选择
    S(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.78, 2.99, 11.78, 1.66, fill=TINT_Y, adj=0.18)
    tf = box(s, 1.10, 3.07, 6.0, 0.42, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "2. 元素选择", 16, TEXT, bold=True, first=True)
    elems = [("图标", True), ("人物贴图", True), ("数据图表", False), ("纯文字", False)]
    ex, ew, eg = 1.10, 2.30, 0.55
    for i, (label, checked) in enumerate(elems):
        x = ex + i * (ew + eg)
        if checked:
            S(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 3.72, 0.40, 0.40, fill=ORANGE, adj=0.28)
            tf = box(s, x, 3.72, 0.40, 0.40, anchor=MSO_ANCHOR.MIDDLE)
            para(tf, "✓", 17, WHITE, bold=True, align=PP_ALIGN.CENTER, first=True)
        else:
            S(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 3.72, 0.40, 0.40, fill=WHITE,
              line_c="D8CFC2", lw=1.2, adj=0.28)
        tf = box(s, x + 0.52, 3.72, ew - 0.52, 0.40, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, label, 15, TEXT, bold=False, first=True)

    # ---- 区 3 配色选择（色块在上，标签在下，不重叠）
    S(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.78, 4.80, 11.78, 1.66, fill=TINT_G, adj=0.18)
    tf = box(s, 1.10, 4.88, 6.0, 0.42, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "3. 配色选择", 16, TEXT, bold=True, first=True)
    sw = [ORANGE, YELLOW, GREEN, BLUE, PURPLE]
    d, gap = 0.62, 0.30
    total = len(sw) * d + (len(sw) - 1) * gap
    cx0 = (SW - total) / 2
    for i, c in enumerate(sw):
        S(s, MSO_SHAPE.OVAL, cx0 + i * (d + gap), 5.30, d, d, fill=c)
    tf = box(s, 1.10, 5.98, 11.18, 0.38, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "糖果撞色", 14, TEXT_LT, bold=False, align=PP_ALIGN.CENTER, first=True)


def page5(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    corner_deco(s, 5)
    tf = box(s, 0.75, 0.42, 6.5, 0.72, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "04 · 它怎么用", 27, TEXT, bold=True, first=True)

    steps = [(ORANGE, "mail", "1", "需求+素材"), (YELLOW, "doc", "2", "结构确认"),
             (GREEN, "pencil", "3", "ASCII 草图"), (BLUE, "eye", "4", "终稿预览")]
    d = 1.72
    cx = 2.28
    gap = 2.86
    cy = 2.70
    for i, (col, ic, num, label) in enumerate(steps):
        x = cx + i * gap
        S(s, MSO_SHAPE.OVAL, x - d / 2, cy - d / 2, d, d, fill=col)
        icon(s, ic, x, cy - 0.30, 0.72, WHITE, col)
        tf = box(s, x - d / 2, cy + 0.02, d, 0.62, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, num, 30, WHITE, bold=True, align=PP_ALIGN.CENTER, first=True)
        tf = box(s, x - 1.20, cy + d / 2 + 0.12, 2.40, 0.52, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, label, 17, TEXT, bold=False, align=PP_ALIGN.CENTER, first=True)
        if i < len(steps) - 1:
            ax = x + d / 2 + 0.06
            S(s, MSO_SHAPE.RIGHT_ARROW, ax, cy - 0.135, gap - d - 0.12, 0.27, fill=NAVY)

    S(s, MSO_SHAPE.ROUNDED_RECTANGLE, 4.62, 4.92, 4.10, 0.70, fill=PURPLE, adj=0.5)
    tf = box(s, 4.62, 4.92, 4.10, 0.70, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "✦  5 · 最终可编辑 PPT", 18, WHITE, bold=True,
         align=PP_ALIGN.CENTER, first=True)

    tf = box(s, 1.5, 6.10, 10.33, 0.52, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "关键：每一步都有用户确认动作", 17, TEXT, bold=False,
         align=PP_ALIGN.CENTER, first=True)


def page6(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    corner_deco(s, 6)
    tf = box(s, 0.75, 0.42, 6.5, 0.72, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "05 · 开始使用", 27, TEXT, bold=True, first=True)

    lines = [(ORANGE, "balloon", "给我一个标题 + 你的素材"),
             (GREEN, "doc", "我按 5 步推进，每步等你确认"),
             (PURPLE, "gift", "5 步后，你拿到的不是 PNG，是能改的 PPTX")]
    ph, pg = 0.86, 0.24
    py = 1.56
    px0, pw = 1.90, 10.05
    for i, (col, ic, text) in enumerate(lines):
        y = py + i * (ph + pg)
        S(s, MSO_SHAPE.ROUNDED_RECTANGLE, px0, y, pw, ph, fill=NAVY, adj=0.5)
        ccx = px0 + ph / 2
        S(s, MSO_SHAPE.OVAL, ccx - ph / 2, y, ph, ph, fill=col)
        icon(s, ic, ccx, y + ph / 2, 0.50, WHITE, col)
        tf = box(s, px0 + ph + 0.28, y, pw - ph - 0.5, ph, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, text, 17, WHITE, bold=False, first=True)

    S(s, MSO_SHAPE.ROUNDED_RECTANGLE, 3.95, 5.32, 5.45, 0.98, fill=ORANGE, adj=0.5)
    tf = box(s, 3.95, 5.32, 5.45, 0.98, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "▶  开始对话 →", 26, WHITE, bold=True, align=PP_ALIGN.CENTER, first=True)

    tf = box(s, 1.5, 6.52, 10.33, 0.50, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "适用：内训分享 / 制度说明 / 项目汇报 / 培训材料", 14, TEXT_LT,
         bold=False, align=PP_ALIGN.CENTER, first=True)


def main():
    prs = Presentation()
    prs.slide_width = E(SW)
    prs.slide_height = E(SH)
    for fn in (page1, page2, page3, page4, page5, page6):
        fn(prs)
    out = safe_save_path(OUTPUT)
    prs.save(out)
    print("saved", out)
    alt = os.path.join(_OUTPUT_DIR, "build", "deck_qianti_v2_render.pptx")
    os.makedirs(os.path.dirname(alt), exist_ok=True)
    alt = safe_save_path(alt)
    prs.save(alt)
    print("saved", alt)


if __name__ == "__main__":
    main()
