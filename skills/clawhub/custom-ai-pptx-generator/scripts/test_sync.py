# -*- coding: utf-8 -*-
"""在真实尺寸下并列比较 sync 图标候选方案（圆底 0.98"，图标 sz=0.588"）。"""
import math
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn, nsdecls
from pptx.oxml import parse_xml

_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.environ.get("QIANTI_OUT", os.path.join(_SCRIPT_DIR, "output"))
os.makedirs(OUT_DIR, exist_ok=True)


def safe_save_path(path):
    """Return a non-existing path: bump _v2/_v3... suffix instead of overwriting."""
    base, ext = os.path.splitext(path)
    cand, n = path, 1
    while os.path.exists(cand):
        n += 1
        cand = "%s_v%d%s" % (base, n, ext)
    return cand
NAVY = "102850"
WHITE = "FFFFFF"
RED = "E3000B"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid()
s.background.fill.fore_color.rgb = RGBColor.from_string("FBF9F4")


def label(x, y, t):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(2.0), Inches(0.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = t
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    rPr = r._r.get_or_add_rPr()
    for tag in ('a:latin', 'a:ea', 'a:cs'):
        rPr.append(parse_xml('<%s %s typeface="微软雅黑"/>' % (tag, nsdecls('a'))))


def shape(kind, x, y, w, h, fill, adj=None):
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    sh.line.fill.background()
    sh.shadow.inherit = False
    if adj:
        for i, v in enumerate(adj):
            try:
                sh.adjustments[i] = v
            except Exception:
                pass
    return sh


def two_arc(cx, cy, sz, Rf, th, hf, a0a, a1a, a0b, a1b):
    R = sz * Rf
    rm = R * (1 + th) / 2
    for a0, a1 in ((a0a, a1a), (a0b, a1b)):
        shape(MSO_SHAPE.BLOCK_ARC, cx - R, cy - R, 2 * R, 2 * R, WHITE,
              adj=[a0 * 0.6, a1 * 0.6, th])
        st = math.radians(a1)
        px, py = cx + rm * math.cos(st), cy + rm * math.sin(st)
        rot = math.degrees(math.atan2(math.cos(st), -math.sin(st))) + 90
        h = sz * hf
        shape(MSO_SHAPE.ISOSCELES_TRIANGLE, px - h / 2, py - h / 2, h, h, WHITE)
        s.shapes[-1].rotation = rot


CD = 0.98
SZ = CD * 0.60
cases = []


def head(cx, cy, rm, a, sz, hf):
    st = math.radians(a)
    px, py = cx + rm * math.cos(st), cy + rm * math.sin(st)
    rot = math.degrees(math.atan2(math.cos(st), -math.sin(st))) + 90
    h = sz * hf
    sh = shape(MSO_SHAPE.ISOSCELES_TRIANGLE, px - h / 2, py - h / 2, h, h, WHITE)
    sh.rotation = rot


def one_arc(cx, cy, sz, Rf, th, hf, a0, a1):
    R = sz * Rf
    rm = R * (1 + th) / 2
    shape(MSO_SHAPE.BLOCK_ARC, cx - R, cy - R, 2 * R, 2 * R, WHITE,
          adj=[a0 * 0.6, a1 * 0.6, th])
    head(cx, cy, rm, a1, sz, hf)


cases.append(("G 双弧 R.34 th.50 箭头.22",
              lambda cx, cy: two_arc(cx, cy, SZ, 0.34, 0.50, 0.22, 192, 348, 12, 168)))
cases.append(("H 双弧 R.32 th.55 箭头.24",
              lambda cx, cy: two_arc(cx, cy, SZ, 0.32, 0.55, 0.24, 196, 344, 16, 164)))
cases.append(("I 双弧 R.34 th.62 箭头.20",
              lambda cx, cy: two_arc(cx, cy, SZ, 0.34, 0.62, 0.20, 190, 350, 10, 170)))
cases.append(("J 单弧250 R.36 th.45 箭头.40",
              lambda cx, cy: one_arc(cx, cy, SZ, 0.36, 0.45, 0.40, 200, 90)))
cases.append(("K 单弧240 R.34 th.50 箭头.34",
              lambda cx, cy: one_arc(cx, cy, SZ, 0.34, 0.50, 0.34, 190, 70)))


def variant_d(cx, cy):
    r = SZ * 0.34
    shape(MSO_SHAPE.DONUT, cx - r, cy - r, 2 * r, 2 * r, WHITE, adj=[0.55])
    for a in (305, 125):
        head(cx, cy, r * 0.78, a, SZ, 0.22)


cases.append(("L 环形 + 两切向箭头", variant_d))

for i, (name, fn) in enumerate(cases):
    cx = 1.5 + (i % 6) * 1.95
    cy = 2.6
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - CD / 2), Inches(cy - CD / 2),
                           Inches(CD), Inches(CD))
    c.fill.solid()
    c.fill.fore_color.rgb = RGBColor.from_string(RED)
    c.line.fill.background()
    c.shadow.inherit = False
    fn(cx, cy)
    label(cx - 1.0, cy + 0.62, name)

out = safe_save_path(os.path.join(OUT_DIR, "test_sync.pptx"))
prs.save(out)
print("saved", out)
