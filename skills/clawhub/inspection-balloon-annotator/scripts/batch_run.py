# -*- coding: utf-8 -*-
"""批量编排：对每个图号串起 渲染→OCR→标注图→表A→表B→自检，并产出总汇总表。
用法: python batch_run.py --manifest manifest.json --out-base <输出根目录> [--dpi 400]
manifest.json: [ {"pdf":"a.pdf","dims":"a_dims.py","meta":"a_meta.json(可选)","name":"图号A"}, ... ]
"""
import argparse, json, os, sys, subprocess

PY = sys.executable
HERE = os.path.dirname(os.path.abspath(__file__))

def run(cmd):
    print(">>", " ".join(cmd))
    r = subprocess.run(cmd, capture_output=True, text=True)
    if r.returncode != 0:
        print("!! 失败:", r.stderr[-800:])
    return r.returncode == 0

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--manifest", required=True)
    ap.add_argument("--out-base", required=True)
    ap.add_argument("--dpi", type=int, default=400)
    a = ap.parse_args()

    jobs = json.load(open(a.manifest, encoding="utf-8"))
    os.makedirs(a.out_base, exist_ok=True)
    rows = []
    for job in jobs:
        name = job["name"]; pdf = job["pdf"]; dims = job["dims"]
        meta = job.get("meta")
        stem = os.path.splitext(os.path.basename(pdf))[0]
        wd = os.path.join(a.out_base, name, "work")
        od = os.path.join(a.out_base, name, "outputs")
        os.makedirs(wd, exist_ok=True); os.makedirs(od, exist_ok=True)
        render = os.path.join(wd, "render.png")
        labels = os.path.join(wd, "labels.json")
        balloons = os.path.join(wd, "balloons.json")
        cov = os.path.join(wd, "coverage_report.json")
        png = os.path.join(od, f"{stem}_气泡标注图_p1.png")
        xlsxA = os.path.join(od, f"{name}_尺寸明细检验表.xlsx")
        xlsxB = os.path.join(od, f"{name}_检验记录表.xlsx")
        ok = True
        ok &= run([PY, os.path.join(HERE, "render_pdf.py"), pdf, render, "--dpi", str(a.dpi)])
        ok &= run([PY, os.path.join(HERE, "ocr_labels.py"), render, labels])
        ok &= run([PY, os.path.join(HERE, "build_balloons.py"), "--dims", dims,
                   "--pdf", pdf, "--png", render, "--ocr", labels, "--out", od, "--work", wd])
        ok &= run([PY, os.path.join(HERE, "gen_excel.py"), "--balloons", balloons, "--out", xlsxA])
        qc = [PY, os.path.join(HERE, "gen_qc_items.py"), "--balloons", balloons,
              "--png", png, "--out", xlsxB]
        if meta and os.path.exists(meta): qc += ["--meta", meta]
        ok &= run(qc)
        ok &= run([PY, os.path.join(HERE, "validate_layout.py"), "--balloons", balloons,
                   "--ocr", labels, "--pdf", pdf, "--png", render, "--out", cov])
        ok &= run([PY, os.path.join(HERE, "finalize_report.py"), "--report", cov])

        stat = {"name": name, "pdf": pdf, "ok": ok}
        if os.path.exists(cov):
            rep = json.load(open(cov, encoding="utf-8"))
            lq = rep.get("layout_quality", {})
            covr = rep.get("coverage", {})
            stat.update({
                "balloons": lq.get("balloons", 0),
                "overlaps": lq.get("bubble_overlaps", 0),
                "leader_text": lq.get("leader_text_violations", 0),
                "circle_text": lq.get("circle_text_violations", 0),
                "to_confirm": len(rep.get("to_be_confirmed", [])),
                "coverage": covr.get("coverage_pct", ""),
                "xlsxA": xlsxA, "xlsxB": xlsxB, "png": png,
            })
        rows.append(stat)
        print(f"[{'OK' if ok else 'FAIL'}] {name}")

    # 总汇总表
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb = Workbook(); ws = wb.active; ws.title = "总汇总表"
    hdr = ["图号/名称", "源PDF", "状态", "气泡数", "重叠", "引线穿字", "圆盘压字",
           "待确认", "覆盖性", "标注图", "表A", "表B"]
    HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
    HEAD_FONT = Font(name="Microsoft YaHei", bold=True, color="FFFFFF", size=10)
    thin = Side(style="thin", color="BFBFBF"); BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
    for i, h in enumerate(hdr, 1):
        c = ws.cell(1, i, h); c.fill = HEAD_FILL; c.font = HEAD_FONT; c.border = BORD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r, s in enumerate(rows, 2):
        vals = [s.get("name", ""), s.get("pdf", ""), "成功" if s.get("ok") else "失败",
                s.get("balloons", ""), s.get("overlaps", ""), s.get("leader_text", ""),
                s.get("circle_text", ""), s.get("to_confirm", ""), s.get("coverage", ""),
                s.get("png", ""), s.get("xlsxA", ""), s.get("xlsxB", "")]
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v); c.border = BORD
            c.alignment = Alignment(horizontal="left" if i in (1,2,10,11,12) else "center", vertical="center", wrap_text=True)
    for i, w in enumerate([16, 30, 8, 8, 7, 9, 9, 8, 9, 40, 36, 36], 1):
        ws.column_dimensions[__import__("openpyxl.utils", fromlist=["get_column_letter"]).get_column_letter(i)].width = w
    out = os.path.join(a.out_base, "总汇总表.xlsx")
    wb.save(out)
    print("saved 总汇总表:", out, "| 图号数:", len(rows))

if __name__ == "__main__":
    main()
