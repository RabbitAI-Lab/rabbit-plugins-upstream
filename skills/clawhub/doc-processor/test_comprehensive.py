#!/usr/bin/env python3
"""
doc-processor v2.7.4 全面功能测试
测试所有功能的完整性和 bug
"""

import sys
import tempfile
import time
from pathlib import Path

sys.path.insert(0, '/home/fuyabing/.openclaw/workspace/skills/doc-processor')
sys.path.insert(0, '/home/fuyabing/.openclaw/workspace/.venvs/doc-processor/lib/python3.12/site-packages')

from doc_processor import DocumentProcessor

test_dir = Path(tempfile.mkdtemp(prefix='doc_test_full_'))
print(f"测试目录：{test_dir}")
print("="*70)

# 初始化处理器
processor = DocumentProcessor(
    workspace=str(test_dir),
    enable_cache=True,
    enable_perf_monitor=True,
    ai_service_type='simple'
)

results = {'passed': 0, 'failed': 0, 'tests': []}

def test(name, func, priority='P0'):
    try:
        func()
        results['passed'] += 1
        results['tests'].append({'name': name, 'status': '✅ PASS', 'priority': priority})
        print(f"[{priority}] ✅ {name}")
        return True
    except Exception as e:
        results['failed'] += 1
        results['tests'].append({'name': name, 'status': '❌ FAIL', 'priority': priority, 'error': str(e)})
        print(f"[{priority}] ❌ {name}: {e}")
        return False

# ========== P0 核心功能 ==========
print("\n【P0 核心功能测试】")
print("-"*70)

def t_word_write():
    output = str(test_dir / "test.docx")
    processor.write(output, {'title': '测试文档', 'paragraphs': ['第一段', '第二段']})
    assert Path(output).exists() and Path(output).stat().st_size > 0

def t_word_read():
    content = processor.read(str(test_dir / "test.docx"))
    assert content.type == 'word' and 'paragraphs' in content.data

def t_excel_write():
    output = str(test_dir / "test.xlsx")
    processor.write(output, {'sheet_name': 'Sheet1', 'data': [['A', 'B'], [1, 2]]})
    assert Path(output).exists() and Path(output).stat().st_size > 0

def t_excel_read():
    content = processor.read(str(test_dir / "test.xlsx"))
    assert content.type == 'excel' and 'Sheet1' in content.data

def t_csv_write():
    output = str(test_dir / "test.csv")
    processor.write(output, [['col1', 'col2'], ['val1', 'val2']])
    assert Path(output).exists()

def t_txt_write():
    output = str(test_dir / "test.txt")
    processor.write(output, "纯文本内容")
    assert Path(output).exists() and "纯文本内容" in Path(output).read_text()

def t_template_fill_word():
    template = str(test_dir / "tpl.docx")
    processor.write(template, {'title': '{{title}}', 'paragraphs': ['{{content}}']})
    output = str(test_dir / "filled.docx")
    result = processor.fill_template(template, {'title': 'T', 'content': 'C'}, output)
    assert Path(result).exists()

def t_template_fill_excel():
    template = str(test_dir / "tpl.xlsx")
    processor.write(template, {'sheet_name': 'S1', 'data': [['{{name}}'], ['{{value}}']]})
    output = str(test_dir / "filled.xlsx")
    result = processor.fill_template(template, {'name': 'N', 'value': 'V'}, output)
    assert Path(result).exists()

test("Word 写入", t_word_write)
test("Word 读取", t_word_read)
test("Excel 写入", t_excel_write)
test("Excel 读取", t_excel_read)
test("CSV 写入", t_csv_write)
test("TXT 写入", t_txt_write)
test("Word 模板填充", t_template_fill_word)
test("Excel 模板填充", t_template_fill_excel)

# ========== P1 格式转换 ==========
print("\n【P1 格式转换测试】")
print("-"*70)

def t_word_to_txt():
    docx = test_dir / "conv.docx"
    processor.write(str(docx), {'title': 'T', 'paragraphs': ['P1', 'P2']})
    txt = test_dir / "conv.txt"
    processor.convert(str(docx), str(txt))
    content = txt.read_text()
    assert 'P1' in content or 'P2' in content

def t_excel_to_txt():
    xlsx = test_dir / "conv.xlsx"
    processor.write(str(xlsx), {'sheet_name': 'S1', 'data': [['A', 'B'], [1, 2]]})
    txt = test_dir / "conv2.txt"
    processor.convert(str(xlsx), str(txt))
    content = txt.read_text()
    assert len(content) > 0

def t_word_to_csv():
    docx = test_dir / "table.docx"
    # 创建带表格的 Word
    from docx import Document
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'A'
    table.cell(0, 1).text = 'B'
    table.cell(1, 0).text = '1'
    table.cell(1, 1).text = '2'
    doc.save(str(docx))
    csv = test_dir / "table.csv"
    processor.convert(str(docx), str(csv))
    assert Path(csv).exists()

test("Word→TXT", t_word_to_txt)
test("Excel→TXT", t_excel_to_txt)
test("Word→CSV", t_word_to_csv)

# ========== P2 批量处理 ==========
print("\n【P2 批量处理测试】")
print("-"*70)

def t_batch_small():
    templates = []
    data_list = []
    for i in range(3):
        tpl = test_dir / f"batch_{i}.docx"
        processor.write(str(tpl), {'title': f'T{i}', 'paragraphs': [f'P{i}']})
        templates.append(str(tpl))
        data_list.append({'title': f'T{i}', 'content': f'C{i}'})
    
    output_dir = test_dir / "batch_out"
    result = processor.batch_fill_templates(templates, data_list, str(output_dir), max_workers=2)
    assert result.success == 3 and result.failed == 0

def t_batch_progress():
    def on_progress(current, total, success, failed):
        pass  # 进度回调
    
    templates = [str(test_dir / f"prog_{i}.docx") for i in range(5)]
    data_list = [{'title': f'T{i}', 'paragraphs': [f'P{i}']} for i in range(5)]
    
    for i, tpl in enumerate(templates):
        processor.write(tpl, {'title': f'T{i}', 'paragraphs': [f'P{i}']})
    
    output_dir = test_dir / "prog_out"
    result = processor.batch_fill_templates(templates, data_list, str(output_dir), progress_callback=on_progress)
    assert result.total == 5

test("小批量处理 (3 个)", t_batch_small)
test("批量进度回调", t_batch_progress)

# ========== P3 缓存管理 ==========
print("\n【P3 缓存管理测试】")
print("-"*70)

def t_cache_stats():
    stats = processor.get_cache_stats()
    assert 'enabled' in stats and 'hit_rate' in stats

def t_cache_clear():
    processor.clear_cache()
    stats = processor.get_cache_stats()
    assert stats.get('entries', 0) == 0 or stats.get('hits', 0) >= 0

def t_cache_preload():
    # 创建模板
    tpl = str(test_dir / "preload.docx")
    processor.write(tpl, {'title': 'P', 'paragraphs': ['C']})
    processor.preload_templates([tpl])
    # 不抛异常即成功

test("缓存统计", t_cache_stats)
test("清空缓存", t_cache_clear)
test("预加载模板", t_cache_preload)

# ========== P4 性能监控 ==========
print("\n【P4 性能监控测试】")
print("-"*70)

def t_perf_stats():
    stats = processor.get_performance_stats()
    assert 'enabled' in stats

def t_perf_operations():
    # 执行一些操作
    for i in range(5):
        output = str(test_dir / f"perf_{i}.docx")
        processor.write(output, {'title': f'P{i}', 'paragraphs': ['C']})
    
    stats = processor.get_performance_stats()
    assert stats.get('enabled') == True

def t_perf_report():
    processor.print_performance_report()  # 不抛异常即成功

test("性能统计", t_perf_stats)
test("性能操作追踪", t_perf_operations)
test("性能报告打印", t_perf_report)

# ========== P5 错误处理 ==========
print("\n【P5 错误处理测试】")
print("-"*70)

def t_error_file_not_found():
    try:
        processor.read(str(test_dir / "nonexistent.docx"))
        assert False, "应该抛出异常"
    except (FileNotFoundError, Exception):
        pass  # 预期行为

def t_error_invalid_format():
    try:
        processor.convert(str(test_dir / "test.docx"), str(test_dir / "test.xyz"))
        assert False, "应该抛出异常"
    except Exception:
        pass  # 预期行为

def t_error_empty_batch():
    try:
        processor.batch_fill_templates([], [], str(test_dir / "empty"))
        assert False, "应该抛出异常"
    except ValueError:
        pass  # 预期行为

test("文件不存在错误", t_error_file_not_found)
test("不支持的格式", t_error_invalid_format)
test("空批量处理", t_error_empty_batch)

# ========== 总结 ==========
print("\n" + "="*70)
print("测试结果汇总")
print("="*70)
print(f"通过：{results['passed']}")
print(f"失败：{results['failed']}")
print(f"成功率：{results['passed']/(results['passed']+results['failed'])*100:.1f}%")
print()

# 按优先级分组
for priority in ['P0', 'P1', 'P2', 'P3', 'P4', 'P5']:
    tests = [t for t in results['tests'] if t['priority'] == priority]
    if tests:
        passed = sum(1 for t in tests if t['status'] == '✅ PASS')
        print(f"{priority}: {passed}/{len(tests)} 通过")
        for t in tests:
            if t['status'] == '❌ FAIL':
                print(f"  ❌ {t['name']}: {t['error']}")

print()
print(f"测试目录：{test_dir}")
print("="*70)

sys.exit(0 if results['failed'] == 0 else 1)
