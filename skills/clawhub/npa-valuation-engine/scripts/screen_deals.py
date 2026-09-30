#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""NPA 标的批量筛选器 — 一堆标的进去，按风险调整后回报排名出来

用法：
    python screen_deals.py --deals deals.json
    python screen_deals.py --deals deals.json --min-conservative-irr 0.15 --json
    python screen_deals.py --deals deals.json --out-xlsx 筛选结果.xlsx

排名口径：默认按【保守情景 IRR】降序 —— 决策看保守情景，乐观只估上行空间。
"""
import argparse
import json
import os
import sys

import numpy_financial as npf

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_model import load_config  # noqa: E402


def build_cf(C, P, hc, f, months, mode, ac):
    n = -(-months // 3)
    cf = [-C]
    for t in range(1, n + 1):
        v = -C * hc
        if mode == 2 and n >= 3:
            w = 0.4 if t == n - 2 else (0.3 if t in (n - 1, n) else 0)
        else:
            w = 1.0 if t == n else 0
        v += P * f * w
        if t == n:
            v -= C * ac
        cf.append(v)
    return cf


def evaluate(cfg):
    d, ho, sc = cfg["deal"], cfg["holding"], cfg["scenarios"]
    acs = cfg["acquisition_costs"]
    fixed = sum(acs.values())
    C = d["expected_deal_price"] + fixed
    P = d["market_price"]
    hc, thr = ho["quarterly_holding_rate"], ho["target_irr"]
    out = {"name": d.get("name", ""), "city": d.get("city", ""), "asset_type": d.get("asset_type", ""),
           "deal_price": d["expected_deal_price"], "C": C, "market_price": P,
           "nominal_discount": d["expected_deal_price"] / d["court_appraisal"],
           "real_discount": d["expected_deal_price"] / P, "scenarios": {}}
    for key, label in (("conservative", "保守"), ("base", "基准"), ("optimistic", "乐观")):
        s = sc[key]
        cf = build_cf(C, P, hc, s["exit_price_ratio"], s["disposal_months"],
                      s["collection_mode"], s["extra_cost_ratio"])
        irr = (1 + npf.irr(cf)) ** 4 - 1
        moic = sum(x for x in cf if x > 0) / abs(sum(x for x in cf if x < 0))
        out["scenarios"][label] = {"irr": irr, "moic": moic,
                                   "pass": irr >= thr, "months": s["disposal_months"]}
    # 打分卡：可修复折价占比
    sc_items = cfg["scorecard"] or []
    total_d = sum(i["ratio"] for i in sc_items)
    repairable = sum(i["ratio"] * i["repairability"] / 5 for i in sc_items)
    out["repairable_ratio"] = (repairable / total_d) if total_d else None
    # 投标上限（保守情景 IRR = 门槛）
    s = sc["conservative"]
    lo, hi = 1000.0, P * 2
    for _ in range(100):
        mid = (lo + hi) / 2
        cf = build_cf(mid + fixed, P, hc, s["exit_price_ratio"], s["disposal_months"],
                      s["collection_mode"], s["extra_cost_ratio"])
        try:
            r = (1 + npf.irr(cf)) ** 4 - 1
        except Exception:
            r = -1
        lo, hi = (mid, hi) if r > thr else (lo, mid)
    out["bid_cap"] = (lo + hi) / 2
    out["headroom"] = out["bid_cap"] - d["expected_deal_price"]
    cons = out["scenarios"]["保守"]
    out["verdict"] = ("可做 ✓" if cons["pass"] else
                      ("临界 △" if cons["irr"] > thr * 0.6 else "放弃 ✗"))
    return out


def main():
    ap = argparse.ArgumentParser(description="NPA 标的批量筛选器")
    ap.add_argument("--deals", required=True, help="标的 JSON（数组或含 deals 键的对象）")
    ap.add_argument("--min-conservative-irr", type=float, default=None, help="保守情景 IRR 下限过滤")
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--out-xlsx", help="导出结果 xlsx")
    a = ap.parse_args()
    raw = json.load(open(a.deals, encoding="utf-8"))
    items = raw.get("deals", raw) if isinstance(raw, dict) else raw
    results = [evaluate(load_config_from_dict(it)) for it in items]
    results.sort(key=lambda x: x["scenarios"]["保守"]["irr"], reverse=True)
    if a.min_conservative_irr is not None:
        results = [r for r in results if r["scenarios"]["保守"]["irr"] >= a.min_conservative_irr]
    if a.json:
        print(json.dumps(results, ensure_ascii=False, indent=2, default=float))
    else:
        print(f"{'标的':<24}{'名义折扣':>9}{'真实折扣':>9}{'保守IRR':>10}{'基准IRR':>10}{'保守MOIC':>10}{'投标上限':>12}{'判定':>8}")
        print("-" * 96)
        for r in results:
            print(f"{r['name'][:22]:<24}{r['nominal_discount']:>8.1%}{r['real_discount']:>9.1%}"
                  f"{r['scenarios']['保守']['irr']:>10.1%}{r['scenarios']['基准']['irr']:>10.1%}"
                  f"{r['scenarios']['保守']['moic']:>10.2f}{r['bid_cap']:>12,.0f}{r['verdict']:>8}")
        print(f"\n共 {len(results)} 个标的；判定口径：保守情景 IRR 是否达标（决策看保守，不看乐观）")
    if a.out_xlsx:
        export_xlsx(results, a.out_xlsx)


def load_config_from_dict(d):
    """把单个标的 dict 合进默认配置"""
    from build_model import DEFAULT_CONFIG, deep_merge
    return load_config_dict(deep_merge(DEFAULT_CONFIG, d))


def load_config_dict(cfg):
    from build_model import ASSET_PROFILES
    if not cfg.get("scorecard"):
        at = str(cfg["deal"].get("asset_type", "住宅")).split("-")[0]
        prof = ASSET_PROFILES.get(at, ASSET_PROFILES["住宅"])
        cfg["scorecard"] = [{"type": t, "ratio": r, "repairability": p, "cost": 30000, "lever": l}
                            for t, r, p, l in prof]
    return cfg


def export_xlsx(results, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "NPA筛选结果"
    heads = ["排名", "标的", "城市", "类型", "成交价", "市场价", "名义折价率", "真实折价率",
             "保守IRR", "基准IRR", "乐观IRR", "保守MOIC", "投标上限", "价差空间", "可修复占比", "判定"]
    for j, h in enumerate(heads, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="8B6914")
        c.alignment = Alignment(horizontal="center")
    for i, r in enumerate(results, 1):
        row = [i, r["name"], r["city"], r["asset_type"], r["deal_price"], r["market_price"],
               r["nominal_discount"], r["real_discount"],
               r["scenarios"]["保守"]["irr"], r["scenarios"]["基准"]["irr"], r["scenarios"]["乐观"]["irr"],
               r["scenarios"]["保守"]["moic"], r["bid_cap"], r["headroom"],
               r["repairable_ratio"], r["verdict"]]
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i + 1, column=j, value=v)
            if j in (7, 8, 9, 10, 11):
                c.number_format = '0.0%'
            elif j in (15,):
                c.number_format = '0.0%'
            elif j in (5, 6, 13, 14):
                c.number_format = '#,##0'
            elif j == 12:
                c.number_format = '0.00"x"'
    for j in range(1, len(heads) + 1):
        ws.column_dimensions[chr(64 + j) if j <= 26 else 'A'].width = 13
    ws.freeze_panes = "A2"
    wb.save(path)
    print(f"已导出: {os.path.abspath(path)}")


if __name__ == "__main__":
    main()
