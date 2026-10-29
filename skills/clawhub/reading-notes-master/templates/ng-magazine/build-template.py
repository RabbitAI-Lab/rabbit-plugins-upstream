# -*- coding: utf-8 -*-
"""NG 杂志竖版笔记 · 合并+校验+截图+PDF+PPTX+打包 流水线模板 (v2.6)

用法：复制本文件到项目 output/ 下改名为 build_mag.py，修改顶部 6 个参数后运行。
依赖：playwright(chromium) + pypdf + python-pptx + pillow
  推荐解释器: ~/.workbuddy/binaries/python/envs/default/bin/python
  生成器脚本用: ~/.workbuddy/binaries/python/versions/3.13.12/bin/python3

★ 踩坑修复已内置：
  1. PDF 用 in 单位（Playwright pdf() 不认 'pt'）
  2. mkdir 判存在再建（WorkBuddy 沙箱 shim 对 exist_ok=True 报 EEXIST）
  3. 裁切校验跳过 SVG 不可渲染元素（用 toUpperCase() 大小写不敏感比较）
  4. PPTX 用 JPEG q87 重建贴图（PNG 44MB → JPEG 11MB）
"""
import re, shutil, glob
from pathlib import Path

# ── 0) 项目参数（只改这里）─────────────────────────────
ROOT = Path(__file__).resolve().parent.parent          # 项目根
OUT = Path(__file__).resolve().parent                  # output 目录
FRAGS = ['_mag_a.html', '_mag_a2.html', '_mag_b1.html',
         '_mag_b2.html', '_mag_c1.html', '_mag_c2.html']   # 生成器产物片段
N_PAGES = 60                                           # 目标页数（必须与页面地图一致）
TITLE = '读书笔记-书名-v2.0'                            # 交付文件名
PDF_W_IN, PDF_H_IN = 6.875, 10.0                       # 495×720pt = 6.875×10in
CSS_NAME = 'mag.css'                                   # 设计系统文件（在 OUT 目录）

# ── 1) 合并片段 ────────────────────────────────────────
secs = []
for f in FRAGS:
    t = (OUT / f).read_text(encoding='utf-8')
    secs += re.findall(r'<section class="page.*?</section>', t, re.S)
print('merged sections:', len(secs))
assert len(secs) == N_PAGES, f'页数必须是 {N_PAGES}！'
html = ('<!DOCTYPE html><html lang="zh"><head><meta charset="utf-8">'
        '<title>' + TITLE + '</title>'
        f'<link rel="stylesheet" href="{CSS_NAME}"></head><body>\n'
        + '\n'.join(secs) + '\n</body></html>')
(OUT / '_mag_merged.html').write_text(html, encoding='utf-8')
print('merged html ok')

# ── 2) Playwright 校验 + 截图 + PDF ────────────────────
from playwright.sync_api import sync_playwright
PNG = OUT / 'v2_png'
if not PNG.exists():                 # 踩坑 2：不用 mkdir(exist_ok=True)
    PNG.mkdir(parents=True)
for old in PNG.glob('*.png'):
    old.unlink()

with sync_playwright() as p:
    b = p.chromium.launch()
    pg = b.new_page(viewport={'width': 660, 'height': 960}, device_scale_factor=2)
    pg.goto((OUT / '_mag_merged.html').as_uri())
    pg.wait_for_timeout(1500)
    n = pg.eval_on_selector_all('.page', 'els => els.length')
    print('pages in dom:', n)

    # 双重校验：破图 + 裁切行走（踩坑 3：SVG 元素 tagName 是小写，统一 toUpperCase）
    issues = pg.evaluate('''() => {
      const out = [];
      const tol = 2.5;
      const SKIP = new Set(['DEFS','STYLE','SCRIPT','CLIPPATH','LINEARGRADIENT',
                            'RADIALGRADIENT','PATTERN','MASK','FILTER','SYMBOL']);
      document.querySelectorAll('.page').forEach((pg, i) => {
        const pr = pg.getBoundingClientRect();
        pg.querySelectorAll('img').forEach(im => {
          if (!im.naturalWidth) out.push('P' + (i + 1) + ' broken-img: ' + im.getAttribute('src'));
        });
        const walk = el => {
          for (const c of el.children) {
            if (SKIP.has(c.tagName.toUpperCase())) continue;
            const r = c.getBoundingClientRect();
            if (r.bottom > pr.bottom + tol || r.right > pr.right + tol || r.top < pr.top - tol)
              out.push('P' + (i + 1) + ' clip <' + c.tagName.toLowerCase() + ' class="' +
                (c.className + '').slice(0, 36) + '"> b+' + (r.bottom - pr.bottom).toFixed(0) +
                ' r+' + (r.right - pr.right).toFixed(0) + ' t' + (r.top - pr.top).toFixed(0));
            walk(c);
          }
        };
        walk(pg);
      });
      return out;
    }''')
    print('issues:', len(issues))
    for x in issues[:50]:
        print('  -', x)

    # 挤压检测：内容容器内容高度超过可见高度 = 文字被裁
    sq = pg.evaluate('''() => {
      const out = [];
      document.querySelectorAll('.page').forEach((pg, i) => {
        pg.querySelectorAll('.bd,.chartbox,.col,.row,.band-body,.pts').forEach(el => {
          if (el.scrollHeight > el.clientHeight + 3)
            out.push('P' + (i + 1) + ' squeeze <' + el.className + '> ' +
                     el.scrollHeight + '>' + el.clientHeight);
        });
      });
      return out;
    }''')
    print('squeeze issues:', len(sq))
    for x in sq[:30]:
        print('  -', x)

    pages = pg.query_selector_all('.page')
    for i, el in enumerate(pages):
        el.screenshot(path=str(PNG / ('p%02d.png' % (i + 1))))
    print('screenshots:', len(pages))

    # 踩坑 1：PDF 用 in 单位，不认 'pt'
    pg.pdf(path=str(OUT / (TITLE + '.pdf')),
           width=f'{PDF_W_IN}in', height=f'{PDF_H_IN}in',
           print_background=True, prefer_css_page_size=False,
           margin={'top': '0', 'bottom': '0', 'left': '0', 'right': '0'})
    print('pdf ok')
    b.close()

# ── 3) PDF 页数校验 ────────────────────────────────────
from pypdf import PdfReader
r = PdfReader(str(OUT / (TITLE + '.pdf')))
print('pdf pages:', len(r.pages))
assert len(r.pages) == N_PAGES, f'PDF 页数不是 {N_PAGES}！'

# ── 4) PPTX（JPEG q87 重建，体积减 75%）────────────────
from pptx import Presentation
from pptx.util import Inches
from PIL import Image
JPG = OUT / 'v2_jpg'
if not JPG.exists():
    JPG.mkdir(parents=True)
for f in sorted(glob.glob(str(PNG / 'p*.png'))):
    im = Image.open(f).convert('RGB')
    im.save(str(JPG / (Path(f).stem + '.jpg')), quality=87, optimize=True)
prs = Presentation()
prs.slide_width = Inches(PDF_W_IN)
prs.slide_height = Inches(PDF_H_IN)
blank = prs.slide_layouts[6]
for f in sorted(glob.glob(str(JPG / 'p*.jpg'))):
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(f, 0, 0, width=prs.slide_width, height=prs.slide_height)
prs.save(str(OUT / (TITLE + '.pptx')))
print('pptx ok, slides:', len(prs.slides._sldIdLst))

# ── 5) 空白页检测（PIL 像素方差）───────────────────────
import numpy as np
for f in sorted(glob.glob(str(PNG / 'p*.png'))):
    g = np.asarray(Image.open(f).convert('L'), dtype=float)
    std = g.std()
    if std < 18:
        print(f'  low-variance page: {Path(f).name} std={std:.1f}（黑底金句页属设计使然可忽略，其余需人工检查）')

# ── 6) 打包交付目录 ────────────────────────────────────
pkg = OUT / TITLE
if not pkg.exists():
    pkg.mkdir()
shutil.copy(OUT / (TITLE + '.pdf'), pkg)
shutil.copy(OUT / (TITLE + '.pptx'), pkg)
pack_html = html.replace('src="../assets/', 'src="assets/')
(pkg / 'notes.html').write_text(pack_html, encoding='utf-8')
shutil.copy(OUT / CSS_NAME, pkg / CSS_NAME)
ast = pkg / 'assets'
if not ast.exists() and (OUT.parent / 'assets').exists():
    shutil.copytree(OUT.parent / 'assets', ast, ignore=shutil.ignore_patterns('.DS_Store'))
print('package ok:', pkg)
print('ALL DONE ✓')
