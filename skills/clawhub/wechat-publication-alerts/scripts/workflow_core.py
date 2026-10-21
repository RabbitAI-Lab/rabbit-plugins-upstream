"""Portable gateway workflows: live prices, durable requests and evidence exports."""
import hashlib
import html
import json
import os
from pathlib import Path
import tempfile
import time
import urllib.error
import urllib.parse
import urllib.request
import uuid
from datetime import datetime, timezone
from decimal import Decimal
from contextlib import contextmanager

from platform_contract import catalog, failure, identity_headers, price_micros, product_path
from data_safety import public_data

ORIGIN = "https://api.we-media.cn"


def now():
    return datetime.now(timezone.utc).isoformat()


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False,
                                     separators=(",", ":"), default=str).encode()).hexdigest()


def micros(value):
    n = Decimal(str(value)) * 1000000
    if not n.is_finite() or n < 0 or n != n.to_integral_value():
        raise ValueError("金额须为非负数，最多六位小数")
    return int(n)


def atomic_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile("w", encoding="utf-8", dir=path.parent, delete=False) as out:
        json.dump(public_data(value), out, ensure_ascii=False, indent=2)
        temp = out.name
        out.flush()
        os.fsync(out.fileno())
    os.replace(temp, path)


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


def key(required=True):
    value = next((os.environ[n].strip() for n in ("WM_API_KEY", "MANGYUN_API_KEY", "MANGE_API_KEY")
                  if os.environ.get(n, "").strip()), "")
    if required and not value:
        channel = identity_headers()["X-Source"]
        slug = identity_headers()["X-WM-Source"]
        raise ValueError("请在本机环境变量 WM_API_KEY 设置曼格云 API Key，不要发到聊天。"
                         "注册和创建 Key：" + ORIGIN + "?" + urllib.parse.urlencode(
                             {"source": channel, "skill": slug}) + "；具体采集/订阅按平台价格计费。")
    return value


class APIError(Exception):
    def __init__(self, code, request_id="", status=0, details=None):
        self.code, self.request_id, self.status = str(code), str(request_id), status
        self.details = details or {}
        super().__init__(self.code + (" (requestId=" + self.request_id + ")" if self.request_id else ""))


class NoRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        raise APIError("API_REDIRECT_REFUSED", status=code)


class API:
    def __init__(self, require_key=True):
        self.key = key(require_key)
        self.identity = digest(self.key) if self.key else "public"
        self.products = None
        self.opener = urllib.request.build_opener(NoRedirect())

    def load_catalog(self):
        if self.products is None:
            self.products = {p["slug"]: p for p in catalog(ORIGIN, self.key)}
        return self.products

    def unit(self, slug, params):
        p = self.load_catalog().get(slug)
        if not p:
            raise ValueError("当前账户目录没有此商品：" + slug)
        return price_micros(p, params)

    def request(self, slug, params, idem):
        product = self.load_catalog().get(slug)
        if not product:
            raise ValueError("当前账户目录没有此商品：" + slug)
        names = {p["name"] for p in product.get("parameters") or []}
        if set(params) - names:
            raise ValueError("接口不支持这些参数：" + ", ".join(sorted(set(params) - names)))
        method = product["publicMethod"].upper()
        url = ORIGIN + product_path(product)
        headers = {**identity_headers(), "X-API-Key": self.key, "Content-Type": "application/json",
                   "Idempotency-Key": idem}
        body = json.dumps(params, ensure_ascii=False).encode() if method != "GET" else None
        if method == "GET":
            url += "?" + urllib.parse.urlencode(params)
        uncertain_retry = False
        attempts = 1 if params.get("action") in ("channel.test", "channel.test-draft", "channel.save", "delivery.retry") else 3
        for attempt in range(attempts):
            try:
                req = urllib.request.Request(url, data=body, headers=headers, method=method)
                with self.opener.open(req, timeout=100) as response:
                    raw = response.read(16 * 1024 * 1024 + 1)
                if len(raw) > 16 * 1024 * 1024:
                    raise APIError("RESPONSE_TOO_LARGE")
                payload = json.loads(raw, parse_float=Decimal)
                if failure(payload) or payload.get("ok") is False:
                    raise APIError(payload.get("code", "UPSTREAM_FAILED"), payload.get("requestId", ""),
                                   details=payload.get("details"))
                data = payload.get("data", payload)
                if isinstance(data, dict) and data.get("ok") is False:
                    raise APIError("UPSTREAM_FAILED", payload.get("requestId", ""))
                return data, {"requestId": payload.get("requestId", ""), "uncertainRetry": uncertain_retry,
                              "consumptionMicros": micros(payload.get("consumption", 0)), "capturedAt": now()}
            except urllib.error.HTTPError as error:
                try:
                    payload = json.loads(error.read(65536))
                except (ValueError, OSError):
                    payload = {}
                if error.code == 429 and attempt + 1 < attempts:
                    try:
                        delay = max(1, int(error.headers.get("Retry-After", "2")))
                    except ValueError:
                        delay = 2
                    if delay <= 15:
                        time.sleep(delay)
                        continue
                raise APIError(payload.get("code", "HTTP_" + str(error.code)),
                               payload.get("requestId", ""), error.code, payload.get("details")) from None
            except (urllib.error.URLError, TimeoutError, ConnectionError, OSError):
                uncertain_retry = True
                if attempt + 1 < attempts:
                    time.sleep(1 + attempt)
                    continue
                raise APIError("NETWORK_RESULT_UNKNOWN") from None
            except (ValueError, UnicodeError):
                raise APIError("INVALID_RESPONSE_RESULT_UNKNOWN") from None


@contextmanager
def workspace_lock(directory):
    """OS lock releases on crash; no stale PID files or concurrent writers."""
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    with (directory / ".workflow.lock").open("a+b") as handle:
        handle.seek(0)
        if os.name == "nt":
            import msvcrt
            if handle.read(1) == b"":
                handle.write(b"0")
                handle.flush()
            handle.seek(0)
            try:
                msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
            except OSError:
                raise ValueError("此工作区已有任务运行，请等待完成") from None
        else:
            import fcntl
            try:
                fcntl.flock(handle, fcntl.LOCK_EX | fcntl.LOCK_NB)
            except OSError:
                raise ValueError("此工作区已有任务运行，请等待完成") from None
        try:
            yield
        finally:
            handle.seek(0)
            if os.name == "nt":
                msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                fcntl.flock(handle, fcntl.LOCK_UN)


class Job:
    def __init__(self, directory, api, scope, budget):
        self.directory, self.api = Path(directory), api
        self.path = self.directory / "job.json"
        scope_hash = digest(scope)
        if self.path.exists():
            self.state = read_json(self.path)
            if self.state["scopeHash"] != scope_hash or self.state["owner"] != api.identity:
                raise ValueError("工作区对应其他任务或 API Key，请指定新的工作区")
        else:
            self.state = {"schema": 1, "scopeHash": scope_hash, "owner": api.identity, "scope": scope,
                          "createdAt": now(), "requests": {}, "errors": []}
        self.budget = int(budget)
        self.save()

    def save(self):
        atomic_json(self.path, self.state)

    def totals(self):
        rows = list(self.state["requests"].values())
        return {"confirmedMicros": sum(r.get("consumptionMicros", 0) for r in rows),
                "unconfirmedReserveMicros": sum(r["reserveMicros"] for r in rows if r.get("chargeUnknown")),
                "requestCount": len(rows), "budgetMicros": self.budget}

    def call(self, slug, params, project=lambda d: d, allowance=None):
        ident = digest([slug, params])
        record = self.state["requests"].get(ident)
        if record and record["state"] == "done":
            return record["data"]
        if record and record["state"] == "failed":
            raise APIError(record["error"], record.get("requestId", ""), details=record.get("details"))
        if record is None:
            amount = self.api.unit(slug, params) if allowance is None else allowance
            totals = self.totals()
            if totals["confirmedMicros"] + totals["unconfirmedReserveMicros"] + amount > self.budget:
                raise APIError("BUDGET_LIMIT")
            record = {"slug": slug, "state": "pending", "idempotencyKey": uuid.uuid4().hex,
                      "reserveMicros": amount, "chargeUnknown": True, "createdAt": now()}
            self.state["requests"][ident] = record
            self.save()
            previously_unknown = False
        else:
            previously_unknown = True
        try:
            data, receipt = self.api.request(slug, params, record["idempotencyKey"])
            projected = project(data)
            # A zero-cost replay after a lost response cannot prove the original charge was zero.
            unknown = (previously_unknown or receipt.get("uncertainRetry")) and receipt["consumptionMicros"] == 0 and record["reserveMicros"] > 0
            # Subscription replays carry the original immutable order charge.
            if slug == "wechat-publication-alerts-api" and isinstance(data, dict) and isinstance(data.get("chargeMicros"), int):
                receipt["consumptionMicros"], unknown = data["chargeMicros"], False
            record.update(receipt, data=projected, state="done", chargeUnknown=unknown)
            self.save()
            return projected
        except APIError as error:
            details = {k: error.details[k] for k in ("collectionId", "nextPage", "nextCursor")
                       if isinstance(error.details, dict) and k in error.details}
            uncertain = error.status == 0 or error.status >= 500 or error.code in ("IDEMPOTENCY_IN_PROGRESS", "REQUEST_IN_PROGRESS")
            record.update(error=error.code, requestId=error.request_id, details=details,
                          state="pending" if uncertain else "failed", chargeUnknown=uncertain)
            self.save()
            raise


def list_items(data, names=("items", "videos", "accounts", "articles", "list")):
    if isinstance(data, list):
        return data
    if isinstance(data, dict):
        for name in names:
            if isinstance(data.get(name), list):
                return data[name]
    raise ValueError("响应缺少预期列表，不能将接口结构变化当成零结果")


def next_params(data, params):
    """Honor page/cursor contract, including empty pages with hasMore=true."""
    if data.get("hasMore") is False:
        return None
    result = dict(params)
    if data.get("collectionId") and data.get("nextPage") is not None:
        for name in ("offset", "limit", "cursor"):
            result.pop(name, None)
        result.update(collectionId=data["collectionId"], page=data["nextPage"])
    elif data.get("nextCursor", data.get("cursor")) not in (None, "") and data.get("hasMore"):
        result["cursor"] = data.get("nextCursor", data.get("cursor"))
    elif data.get("hasMore") and data.get("nextOffset") is not None and "offset" in params:
        result["offset"] = data["nextOffset"]
    elif data.get("hasMore"):
        raise APIError("PAGINATION_POSITION_MISSING")
    else:
        return None
    if result == params:
        raise APIError("PAGINATION_CURSOR_REPEATED")
    return result


def select(data, fields):
    return {k: data.get(k) for k in fields if k in data}


def page_project(data, project):
    return {**select(data, ("hasMore", "nextCursor", "cursor", "collectionId", "nextPage", "nextOffset",
                           "count", "totalCount", "capturedAt", "warnings", "status", "partial")),
            "account": select(data.get("account") or {}, ("accountId", "accountName", "name", "ghid", "biz", "alias")),
            "items": [project(r) for r in list_items(data)]}


def excel_ready():
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        raise ValueError("生成 Excel 需要 openpyxl，请先运行 python -m pip install openpyxl；尚未采集或扣费") from None


def export_bundle(directory, title, sheets, notes, raw, findings=None):
    """Exact IDs, formula-safe cells, all rows and JSON evidence retained."""
    excel_ready()
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    atomic_json(directory / "evidence.json", raw)
    book = Workbook()
    book.remove(book.active)
    text = ["# " + title, "", *notes, ""]
    sheets = {**sheets, "说明": [{"说明": x} for x in notes]}
    if findings:
        sheets["分析与建议"] = [{"分析（含证据编号）": item} for item in findings.split("\n\n")]
    for name, rows in sheets.items():
        rows = rows or [{"说明": "本次没有可用记录"}]
        columns = list(dict.fromkeys(k for r in rows for k in r))
        if len(rows) > 1048575:
            raise ValueError("单工作表超过 Excel 行数上限，请缩小采集范围")
        sheet = book.create_sheet(name[:31])
        sheet.append(columns)
        for row in rows:
            sheet.append([None] * len(columns))
            for index, col in enumerate(columns, 1):
                value = row.get(col)
                if isinstance(value, (dict, list)):
                    value = json.dumps(value, ensure_ascii=False)
                if isinstance(value, Decimal):
                    value = str(value)
                cell = sheet.cell(sheet.max_row, index)
                if isinstance(value, str):
                    # Full text remains in evidence.json; no formula interpretation.
                    value = "".join(c for c in value if ord(c) >= 32 or c in "\t\n\r")
                    cell.value, cell.data_type = value[:32767], "s"
                else:
                    cell.value = value
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="1D4E89")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, col in enumerate(columns, 1):
            sheet.column_dimensions[get_column_letter(index)].width = 44 if any(x in col.lower() for x in ("content", "内容", "标题", "证据", "说明")) else 24
        sheet.row_dimensions[1].height = 26
        text.extend(["## " + name, "", "记录数：" + str(len(rows)), ""])
        if name not in ("说明", "分析与建议"):
            for row in rows[:5]:
                text.append("- " + "; ".join(str(k) + ": " + str(v)[:240] for k, v in list(row.items())[:8]))
            text.append("")
    if findings:
        text.extend(["## 分析与建议", "", findings, ""])
    text.extend(["数据与完整明细见 report.xlsx、evidence.json。", "",
                 "曼格云 API：" + ORIGIN + "；报告修订可复用本次数据，不需要重复采集。"])
    book.save(directory / "report.xlsx")
    book.close()
    (directory / "report.md").write_text("\n".join(text), encoding="utf-8")
    (directory / "report.html").write_text(
        '<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
        '<title>' + html.escape(title) + '</title><style>body{font:16px/1.7 system-ui;max-width:960px;margin:32px auto;padding:0 20px;color:#243247}pre{white-space:pre-wrap;overflow-wrap:anywhere}</style><pre>'
        + html.escape("\n".join(text)) + "</pre></html>", encoding="utf-8")
    return {ext: str((directory / ("report." + ext)).resolve()) for ext in ("xlsx", "md", "html")}


def analysis_text(path, evidence_ids, facts_hash):
    if not path:
        return None
    analysis = read_json(path)
    if analysis.get("factsHash") != facts_hash:
        raise ValueError("分析对应其他数据或旧口径，请按本次 analysis-input.json 重新分析")
    result = []
    for finding in analysis.get("findings", []):
        refs = finding.get("evidenceIds")
        if not refs or any(ref not in evidence_ids for ref in refs):
            raise ValueError("每条分析必须引用本次真实 evidenceIds")
        result.append(str(finding["finding"]) + "\n建议：" + str(finding.get("action", "")) + "\n依据：" + ", ".join(refs))
    if not result:
        raise ValueError("分析为空，不能当作完成的解读报告")
    return "\n\n".join(result)
