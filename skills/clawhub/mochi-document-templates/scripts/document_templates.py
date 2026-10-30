"""Literal, non-executable template replacement across DOCX runs and tables."""
import html,re,zipfile
from core import *
PATTERN=re.compile(r'\{\{\s*([A-Za-z_][A-Za-z0-9_.-]{0,99})\s*\}\}')
def fill(text,data,escape=False):
    missing=sorted(set(PATTERN.findall(text))-set(data))
    if missing: raise ValueError('Missing template fields: '+', '.join(missing))
    def sub(m):
        value=str(data[m.group(1)])
        return html.escape(value) if escape else value
    return PATTERN.sub(sub,text)
def paragraphs(doc):
    for p in doc.paragraphs: yield p
    def tables(tables):
        for t in tables:
            for row in t.rows:
                for cell in row.cells:
                    yield from cell.paragraphs
                    yield from tables_inner(cell.tables)
    def tables_inner(ts): return tables(ts)
    yield from tables(doc.tables)
    for section in doc.sections:
        for part in (section.header,section.footer,section.first_page_header,section.first_page_footer,section.even_page_header,section.even_page_footer):
            yield from part.paragraphs; yield from tables(part.tables)
def replace_paragraph(p,data):
    text=''.join(run.text for run in p.runs)
    # Replace from right to left; retain surrounding run formatting and apply the
    # first token run's formatting to its inserted value.
    for m in reversed(list(PATTERN.finditer(text))):
        if m.group(1) not in data: raise ValueError('Missing template field: '+m.group(1))
        positions=[]; n=0
        for i,run in enumerate(p.runs): positions.append((n,n+len(run.text),i)); n+=len(run.text)
        hits=[(a,b,i) for a,b,i in positions if a<m.end() and b>m.start()]
        if not hits: continue
        first,last=hits[0],hits[-1]; a,b,i=first; x,y,j=last
        prefix=p.runs[i].text[:m.start()-a]; suffix=p.runs[j].text[m.end()-x:]
        p.runs[i].text=prefix+str(data[m.group(1)])+(suffix if i==j else '')
        for k in range(i+1,j): p.runs[k].text=''
        if j!=i: p.runs[j].text=suffix
def handle(st,r):
    p=local(st.root,r['template'],True)
    if p.stat().st_size>20*1024*1024: raise ValueError('Template <=20 MiB')
    docx=p.suffix.lower()=='.docx'
    if docx:
        from docx import Document
        with zipfile.ZipFile(p) as z:
            if sum(x.file_size for x in z.infolist())>100*1024*1024: raise ValueError('Expanded DOCX too large')
            if any('vbaProject' in x.filename for x in z.infolist()): raise ValueError('Macro-enabled template rejected')
        doc=Document(p); text='\n'.join(''.join(x.text for x in q.runs) for q in paragraphs(doc))
    else:
        if p.suffix.lower() not in ('.txt','.md','.html','.csv'): raise ValueError('Template must be DOCX/TXT/MD/HTML/CSV')
        text=p.read_text(encoding='utf-8-sig')
    fields=sorted(set(PATTERN.findall(text)))
    if r['action']=='inspect': return {'fields':fields,'templateSha256':digest(p.read_bytes()),'format':p.suffix}
    if r['action']!='render': raise ValueError('action inspect or render')
    rows=r.get('rows') or [r['data']]
    if not isinstance(rows,list) or not 1<=len(rows)<=200: raise ValueError('Use 1..200 data rows')
    prepared=[]
    for row in rows:
        data=row.get('data',row); fill(text,data)
        name=ident(row.get('outputName',str(len(prepared)+1)))+p.suffix.lower(); target=local(st.root,r.get('outputDirectory','exports/documents')+'/'+name)
        if target.exists() or target in [x[0] for x in prepared]: raise ValueError('Output exists or repeated output name')
        prepared.append((target,data))
    out=[]
    for target,data in prepared:
        if docx:
            import io
            doc=Document(p)
            for q in paragraphs(doc): replace_paragraph(q,data)
            b=io.BytesIO(); doc.save(b); raw=b.getvalue()
        else: raw=fill(text,data,p.suffix.lower()=='.html').encode('utf8')
        atomic(target,raw); out.append({'file':str(target),'sha256':digest(raw),'fields':fields})
    return {'files':out,'templateSha256':digest(p.read_bytes())}
if __name__=='__main__': runmain(handle,'document-templates')
