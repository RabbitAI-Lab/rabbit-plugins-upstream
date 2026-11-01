import datetime as dt,io,json,os,subprocess,tempfile,unittest,zipfile,uuid,sys
from pathlib import Path
from unittest.mock import patch,MagicMock
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/"scripts"))
from core import *
import document_templates as docs

class Base(unittest.TestCase):
    def setUp(self):
        self.root=Path(os.environ.get('MOCHI_TEST_TEMP',tempfile.gettempdir()))/('case-'+uuid.uuid4().hex); self.root.mkdir(); self.st=Store(self.root,'test')
    def tearDown(self): self.st.db.close()
    def file(self,name,data):
        p=self.root/name; p.parent.mkdir(parents=True,exist_ok=True); p.write_bytes(data if isinstance(data,bytes) else data.encode()); return p

class DocumentTests(Base):
    def test_text_batch(self):
        self.file('t.txt','Hello {{name}} / {{amount}}'); result=docs.handle(self.st,{'action':'render','template':'t.txt','rows':[{'outputName':'a','data':{'name':'Alice','amount':'9.90'}},{'outputName':'b','data':{'name':'张三','amount':'20'}}]}); self.assertEqual(len(result['files']),2); self.assertIn('张三',Path(result['files'][1]['file']).read_text())
    def test_html_escaped(self): self.assertEqual(docs.fill('{{name}}',{'name':'<script>'},True),'&lt;script&gt;')
    def test_missing_fields_no_output(self):
        self.file('t.txt','{{name}} {{unknown}}')
        with self.assertRaises(ValueError): docs.handle(self.st,{'action':'render','template':'t.txt','data':{'name':'a'}})
        self.assertFalse((self.root/'exports/documents').exists())
    def test_docx_runs_and_table(self):
        from docx import Document
        d=Document(); p=d.add_paragraph(); p.add_run('Dear {{na').bold=True; p.add_run('me}}, welcome'); d.add_table(rows=1,cols=1).cell(0,0).text='{{amount}}'; d.save(self.root/'t.docx'); result=docs.handle(self.st,{'action':'render','template':'t.docx','data':{'name':'Alice','amount':'9.90'}}); out=Document(result['files'][0]['file']); self.assertEqual(out.paragraphs[0].text,'Dear Alice, welcome'); self.assertTrue(out.paragraphs[0].runs[0].bold); self.assertEqual(out.tables[0].cell(0,0).text,'9.90')
    def test_output_conflict(self):
        self.file('t.txt','{{x}}'); r={'action':'render','template':'t.txt','data':{'x':'value'}}; docs.handle(self.st,r)
        with self.assertRaises(ValueError): docs.handle(self.st,r)

if __name__=="__main__": unittest.main(verbosity=2)
