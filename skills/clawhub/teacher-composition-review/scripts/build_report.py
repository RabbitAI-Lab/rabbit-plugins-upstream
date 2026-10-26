"""把已完成的作文批改结果交付为 Markdown 或 A4 Word，不执行识别、评分或建档。

默认学生版隐藏教师诊断和练习参考答案；只有 --teacher-copy 才导出这些内容。
Word 依赖不可用时生成真实 Markdown，并在 JSON 结果中告知实际路径和降级原因。
"""

from __future__ import annotations

import argparse
import json
import os
import tempfile
from pathlib import Path


RATINGS = {"优秀", "良好", "合格", "待提高"}
FIELDS = {"title", "grade_level", "composition_type", "rating", "rating_reason", "dimensions",
          "highlights", "annotations", "student_comment", "focus", "exercise", "teacher_reference", "second_draft"}


def _text(value, label):
    """要求有实际内容的字符串，避免空证据进入交付报告。"""
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{label}必须是非空文字。")
    return value


def _object(value, fields, label):
    """严格限制对象字段，防止把全文、身份信息或档案私有字段意外导出。"""
    if not isinstance(value, dict) or set(value) != fields:
        raise ValueError(f"{label}字段必须且只能包含：{', '.join(sorted(fields))}。")


def _list(value, label):
    """数组顺序就是报告顺序，不排序、不自动补齐模型未提供的内容。"""
    if not isinstance(value, list):
        raise ValueError(f"{label}必须是列表。")
    return value


def _rating(value):
    """未知题目等情形允许 null，显示为暂不评级，不把缺信息换成低等级。"""
    if value is not None and (not isinstance(value, str) or value not in RATINGS):
        raise ValueError("等级只能是固定四级之一或 null。")


def validate_report(payload):
    """先验证完整报告再写文件，任何错误均不得覆盖既有输出。"""
    _object(payload, FIELDS, "报告")
    for name in ("title", "composition_type", "rating_reason", "student_comment", "teacher_reference"):
        _text(payload[name], name)
    if type(payload["grade_level"]) is not int or payload["grade_level"] not in {3, 4, 5, 6}:
        raise ValueError("年级必须是三至六年级整数。")
    _rating(payload["rating"])
    dimensions = _list(payload["dimensions"], "四维评价")
    names = []
    for item in dimensions:
        _object(item, {"name", "rating", "evidence"}, "维度")
        names.append(_text(item["name"], "维度名"))
        _rating(item["rating"])
        _text(item["evidence"], "评级依据")
    if names != ["内容", "结构", "语言", "规范"]:
        raise ValueError("四维评价按内容、结构、语言、规范提供且不得缺项。")
    for item in _list(payload["highlights"], "亮点"):
        _object(item, {"quote", "reason"}, "亮点")
        for field in item:
            _text(item[field], "亮点原句或理由")
    ids = []
    suggestions = 0
    for item in _list(payload["annotations"], "订正"):
        _object(item, {"id", "kind", "location", "quote", "problem", "revision", "reason"}, "订正")
        if type(item["id"]) is not int or item["id"] <= 0:
            raise ValueError("订正编号必须是正整数。")
        ids.append(item["id"])
        if item["kind"] not in ("correction", "suggestion"):
            raise ValueError("订正类别只能是 correction 或 suggestion。")
        suggestions += item["kind"] == "suggestion"
        for field in ("location", "quote", "problem", "revision", "reason"):
            _text(item[field], "订正内容")
    if ids != list(range(1, len(ids) + 1)):
        raise ValueError("订正编号必须从 1 连续且唯一。")
    if suggestions > 5:
        raise ValueError("可选提升建议最多五条。")
    focus = _list(payload["focus"], "修改重点")
    if not 1 <= len(focus) <= 2:
        raise ValueError("修改重点保留一至两项。")
    for goal in focus:
        _text(goal, "修改重点")
    exercise = payload["exercise"]
    _object(exercise, {"minutes", "task", "hint", "success_criteria", "teacher_reference"}, "练习")
    if type(exercise["minutes"]) is not int or not 3 <= exercise["minutes"] <= 5:
        raise ValueError("练习时长必须是三至五分钟整数。")
    for field in ("task", "hint", "teacher_reference"):
        _text(exercise[field], "练习内容")
    criteria = _list(exercise["success_criteria"], "完成标准")
    if not criteria:
        raise ValueError("练习必须给出完成标准。")
    for criterion in criteria:
        _text(criterion, "完成标准")
    for item in _list(payload["second_draft"], "二稿复查"):
        _object(item, {"goal", "status", "evidence", "authorship"}, "二稿目标")
        if item["status"] not in ("已解决", "部分解决", "仍需练习", "无法核实"):
            raise ValueError("二稿状态不合法；稿件改好不等于能力已掌握。")
        if item["authorship"] not in ("学生独立修改", "教师示范修改", "共同修改", "来源未确认"):
            raise ValueError("二稿修改来源不合法。")
        _text(item["goal"], "二稿目标")
        _text(item["evidence"], "二稿依据")
    return payload


def report_blocks(report, teacher_copy=False):
    """共享内容块同时供两种格式使用，确保学生版不会夹带教师专用内容。"""
    blocks = [("title", "小学作文批改报告" + ("（教师版）" if teacher_copy else "（学生版）")),
              ("text", f"题目：{report['title']}　年级：{report['grade_level']} 年级　类型：{report['composition_type']}"),
              ("text", f"建议等级：{report['rating'] or '暂不评级'}（由老师复核）"),
              ("text", "依据：" + report["rating_reason"]), ("heading", "四维评价")]
    for item in report["dimensions"]:
        blocks.append(("text", f"{item['name']}：{item['rating'] or '暂不评级'}；{item['evidence']}"))
    if report["highlights"]:
        blocks.append(("heading", "值得保留的亮点"))
        for item in report["highlights"]:
            blocks.append(("text", f"原句：{item['quote']}\n好在哪里：{item['reason']}"))
    blocks.extend([("heading", "给你的评语"), ("text", report["student_comment"]), ("heading", "本次修改重点")])
    blocks.extend(("bullet", goal) for goal in report["focus"])
    if report["annotations"]:
        blocks.append(("heading", "编号订正与建议"))
        for item in report["annotations"]:
            kind = "明确错误" if item["kind"] == "correction" else "可选建议"
            blocks.append(("text", f"{item['id']}. {kind} · {item['location']}\n原句：{item['quote']}\n问题：{item['problem']}\n最小修改或方向：{item['revision']}\n原因：{item['reason']}"))
    exercise = report["exercise"]
    blocks.extend([("heading", f"一个 {exercise['minutes']} 分钟小练习"), ("text", exercise["task"]),
                   ("text", "提示：" + exercise["hint"]), ("text", "完成标准：" + "；".join(exercise["success_criteria"]))])
    if report["second_draft"]:
        blocks.append(("heading", "二稿复查"))
        for item in report["second_draft"]:
            blocks.append(("text", f"目标：{item['goal']}\n结果：{item['status']}\n依据：{item['evidence']}\n修改来源：{item['authorship']}"))
        blocks.append(("text", "稿件改好不等于已经掌握；教师示范或共同修改后，还需学生独立尝试。"))
    if teacher_copy:
        blocks.extend([("heading", "教师参考（请勿直接发给学生）"), ("text", report["teacher_reference"]),
                       ("text", "练习参考：" + exercise["teacher_reference"])])
    return blocks


def _save_blocks(blocks, path, output_format):
    """通过同目录临时文件原子替换，不在失败时留下半份报告。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(dir=path.parent, suffix=path.suffix)
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        if output_format == "md":
            # 正文转义 Markdown 控制符，避免原句中的符号被误解释为标题、链接或 HTML。
            def escape(text):
                for character in ("\\", "`", "*", "_", "[", "]", "<", ">", "#"):
                    text = text.replace(character, "\\" + character)
                return text
            prefixes = {"title": "# ", "heading": "## ", "bullet": "- ", "text": ""}
            temporary.write_text("\n\n".join(prefixes[kind] + escape(text) for kind, text in blocks) + "\n", encoding="utf-8")
        else:
            from docx import Document
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            from docx.shared import Mm, Pt, RGBColor

            document = Document()
            section = document.sections[0]
            section.page_width, section.page_height = Mm(210), Mm(297)
            section.top_margin = section.bottom_margin = Mm(18)
            section.left_margin = section.right_margin = Mm(20)
            for name in ("Normal", "Title", "Heading 1"):
                style = document.styles[name]
                style.font.name = "宋体"
                style.font.color.rgb = RGBColor(0, 0, 0)
                style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
                # 清除默认模板的装饰性段落边框，黑白打印时保持简洁。
                for border in style.element.findall(".//" + qn("w:pBdr")):
                    border.getparent().remove(border)
            title_style = document.styles["Title"]
            title_style.font.size = Pt(18)
            title_style.paragraph_format.space_before = Pt(0)
            title_style.paragraph_format.space_after = Pt(8)
            heading_style = document.styles["Heading 1"]
            heading_style.font.size = Pt(12)
            heading_style.paragraph_format.space_before = Pt(8)
            heading_style.paragraph_format.space_after = Pt(4)
            normal = document.styles["Normal"]
            normal.font.size = Pt(11)
            normal.paragraph_format.space_after = Pt(4)
            normal.paragraph_format.line_spacing = 1.15
            for kind, text in blocks:
                if kind == "title":
                    document.add_heading(text, 0)
                elif kind == "heading":
                    document.add_heading(text, 1)
                else:
                    paragraph = document.add_paragraph(("• " if kind == "bullet" else "") + text)
                    # 一个订正项目用多行表达，保持其原句、改法与原因同页，便于学生照着订正。
                    paragraph.paragraph_format.keep_together = True
            # Word 根据正文长度自动分页，页脚仅添加页码，不承诺强塞在一张纸内。
            footer = section.footer.paragraphs[0]
            footer.alignment = 1
            field = OxmlElement("w:fldSimple")
            field.set(qn("w:instr"), "PAGE")
            footer._p.append(field)
            document.save(temporary)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def build_report(request_path, output_path, output_format=None, teacher_copy=False):
    """返回真实交付路径及降级说明，不把报告导出当作成长档案持久化。"""
    request = Path(request_path).expanduser().resolve()
    target = Path(output_path).expanduser().resolve()
    if target == request or (target.exists() and os.path.samefile(request, target)):
        raise ValueError("报告输出不能覆盖输入请求。")
    if output_format is None and target.suffix.lower() not in (".md", ".docx"):
        raise ValueError("输出后缀只能是 .md 或 .docx，或显式指定 --format。")
    actual_format = output_format or target.suffix.lower()[1:]
    if actual_format not in ("md", "docx"):
        raise ValueError("报告格式只能是 md 或 docx。")
    payload = validate_report(json.loads(request.read_text(encoding="utf-8")))
    warning = ""
    if actual_format == "docx":
        try:
            import docx
        except ImportError:
            actual_format = "md"
            warning = "缺少 python-docx，已交付 Markdown；可安装 python-docx 后重新导出 Word。"
    target = target.with_suffix("." + actual_format).resolve()
    # 格式切换及降级会改变最终文件名，必须在最终路径上再次排除输入别名。
    if target == request or (target.exists() and os.path.samefile(request, target)):
        raise ValueError("最终报告路径不能覆盖输入请求。")
    _save_blocks(report_blocks(payload, teacher_copy), target, actual_format)
    return {"status": "success", "output": str(target), "format": actual_format,
            "audience": "teacher" if teacher_copy else "student", "warning": warning}


def main():
    """CLI 只输出结构化状态；调用方应展示 output 中实际生成的文件链接。"""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--request", required=True, help="已确认的结构化报告 JSON")
    parser.add_argument("--output", required=True, help="输出 .md 或 .docx 路径")
    parser.add_argument("--format", choices=("md", "docx"))
    parser.add_argument("--teacher-copy", action="store_true", help="显式包含教师诊断和练习参考")
    args = parser.parse_args()
    try:
        result = build_report(args.request, args.output, args.format, args.teacher_copy)
    except (ValueError, OSError, TypeError) as error:
        print(json.dumps({"status": "error", "message": str(error)}, ensure_ascii=False))
        return 1
    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
