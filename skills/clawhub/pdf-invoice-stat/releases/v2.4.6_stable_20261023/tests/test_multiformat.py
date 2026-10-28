#!/usr/bin/env python3
"""
pdf-invoice-stat v2.4.6 · 五格式回归测试

目的：防止后续改动破坏任一格式的解析能力或表头样式

覆盖：
  1. 格式识别（PDF / XML / OFD / ZIP / RAR）
  2. 真实文件解析（文件不存在时自动跳过，不误报失败）
  3. OFD vs XML 交叉验证（同一发票两种格式字段必须一致）
  4. 表头/样式一致性（v2.4.0 原版 15 列）
  5. PDF 主代码零改动（md5 基线）
  6. 核心函数单元验证（发票号/日期/税率/金额/损坏PDF）

运行：python3 tests/test_multiformat.py
"""
import sys
import os
import hashlib
import zipfile
import tempfile
import shutil

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)

# 真实测试文件位置（不存在则跳过对应项）
FILES = {
    'pdf_haikou': '/tmp/clawhub/20261022_优格海口_文件夹外发票.pdf',
    'pdf_toll':   '/tmp/clawhub/20261022_优格海口_通行费发票.pdf',
    'pdf_sh':     '/tmp/clawhub/20261022_idunique_合并发票.pdf',
    'ofd_zip':    '/tmp/clawhub/INTL_input-国际发票OFD格式.zip',
    'xml_zip':    '/tmp/clawhub/INTL_input-国际发票XML格式.zip',
    'mixed_zip':  '/tmp/clawhub/mixed.zip',
    'ofd_single': '/tmp/clawhub/intl_single.zip',
}

# PDF 主代码 md5 基线（v2.4.4 起未变）
PDF_MD5_BASELINE = '788f0384d9e242b496132003ae5a9398'

V240_HEADERS = ['序号', '发票号码', '开票日期', '购买方信息-名称',
                '购买方信息-统一社会信用代码', '销售方信息-名称',
                '销售方信息-统一社会信用代码', '项目名称',
                '金额', '税率/征收率', '税额', '价税合计', '备注', '报销人', '报销日期']

passed = 0
failed = 0
skipped = 0


def ok(name, detail=''):
    global passed
    passed += 1
    print(f'  ✅ {name}' + (f'  ({detail})' if detail else ''))


def fail(name, detail=''):
    global failed
    failed += 1
    print(f'  ❌ {name}' + (f'  ({detail})' if detail else ''))


def skip(name, detail=''):
    global skipped
    skipped += 1
    print(f'  ⏭️  {name}  ({detail})')


def check(name, cond, detail=''):
    if cond:
        ok(name, detail)
    else:
        fail(name, detail)


# ============================================================
print('=' * 70)
print(' 1. 版本与模块加载')
print('=' * 70)

try:
    from version import __version__, __formats__
    check('version.py 加载', __version__ == '2.4.6', f'v{__version__}')
    check('格式声明完整', set(__formats__) == {'PDF', 'XML', 'OFD', 'ZIP', 'RAR'}, '/'.join(__formats__))
except Exception as e:
    fail('version.py 加载', str(e))

try:
    from format_router import (extract_invoices, detect_format, to_v240,
                               create_excel_any, extract_folder)
    ok('format_router 加载')
except Exception as e:
    fail('format_router 加载', str(e))

try:
    import xml_parser
    ok('xml_parser 加载')
except Exception as e:
    fail('xml_parser 加载', str(e))

try:
    import ofd_parser
    ok('ofd_parser 加载')
except Exception as e:
    fail('ofd_parser 加载', str(e))

# ============================================================
print()
print('=' * 70)
print(' 2. PDF 主代码零改动')
print('=' * 70)
try:
    h = hashlib.md5(open(os.path.join(ROOT, 'invoice_extractor.py'), 'rb').read()).hexdigest()
    check('invoice_extractor.py md5 未变', h == PDF_MD5_BASELINE, h)
except Exception as e:
    fail('md5 校验', str(e))

# ============================================================
print()
print('=' * 70)
print(' 3. 格式识别')
print('=' * 70)
try:
    if os.path.exists(FILES['pdf_haikou']):
        check('PDF 识别', detect_format(FILES['pdf_haikou']) == 'PDF')
    else:
        skip('PDF 识别', '测试文件不存在')

    # 注意区分两种 zip：
    #   单张 OFD（内部含 OFD.xml）→ 识别为 OFD
    #   批量 zip（内含多个 .ofd）→ 识别为 ZIP（由 router 解包后逐个路由）
    if os.path.exists(FILES['ofd_single']):
        d = detect_format(FILES['ofd_single'])
        check('单张 OFD 识别', d == 'OFD', d)
    else:
        skip('单张 OFD 识别', '测试文件不存在')

    if os.path.exists(FILES['ofd_zip']):
        d = detect_format(FILES['ofd_zip'])
        check('批量 OFD(zip) 识别为 ZIP', d == 'ZIP', d)
    else:
        skip('批量 OFD(zip) 识别', '测试文件不存在')

    if os.path.exists(FILES['xml_zip']):
        d = detect_format(FILES['xml_zip'])
        check('XML(ZIP) 识别', d == 'ZIP', d)
    else:
        skip('XML(ZIP) 识别', '测试文件不存在')

    if os.path.exists(FILES['mixed_zip']):
        d = detect_format(FILES['mixed_zip'])
        check('混合 ZIP 识别', d == 'ZIP', d)
    else:
        skip('混合 ZIP 识别', '测试文件不存在')

    # RAR：仅验证能力（unar 存在性）
    import shutil as _sh
    if _sh.which('unar'):
        ok('RAR 支持（unar 可用）')
    else:
        skip('RAR 支持', 'unar 未安装')
except Exception as e:
    fail('格式识别', str(e))

# ============================================================
print()
print('=' * 70)
print(' 4. 真实文件解析')
print('=' * 70)
EXPECTED = {
    'pdf_haikou': (31, 44156.51),
    'pdf_toll':   (11, 166.55),
    'pdf_sh':     (10, 4300.45),
    'ofd_zip':    (5, 9090.00),
    'xml_zip':    (5, 9090.00),
    'mixed_zip':  (11, 4388.45),
}
try:
    for key, (exp_n, exp_amt) in EXPECTED.items():
        f = FILES[key]
        if not os.path.exists(f):
            skip(f'{key} 解析', '测试文件不存在')
            continue
        invs = extract_invoices(f)
        n = len(invs)
        amt = sum(i.get('价税合计', 0) for i in invs)
        check(f'{key} 张数', n == exp_n, f'{n}/{exp_n}')
        check(f'{key} 金额', abs(amt - exp_amt) < 0.05, f'¥{amt:,.2f}')
        # 必填字段非空
        need = ['发票号码', '开票日期', '购买方信息-名称', '销售方信息-名称', '金额', '价税合计']
        missing = {k: sum(1 for i in invs if not i.get(k)) for k in need}
        bad = {k: v for k, v in missing.items() if v > 0}
        check(f'{key} 必填字段完整', not bad, str(bad) if bad else '6/6 全满')
except Exception as e:
    fail('真实文件解析', str(e))

# ============================================================
print()
print('=' * 70)
print(' 5. OFD vs XML 交叉验证')
print('=' * 70)
try:
    if os.path.exists(FILES['ofd_zip']) and os.path.exists(FILES['xml_zip']):
        o = {i['发票号码']: i for i in extract_invoices(FILES['ofd_zip'])}
        x = {i['发票号码']: i for i in extract_invoices(FILES['xml_zip'])}
        check('发票号集合一致', set(o.keys()) == set(x.keys()),
              f'OFD {len(o)} / XML {len(x)}')
        F = ['开票日期', '购买方信息-名称', '购买方信息-统一社会信用代码',
             '销售方信息-名称', '销售方信息-统一社会信用代码',
             '项目名称', '金额', '税率/征收率', '税额', '价税合计']
        total = passed_n = 0
        diffs = []
        for no in set(o) & set(x):
            for f in F:
                total += 1
                if str(o[no].get(f, '')) == str(x[no].get(f, '')):
                    passed_n += 1
                else:
                    diffs.append(f'{no}.{f}')
        check('OFD/XML 字段全一致', passed_n == total,
              f'{passed_n}/{total}' + (f' 差异: {diffs[:3]}' if diffs else ''))
        # 特定 bug 回归点
        if o:
            any_inv = list(o.values())[0]
            check('购方名非标签行', '统一社会信用代码' not in str(any_inv.get('购买方信息-名称', '')),
                  str(any_inv.get('购买方信息-名称', ''))[:22])
            check('购方税号非发票号', not str(any_inv.get('购买方信息-统一社会信用代码', '')).isdigit(),
                  str(any_inv.get('购买方信息-统一社会信用代码', '')))
    else:
        skip('OFD/XML 交叉验证', '测试文件不存在')
except Exception as e:
    fail('OFD/XML 交叉验证', str(e))

# ============================================================
print()
print('=' * 70)
print(' 6. 表头与样式一致性（v2.4.0 原版）')
print('=' * 70)
try:
    import openpyxl
    from format_router import create_excel_any
    for key in ['pdf_haikou', 'ofd_zip', 'xml_zip']:
        f = FILES[key]
        if not os.path.exists(f):
            skip(f'{key} 表头', '测试文件不存在')
            continue
        invs = extract_invoices(f)
        tmp = tempfile.mktemp(suffix='.xlsx')
        create_excel_any(invs, tmp)
        ws = openpyxl.load_workbook(tmp).active
        hdr = [c.value for c in ws[1]]
        check(f'{key} 15 列表头', hdr == V240_HEADERS, f'{len(hdr)} 列')
        check(f'{key} 表头灰底', str(ws['A1'].fill.fgColor.rgb).endswith('D9D9D9'))
        check(f'{key} 表头加粗', ws['A1'].font.bold is True)
        check(f'{key} 冻结 A2', ws.freeze_panes == 'A2', str(ws.freeze_panes))
        check(f'{key} 发票号文本格式', ws.cell(2, 2).number_format == '@')
        check(f'{key} 日期格式', ws.cell(2, 3).number_format == 'YYYY-MM-DD')
        check(f'{key} 金额千分位', ws.cell(2, 9).number_format == '#,##0.00')
        os.remove(tmp)
except Exception as e:
    fail('表头与样式', str(e))

# ============================================================
print()
print('=' * 70)
print(' 7. 核心函数单元验证')
print('=' * 70)
try:
    from invoice_extractor import (find_invoice_number, normalize_date,
                                   extract_tax_rate, extract_amounts, format_rate)
    r = find_invoice_number('电子发票（普通发票）\n发票号码：26312000006362802556')
    check('发票号（同页）', r == '26312000006362802556', str(r))
    r = find_invoice_number('发票号码：\n开票日期：\n购 销\n买 名称：\n下载次数：1\n26442000010830586171')
    check('发票号（维达跨行）', r == '26442000010830586171', str(r))
    r = normalize_date('开票日期：2026年10月10日')
    check('日期（同页）', r == '2026-10-10', str(r))
    r = normalize_date('发票号码：\n开票日期：\n2026年09月21日')
    check('日期（跨行）', r == '2026-09-21', str(r))
    r = extract_tax_rate('税率 6% 税额 4.98')
    check('单税率', r == '6%', str(r))
    r = extract_tax_rate('税率 9% 13% 免税 税额 282.94')
    check('多税率留空（设计规范）', r == '', repr(r))
    r = extract_tax_rate('不征税')
    check('不征税', r == '不征税', str(r))
    r = extract_amounts('合 计 ¥3089.11 ¥30.89\n价税合计（大写）叁仟壹佰贰拾圆整（小写）¥3120.00')
    check('金额提取', abs(r[0]-3089.11)<0.01 and abs(r[1]-30.89)<0.01 and abs(r[2]-3120.00)<0.01, str(r))
    check('税率格式化', format_rate('6%') == 0.06, str(format_rate('6%')))
except Exception as e:
    fail('核心函数', str(e))

try:
    from invoice_extractor import extract_all_invoices
    bad = tempfile.mktemp(suffix='.pdf')
    open(bad, 'w').write('NOT_A_PDF')
    res = extract_all_invoices(bad)
    check('损坏 PDF 优雅返回', res == [], str(res))
    os.remove(bad)
except Exception as e:
    fail('损坏 PDF', str(e))

# ============================================================
print()
print('=' * 70)
print(f' 结果：通过 {passed} · 失败 {failed} · 跳过 {skipped}')
print('=' * 70)
if failed:
    print('\n❌ 有测试未通过，请检查')
    sys.exit(1)
else:
    print('\n✅ 全部通过（跳过的项目需真实测试文件才能验证）')
