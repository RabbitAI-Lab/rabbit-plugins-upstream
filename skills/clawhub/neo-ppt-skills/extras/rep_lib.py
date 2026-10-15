# -*- coding: utf-8 -*-
"""像素复刻通用库：1280x720 画布 -> 16:9 pptx 原生形状。1px = 9525 EMU。"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
import copy

PX = 9525
W, H = 1280, 720

def C(hexs):
    hexs = hexs.lstrip("#")
    return RGBColor(int(hexs[0:2], 16), int(hexs[2:4], 16), int(hexs[4:6], 16))

def px(v):
    return Emu(int(round(v * PX)))

def new_deck():
    prs = Presentation()
    prs.slide_width = px(W)
    prs.slide_height = px(H)
    return prs

def add_slide(prs, bg="#FFFFFF"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, px(W), px(H))
    r.fill.solid(); r.fill.fore_color.rgb = C(bg)
    r.line.fill.background()
    r.shadow.inherit = False
    return slide

def _noshadow(sh):
    sh.shadow.inherit = False

def rect(slide, x, y, w, h, fill=None, line=None, lw=1.0, radius=None, dash=None):
    shp_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(shp_type, px(x), px(y), px(w), px(h))
    if radius:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = C(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = C(line); sh.line.width = Pt(lw)
        if dash:
            ln = sh.line._get_or_add_ln()
            d = ln.makeelement(qn('a:prstDash'), {'val': dash}); ln.append(d)
    else:
        sh.line.fill.background()
    _noshadow(sh)
    return sh

def oval(slide, x, y, w, h, fill=None, line=None, lw=1.0):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, px(x), px(y), px(w), px(h))
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = C(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = C(line); sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    _noshadow(sh)
    return sh

def shape(slide, kind, x, y, w, h, fill=None, line=None, lw=1.0, rot=0):
    sh = slide.shapes.add_shape(getattr(MSO_SHAPE, kind), px(x), px(y), px(w), px(h))
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = C(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = C(line); sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    if rot:
        sh.rotation = rot
    _noshadow(sh)
    return sh

def line(slide, x1, y1, x2, y2, color="#CCCCCC", w=1.0, dash=None):
    ln = slide.shapes.add_connector(1, px(x1), px(y1), px(x2), px(y2))
    ln.line.color.rgb = C(color); ln.line.width = Pt(w)
    if dash:
        lnel = ln.line._get_or_add_ln()
        d = lnel.makeelement(qn('a:prstDash'), {'val': dash}); lnel.append(d)
    ln.shadow.inherit = False
    return ln

def text(slide, x, y, w, h, runs, size=14, color="#333333", bold=False,
         align="l", anchor="t", font="微软雅黑", spacing=1.0, space_after=0):
    """runs: str 或 [段落]; 段落 = str 或 [(txt,{color/bold/size/font}),...]"""
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(runs, str):
        runs = [runs]
    if runs and isinstance(runs[0], tuple):
        runs = [runs]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT, "j": PP_ALIGN.JUSTIFY}[align]
        p.line_spacing = spacing
        if space_after:
            p.space_after = Pt(space_after)
        if isinstance(para, str):
            para = [(para, {})]
        for txt, st in para:
            r = p.add_run(); r.text = txt
            f = r.font
            f.size = Pt(st.get("size", size))
            f.bold = st.get("bold", bold)
            f.name = st.get("font", font)
            f.color.rgb = C(st.get("color", color))
            ea = f._rPr.makeelement(qn('a:ea'), {'typeface': st.get("font", font)})
            f._rPr.append(ea)
    return tb

def shape_text(sh, runs, size=14, color="#FFFFFF", bold=False, align="c", font="微软雅黑", spacing=1.0):
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = px(6)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    if isinstance(runs, str):
        runs = [runs]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        p.line_spacing = spacing
        if isinstance(para, str):
            para = [(para, {})]
        for txt, st in para:
            r = p.add_run(); r.text = txt
            f = r.font
            f.size = Pt(st.get("size", size))
            f.bold = st.get("bold", bold)
            f.name = st.get("font", font)
            f.color.rgb = C(st.get("color", color))
            ea = f._rPr.makeelement(qn('a:ea'), {'typeface': st.get("font", font)})
            f._rPr.append(ea)

def stars(slide, x, y, n, total=5, size=12, color="#8B1E2D", off="#D9D9D9"):
    text(slide, x, y, 160, size + 6,
         [("★" * n, {"color": color}), ("★" * (total - n), {"color": off})], size=size)

def pageno(slide, n, color="#9A9A9A"):
    text(slide, 1150, 668, 90, 30, f"{n:02d}", size=18, color=color, align="r")

def title_std(slide, txt, y=34, size=28, underline=True):
    text(slide, 96, y, 1090, 44, txt, size=size, color="#1A1A1A", bold=True)
    if underline:
        rect(slide, 98, y + 48, 56, 5, fill="#8B1E2D")

def opinion(slide, txt, y=104, icon="#8B1E2D", color="#8B1E2D", h=46, x=96, w=1088, size=14.5, two=False):
    rect(slide, x, y, w, h, fill="#FFFFFF", line="#E3DDD4", lw=1, radius=0.12)
    oval(slide, x + 18, y + h / 2 - 11, 22, 22, fill=icon)
    oval(slide, x + 25, y + h / 2 - 4, 8, 8, fill="#FFFFFF")
    if two:
        text(slide, x + 54, y + 4, w - 70, h - 8, txt, size=size, color=color, bold=True, anchor="m", spacing=1.05)
    else:
        text(slide, x + 54, y, w - 70, h, txt, size=size, color=color, bold=True, anchor="m")
