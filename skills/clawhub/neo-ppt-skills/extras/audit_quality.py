# -*- coding: utf-8 -*-
"""逐套质量审计：
1) 成品pptx渲染 vs 原图的相似度（降采样灰度Pearson相关，3个中段采样页）
2) 项目checker报告的 warning/error 数
3) 结构指标（页数/文本框密度/图片数）
输出 audit_result.json + 控制台排名表"""
import sys, os, glob, json, subprocess, io
sys.stdout.reconfigure(encoding="utf-8")
from pathlib import Path
from PIL import Image
import pymupdf

BASE = Path(r"C:\Users\w\Documents\zcode_file\z\银行ppt\图片转PPT")
CHENG = BASE / "成品"
IMG = BASE.parent / "图片"
SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"
TMP = Path(os.environ.get("TEMP", r"C:\Users\w\AppData\Local\Temp")) / "gaudit"
TMP.mkdir(exist_ok=True)

def gray(im, size=(160, 90)):
    return im.convert("L").resize(size)

def corr(a, b):
    pa, pb = list(a.getdata()), list(b.getdata())
    na, nb = sum(pa)/len(pa), sum(pb)/len(pb)
    num = sum((x-na)*(y-nb) for x, y in zip(pa, pb))
    da = (sum((x-na)**2 for x in pa)) ** 0.5
    db = (sum((y-nb)**2 for y in pb)) ** 0.5
    return num / (da * db) if da and db else 0.0

def src_dir(prefix):
    for d in glob.glob(str(IMG / (prefix + "_*"))):
        return Path(d)
    return None

decks = []
for f in sorted(CHENG.glob("*.pptx")):
    m = __import__("re").match(r"^(\d+)(?:_|_.*?上篇|_.*?下篇)?", f.stem)
    decks.append((f, m.group(1)))

results = []
for f, prefix in decks:
    tag = "bank" + prefix
    # 渲染
    pdf = TMP / (f.stem + ".pdf")
    if not pdf.exists():
        r = subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf", "--outdir", str(TMP), str(f)],
                           capture_output=True, timeout=300)
        if not pdf.exists():
            results.append({"deck": tag, "file": f.name, "error": "render fail"}); continue
    doc = pymupdf.open(str(pdf))
    npg = doc.page_count
    # 采样3个中段页（避开封面/结尾），deck页i ↔ 源图 i+1（bank10）/ i（其余）
    offsets = [1] if prefix == "10" else [0]
    sims = []
    srcd = src_dir(prefix)
    for pg in sorted({npg // 2, max(2, npg // 3), min(npg - 2, npg * 2 // 3)}):
        for off in offsets:
            si = pg + off
            if not srcd: continue
            pix = doc[pg - 1].get_pixmap(dpi=60)
            mine = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            s = gray(mine)
            best = None
            for w in range(-1, 4):  # 滑动窗口自校正页码偏移
                cand = list(srcd.glob(f"{si + w:03d}.png")) or list(srcd.glob(f"{si + w:03d}.jpg"))
                if not cand: continue
                c = corr(s, gray(Image.open(cand[0])))
                if best is None or c > best:
                    best = c
            if best is not None:
                sims.append(best)
    sim = round(sum(sims) / len(sims), 3) if sims else None
    # checker报告
    proj = BASE / (tag + "_ppt169_*")
    proj_dirs = glob.glob(str(proj))
    warns = errs = None
    for pd in proj_dirs:
        rp = Path(pd) / "validation" / "svg_quality_report.json"
        if rp.exists():
            try:
                rj = json.load(io.open(rp, encoding="utf-8"))
                files = rj.get("files", rj if isinstance(rj, list) else [])
                w = e = 0
                items = files if isinstance(files, list) else files.get("items", [])
                for it in items:
                    w += len(it.get("warnings", [])); e += len(it.get("errors", []))
                warns, errs = w, e
            except Exception:
                pass
    # 结构
    from pptx import Presentation
    p = Presentation(f)
    t = pic = 0
    def walk(shs):
        global t, pic
        for s in shs:
            if s.shape_type == 13: pic += 1
            if getattr(s, "has_text_frame", False) and s.text_frame.text.strip(): t += 1
            if s.shape_type == 6: walk(s.shapes)
    for sl in p.slides: walk(sl.shapes)
    n = len(p.slides.__iter__.__self__._sldIdLst)
    results.append({"deck": tag, "file": f.name, "pages": n, "textframes": t,
                    "tp": round(t / n, 1), "pics": pic, "sim": sim,
                    "warns": warns, "errs": errs})
    print(f"{tag}: pages={n} t/p={t/n:.0f} sim={sim} warn={warns} err={errs}", flush=True)

json.dump(results, io.open(BASE / "audit_result.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print("\n=== 相似度排名（低=复刻质量差，淘汰候选）===")
for r in sorted([r for r in results if r.get("sim") is not None], key=lambda x: x["sim"]):
    print(f"{r['deck']:>8} sim={r['sim']:.3f}  pages={r['pages']:>2}  t/p={r['tp']:>4.0f}  warn={r['warns']}")
