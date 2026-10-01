#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
docx-writer 引擎（配置驱动 / 模板可插拔 / 内容与排版解耦）

设计目标
--------
- 引擎只负责「按模板配置渲染」，不关心具体内容。
- 模板 = 一份 YAML/JSON 配置：页面(页边距/纸张/背景) + 样式表(字体/字号/颜色/对齐/缩进/行距)
  + 脚注(字体/字号/编号方式) + 图片/表格默认值 + 后处理开关。
- 内容既可用 Python API 逐句写，也可用声明式内容文件（见 build_from_content.py）。

快速用法
--------
    from docx_engine import Template, Builder
    tpl = Template.load('academic')          # 名字或 yaml 路径
    b = Builder(tpl)
    b.title('标题')
    b.heading('一、引言', level=1)
    b.para('正文。{F1}', fn_map={'F1': '注释文本'})
    b.image('fig.png', caption='图1  示意')
    b.table(['列1', '列2'], [['a', 'b']], caption='表1  示例')
    b.save('out.docx')

模板放置：docx-writer/templates/<name>.yaml ；`base_docx` 指向承载脚注样式的基底 docx。
"""
import os, re, sys, json, copy

from lxml import etree
import docx as dx
from docx.shared import Pt, Cm, Mm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

try:
    import yaml  # PyYAML
except Exception:  # pragma: no cover
    yaml = None

W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
WP_NS = 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'
R_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
MATH_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
def w(tag): return '{%s}%s' % (W_NS, tag)

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL_DIR = os.path.dirname(HERE)
TEMPLATES_DIR = os.path.join(SKILL_DIR, 'templates')
if HERE not in sys.path:          # 便于 import 同目录的 latex_formula
    sys.path.insert(0, HERE)

ALIGN = {
    'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
    'right': WD_ALIGN_PARAGRAPH.RIGHT, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY,
}

PAGE_SIZE = {'A4': (Cm(21.0), Cm(29.7)), 'LETTER': (Inches(8.5), Inches(11)),
             'A3': (Cm(29.7), Cm(42.0)), 'A5': (Cm(14.8), Cm(21.0))}


# ───────────────────────── 单位/取值解析 ─────────────────────────
def to_len(v, default=None):
    if v is None:
        return default
    if isinstance(v, (int, float)):
        return Pt(float(v))
    s = str(v).strip().lower()
    m = re.match(r'^(-?[\d.]+)\s*(cm|mm|pt|in|px)?$', s)
    if not m:
        return default
    n = float(m.group(1)); u = m.group(2) or 'pt'
    return {'cm': Cm, 'mm': Mm, 'pt': Pt, 'in': Inches, 'px': lambda x: Pt(x * 0.75)}[u](n)


def to_spacing(v, default=None):
    """行距：数字=倍数；'18pt'=固定值。"""
    if v is None:
        return default
    s = str(v).strip().lower()
    if s.endswith('pt'):
        return Pt(float(s[:-2]))
    try:
        return float(s)
    except ValueError:
        return default


def to_color(v, default=None):
    if v is None:
        return default
    s = str(v).strip().lstrip('#')
    if s in ('', 'none', 'null'):
        return None
    return RGBColor.from_string(s.upper())


# ───────────────────────── 模板 ─────────────────────────
STYLE_DEFAULT = {
    'cn': '宋体', 'en': 'Times New Roman', 'size': '10.5pt', 'bold': False,
    'italic': False, 'color': None, 'align': None, 'indent': None,
    'line_spacing': None, 'space_before': None, 'space_after': None,
}


class Template:
    """模板配置。ref = 模板名 或 yaml/json 路径。"""

    def __init__(self, cfg, base_dir):
        self.cfg = cfg
        self.base_dir = base_dir
        self.name = cfg.get('name', os.path.splitext(os.path.basename(base_dir))[0])

    @classmethod
    def load(cls, ref):
        path = ref
        if not os.path.exists(path):
            for ext in ('yaml', 'yml', 'json'):
                cand = os.path.join(TEMPLATES_DIR, '%s.%s' % (ref, ext))
                if os.path.exists(cand):
                    path = cand
                    break
        if not os.path.exists(path):
            raise FileNotFoundError('template not found: %s' % ref)
        with open(path, 'r', encoding='utf-8') as f:
            text = f.read()
        if path.endswith('.json'):
            cfg = json.loads(text)
        else:
            if yaml is None:
                raise RuntimeError('PyYAML not installed; use a .json template')
            cfg = yaml.safe_load(text)
        return cls(cfg or {}, os.path.dirname(os.path.abspath(path)))

    # --- 便捷访问 ---
    @property
    def page(self): return self.cfg.get('page', {})

    def base_docx(self):
        b = self.cfg.get('base_docx')
        if not b:
            return None
        for cand in (os.path.join(self.base_dir, b), os.path.join(SKILL_DIR, b), b):
            if os.path.exists(cand):
                return cand
        return None

    def style(self, name):
        st = dict(STYLE_DEFAULT)
        st.update(self.cfg.get('styles', {}).get(name, {}))
        return st

    def style_names(self):
        return list(self.cfg.get('styles', {}).keys())

    @property
    def footnote(self): return self.cfg.get('footnote', {})

    @property
    def features(self): return self.cfg.get('features', {})


# ───────────────────────── 构建器 ─────────────────────────
class Builder:
    def __init__(self, template):
        self.t = template
        base = template.base_docx()
        self.doc = dx.Document(base) if base else dx.Document()
        self._prepare_body()
        self._clean_footnotes()
        self._setup_page(template.page)
        self._setup_footnote(template.footnote)
        self.fn_map = {}
        self.lx_map = {}

    # ---- 初始化 ----
    def _prepare_body(self):
        body = self.doc.element.body
        for child in list(body):
            if child.tag is etree.Comment:
                continue
            if etree.QName(child).localname != 'sectPr':
                body.remove(child)

    def _clean_footnotes(self):
        try:
            for rel in self.doc.part.rels.values():
                if 'footnotes' in rel.reltype:
                    fn_elem = rel.target_part.element
                    for fn in list(fn_elem):
                        if fn.get(w('type'), '') not in ('separator', 'continuationSeparator'):
                            fn_elem.remove(fn)
                    break
        except Exception:
            pass

    def _setup_page(self, page):
        if not page:
            return
        for sec in self.doc.sections:
            if page.get('size'):
                size = PAGE_SIZE.get(str(page['size']).upper())
                if size:
                    sec.page_width, sec.page_height = size
            m = page.get('margin', {})
            if 'top' in m: sec.top_margin = to_len(m['top'])
            if 'bottom' in m: sec.bottom_margin = to_len(m['bottom'])
            if 'left' in m: sec.left_margin = to_len(m['left'])
            if 'right' in m: sec.right_margin = to_len(m['right'])
            if page.get('orientation') == 'landscape':
                sec.orientation = 1
        bg = page.get('background')
        if bg:
            color = bg.get('color') if isinstance(bg, dict) else bg
            if color:
                el = etree.Element(w('background'))
                el.set(w('color'), str(color).lstrip('#'))
                self.doc.element.insert(0, el)

    def _setup_footnote(self, fn):
        if not fn.get('enabled', True):
            return
        numbering = fn.get('numbering', 'per_page')
        for sec in self.doc.sections:
            sectPr = sec._sectPr
            fpr = sectPr.find(w('footnotePr'))
            if fpr is None:
                fpr = etree.Element(w('footnotePr'))
                sectPr.insert(0, fpr)
            if fn.get('num_fmt'):
                nf = fpr.find(w('numFmt'))
                if nf is None:
                    nf = etree.SubElement(fpr, w('numFmt'))
                nf.set(w('val'), str(fn['num_fmt']))
            if numbering:
                nr = fpr.find(w('numRestart'))
                if nr is None:
                    nr = etree.SubElement(fpr, w('numRestart'))
                nr.set(w('val'), 'eachPage' if numbering == 'per_page' else 'continuous')

    # ---- 样式应用 ----
    def _font(self, run, st):
        cn = st.get('cn', '宋体'); en = st.get('en', 'Times New Roman')
        run.font.name = en
        if st.get('size'):
            run.font.size = to_len(st['size'])
        if st.get('bold'):
            run.font.bold = True
        if st.get('italic'):
            run.font.italic = True
        col = to_color(st.get('color'))
        if col is not None:
            run.font.color.rgb = col
        rPr = run._element.find(w('rPr'))
        if rPr is None:
            rPr = etree.SubElement(run._element, w('rPr'))
        rf = rPr.find(w('rFonts'))
        if rf is None:
            rf = etree.SubElement(rPr, w('rFonts'))
        rf.set(w('eastAsia'), cn); rf.set(w('ascii'), en); rf.set(w('hAnsi'), en)

    def _para_format(self, p, st):
        if st.get('align'):
            p.alignment = ALIGN.get(st['align'], WD_ALIGN_PARAGRAPH.LEFT)
        pf = p.paragraph_format
        if st.get('indent') is not None:
            pf.first_line_indent = to_len(st['indent'], Pt(0))
        ls = to_spacing(st.get('line_spacing'))
        if ls is not None:
            pf.line_spacing = ls
        sb = to_len(st.get('space_before'))
        if sb is not None: pf.space_before = sb
        sa = to_len(st.get('space_after'))
        if sa is not None: pf.space_after = sa

    def _add_runs(self, p, text, st):
        """支持 {F1} 脚注占位 与 {X1} 行内公式占位（配合 self.fn_map / self.lx_map）。"""
        if not text:
            return
        for seg in re.split(r'(\{[A-Za-z0-9_]+\})', text):
            if not seg:
                continue
            m = re.fullmatch(r'\{([A-Za-z0-9_]+)\}', seg)
            if m and m.group(1) in self.fn_map:
                self._footnote(p, self.fn_map[m.group(1)], st)
            elif m and m.group(1) in self.lx_map:
                self._inline_formula(p, self.lx_map[m.group(1)], st)
            else:
                r = p.add_run(seg)
                self._font(r, st)

    # ---- 脚注 ----
    def _footnote(self, p, text, st=None):
        fn = self.t.footnote
        f = p.add_footnote(text)
        for r in f.findall('.//' + w('r')):
            t = r.find(w('t'))
            if t is None or not (t.text and t.text.strip()):
                continue
            rPr = r.find(w('rPr'))
            if rPr is None:
                rPr = etree.SubElement(r, w('rPr'))
            rf = rPr.find(w('rFonts'))
            if rf is None:
                rf = etree.SubElement(rPr, w('rFonts'))
            rf.set(w('eastAsia'), fn.get('cn', '宋体'))
            rf.set(w('ascii'), fn.get('en', 'Times New Roman'))
            rf.set(w('hAnsi'), fn.get('en', 'Times New Roman'))
            sz = to_len(fn.get('size', '9pt'))
            hp = str(int(sz.pt * 2))
            for tag in ('sz', 'szCs'):
                el = rPr.find(w(tag))
                if el is None:
                    el = etree.SubElement(rPr, w(tag))
                el.set(w('val'), hp)

    # ---- 公共 API ----
    def para(self, text='', style='body', fn_map=None, lx_map=None):
        if fn_map:
            self.fn_map.update(fn_map)
        if lx_map:
            self.lx_map.update(lx_map)
        st = self.t.style(style)
        p = self.doc.add_paragraph()
        self._para_format(p, st)
        self._add_runs(p, text, st)
        return p

    def title(self, text, style='title'):
        return self.para(text, style=style)

    def heading(self, text, level=1, style=None):
        return self.para(text, style=style or ('h%d' % level))

    def caption(self, text, style='caption'):
        return self.para(text, style=style)

    def image(self, path_or_bytes, caption=None, width=None,
              style='image', caption_style='caption'):
        st = self.t.style(style)
        wd = to_len(width) if width else to_len(st.get('width'), Cm(14))
        p = self.doc.add_paragraph()
        p.alignment = ALIGN.get(st.get('align', 'center'), WD_ALIGN_PARAGRAPH.CENTER)
        p.paragraph_format.first_line_indent = Pt(0)
        src = self._image_source(path_or_bytes)
        run = p.add_run()
        run.add_picture(src, width=wd)
        self._patch_inline(p)
        if caption:
            self.caption(caption, style=caption_style)
        return p

    @staticmethod
    def _image_source(path_or_bytes):
        """返回可用于 add_picture 的路径；相机 JPEG 先剥离 EXIF。"""
        import io
        from PIL import Image
        if isinstance(path_or_bytes, (bytes, bytearray)):
            data = bytes(path_or_bytes)
        else:
            if not os.path.exists(path_or_bytes):
                raise FileNotFoundError(path_or_bytes)
            with open(path_or_bytes, 'rb') as f:
                data = f.read()
        try:
            im = Image.open(io.BytesIO(data)).convert('RGB')
        except Exception:
            tmp = os.path.join(os.environ.get('TEMP', '.'), 'dw_%d.jpg' % id(data))
            with open(tmp, 'wb') as f:
                f.write(data)
            return tmp
        buf = io.BytesIO()
        im.thumbnail((1600, 1600))
        im.save(buf, 'JPEG', quality=88)
        tmp = os.path.join(os.environ.get('TEMP', '.'),
                           'dw_%d.jpg' % abs(hash(data)))
        with open(tmp, 'wb') as f:
            f.write(buf.getvalue())
        return tmp

    def _patch_inline(self, p):
        for inline in p._p.iter('{%s}inline' % WP_NS):
            for a in ('distT', 'distB', 'distL', 'distR'):
                inline.set(a, '0')
            ext = inline.find('{%s}extent' % WP_NS)
            if ext is not None and inline.find('{%s}effectExtent' % WP_NS) is None:
                ee = etree.Element('{%s}effectExtent' % WP_NS)
                for k in ('l', 't', 'r', 'b'):
                    ee.set(k, '0')
                ext.addnext(ee)

    # ---- LaTeX 公式（图片 或 Word 原生 OMML） ----
    def _formula_mode(self, override=None):
        """决定公式呈现方式：image（XeLaTeX 紧裁图片）| omml（Word 原生可编辑）| auto。"""
        st = self.t.style('formula')
        mode = str(override or st.get('mode') or 'image').lower()
        if mode == 'auto':
            from latex_formula import find_pandoc
            mode = 'omml' if find_pandoc(st.get('pandoc')) else 'image'
        return mode

    def _omml_element(self, latex, display):
        from latex_formula import latex_to_omml
        return etree.fromstring(latex_to_omml(latex, display=display).encode('utf-8'))

    def _formula_png(self, latex, display=True, cache_dir=None, dpi=300, border='3pt'):
        from latex_formula import render_formulas
        cd = cache_dir or os.path.join(os.getcwd(), '_docx_formula_cache')
        return render_formulas([latex], cd, dpi=dpi, display=display, border=border)[0]

    def formula(self, latex, caption=None, width=None, style='formula',
                caption_style='caption', cache_dir=None, dpi=300, scale=None,
                mode=None):
        """插入行间公式。mode=None 用模板 formula 样式的 mode；'omml'=Word 原生可编辑，'image'=紧裁图片。"""
        st = self.t.style(style)
        eff = self._formula_mode(mode)
        p = self.doc.add_paragraph()
        self._para_format(p, st)
        p.paragraph_format.first_line_indent = Pt(0)
        if eff == 'omml':
            p._p.append(self._omml_element(latex, True))
        else:
            it = self._formula_png(latex, True, cache_dir, dpi)
            sc = float(scale if scale is not None else st.get('scale', 1.0))
            wd = to_len(width) if width else Pt(it['width_pt'] * sc)
            run = p.add_run()
            run.add_picture(it['png'], width=wd)
            self._patch_inline(p)
        if caption:
            self.caption(caption, style=caption_style)
        return p

    def _inline_formula(self, p, latex, st, cache_dir=None, dpi=300, scale=None,
                        mode=None):
        """把公式作为行内元素追加到段落 p（原生 OMML 或缩放图片）。"""
        if self._formula_mode(mode) == 'omml':
            p._p.append(self._omml_element(latex, False))
            return None
        fst = self.t.style('formula')
        it = self._formula_png(latex, False, cache_dir, dpi,
                               border=fst.get('inline_border', '1pt'))
        if scale is None:
            scale = (to_len(st.get('size', '10.5pt')).pt / 12.0
                     * float(fst.get('inline_scale', 1.0)))
        run = p.add_run()
        run.add_picture(it['png'], width=Pt(it['width_pt'] * float(scale)))
        self._patch_inline(p)
        return run

    def table(self, headers, rows, caption=None, style='table',
              caption_style='caption'):
        tst = self.t.style(style)
        mode = tst.get('mode', 'three_line')
        if caption:
            self.caption(caption, style=caption_style)
        t = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        t.style = tst.get('word_style', 'Table Grid')
        # 清边框
        for row in t.rows:
            for cell in row.cells:
                self._clear_borders(cell)
        hst = dict(STYLE_DEFAULT); hst.update(tst.get('header', {}))
        bst = dict(STYLE_DEFAULT); bst.update(tst.get('cell', {}))
        for i, x in enumerate(headers):
            cell = t.rows[0].cells[i]
            cell.text = ''
            run = cell.paragraphs[0].add_run(str(x))
            self._font(run, hst)
            cell.paragraphs[0].alignment = ALIGN.get(hst.get('align', 'center'))
            if mode == 'three_line':
                self._edge_border(cell, 'top', 'single', '4')
            self._edge_border(cell, 'bottom', 'single', '4')
        for ri, row in enumerate(rows):
            for ci, v in enumerate(row):
                cell = t.rows[ri + 1].cells[ci]
                cell.text = ''
                run = cell.paragraphs[0].add_run(str(v))
                self._font(run, bst)
                cell.paragraphs[0].alignment = ALIGN.get(bst.get('align', 'center'))
                if mode == 'three_line' and ri == len(rows) - 1:
                    self._edge_border(cell, 'bottom', 'single', '4')
                elif mode == 'grid':
                    for s in ('top', 'left', 'bottom', 'right'):
                        self._edge_border(cell, s, 'single', '4')
        return t

    @staticmethod
    def _clear_borders(cell):
        tcPr = cell._tc.get_or_add_tcPr()
        for b in tcPr.findall(w('tcBorders')):
            tcPr.remove(b)

    @staticmethod
    def _edge_border(cell, edge, val='single', sz='4'):
        tcPr = cell._tc.get_or_add_tcPr()
        b = tcPr.find(w('tcBorders'))
        if b is None:
            b = etree.SubElement(tcPr, w('tcBorders'))
        el = b.find(w(edge))
        if el is None:
            el = etree.SubElement(b, w(edge))
        el.set(w('val'), val); el.set(w('sz'), sz)

    # ---- 收尾 ----
    def save(self, path):
        if self.t.features.get('fix_quotes', True):
            fix_quotes(self.doc.element.body)
        if self.t.features.get('superscript_footnotes', True):
            self._superscript()
        self.doc.save(path)
        return path

    def _superscript(self):
        for p in self.doc.element.body.findall('.//' + w('p')):
            for r in p.findall('.//' + w('r')):
                if r.find(w('footnoteReference')) is not None:
                    rPr = r.find(w('rPr'))
                    if rPr is None:
                        rPr = etree.SubElement(r, w('rPr'))
                    va = rPr.find(w('vertAlign'))
                    if va is None:
                        va = etree.SubElement(rPr, w('vertAlign'))
                    va.set(w('val'), 'superscript')


# ───────────────────────── 引号字体后处理（沿用既有实现） ─────────────────────────
def _has_cjk(text):
    for c in text:
        if '\u4e00' <= c <= '\u9fff' or '\u3000' <= c <= '\u303f' or '\uff00' <= c <= '\uffef':
            return True
    return False


def _fix_pair(text, mark, lq, rq):
    out, i = [], 0
    while i < len(text):
        if text[i] == mark:
            j = text.find(mark, i + 1)
            if j != -1:
                inner = text[i + 1:j]
                out.append((lq + inner + rq) if _has_cjk(inner) else (mark + inner + mark))
                i = j + 1
                continue
        out.append(text[i]); i += 1
    return ''.join(out)


def fix_quotes_in_text(old):
    return _fix_pair(_fix_pair(old, '"', '\u201c', '\u201d'), "'", '\u2018', '\u2019')


def _font_attr(rPr, key, default):
    if rPr is None:
        return default
    rf = rPr.find(w('rFonts'))
    if rf is not None and rf.get(w(key)):
        return rf.get(w(key))
    return default


def _copy_font_el(src, tag, dst):
    s = src.find(w(tag))
    if s is not None:
        d = etree.SubElement(dst, w(tag))
        d.set(w('val'), s.get(w('val')))


def fix_quotes(parent):
    for p in parent.findall('.//' + w('p')):
        for r in list(p.findall(w('r'))):
            t = r.find(w('t'))
            if t is None or t.text is None:
                continue
            new_text = fix_quotes_in_text(t.text)
            if new_text == t.text:
                continue
            qchars = ('\u201c', '\u201d', '\u2018', '\u2019')
            segs, i = [], 0
            while i < len(new_text):
                c = new_text[i]
                if c in qchars:
                    segs.append((c, True)); i += 1
                else:
                    j = i
                    while j < len(new_text) and new_text[j] not in qchars:
                        j += 1
                    if j > i:
                        segs.append((new_text[i:j], False))
                    i = j
            if len(segs) <= 1:
                continue
            orig = r.find(w('rPr'))
            idx = list(p).index(r)
            p.remove(r)
            for seg_text, is_quote in reversed(segs):
                if not seg_text:
                    continue
                nr = etree.Element(w('r'))
                nrPr = etree.SubElement(nr, w('rPr'))
                nrf = etree.SubElement(nrPr, w('rFonts'))
                if is_quote:
                    cjk = _font_attr(orig, 'eastAsia', '宋体')
                    nrf.set(w('eastAsia'), cjk); nrf.set(w('ascii'), cjk); nrf.set(w('hAnsi'), cjk)
                else:
                    nrf.set(w('eastAsia'), _font_attr(orig, 'eastAsia', '宋体'))
                    nrf.set(w('ascii'), _font_attr(orig, 'ascii', 'Times New Roman'))
                    nrf.set(w('hAnsi'), _font_attr(orig, 'hAnsi', 'Times New Roman'))
                if orig is not None:
                    for tag in ('sz', 'szCs'):
                        _copy_font_el(orig, tag, nrPr)
                nt = etree.SubElement(nr, w('t'))
                nt.text = seg_text
                p.insert(idx, nr)
