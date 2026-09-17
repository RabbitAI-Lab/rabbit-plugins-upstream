# -*- coding: utf-8 -*-
"""把已定稿的 PDF 版式逐页封装为 PPTX。

策略：
  · 幻灯片尺寸 = 792pt × 594pt = 11in × 8.25in（与原稿页面完全同尺寸，无异缩放）
  · 每页以 200dpi 位图满幅铺底，保证自绘 SVG 图表 100% 保真
  · 每页写入演讲者备注（该页全文），使 PPTX 可检索、可复制文字
用法： python mkpptx.py <in.pdf> <out.pptx> [dpi=200]
"""
import os
import re
import subprocess
import sys
import tempfile

from pptx import Presentation
from pptx.util import Inches, Pt

PT = 12700  # 1pt = 12700 EMU


def page_count(pdf):
    out = subprocess.run(['pdfinfo', pdf], capture_output=True, text=True).stdout
    return int(re.search(r'^Pages:\s+(\d+)', out, re.M).group(1))


def page_text(pdf, n):
    t = subprocess.run(['pdftotext', '-f', str(n), '-l', str(n), '-layout', pdf, '-'],
                       capture_output=True, text=True).stdout
    # 压掉页脚免责声明与文件编号，备注里保留正文即可
    lines = [re.sub(r'\s+$', '', l) for l in t.splitlines()]
    lines = [l for l in lines if l.strip()]
    drop = ('本文件仅供派发于合格境内机构投资者', 'MKTG')
    lines = [l for l in lines if not l.lstrip().startswith(drop)]
    return '\n'.join(lines)


def render_pages(pdf, dpi, tmpdir):
    subprocess.run(['pdftoppm', '-r', str(dpi), '-png', pdf, os.path.join(tmpdir, 'pg')], check=True)
    files = sorted(f for f in os.listdir(tmpdir) if f.endswith('.png'))
    return [os.path.join(tmpdir, f) for f in files]


def main():
    src = sys.argv[1]
    dst = sys.argv[2]
    dpi = int(sys.argv[3]) if len(sys.argv) > 3 else 200

    n = page_count(src)
    prs = Presentation()
    # 与原稿同尺寸：792×594pt
    prs.slide_width = 792 * PT
    prs.slide_height = 594 * PT
    blank = prs.slide_layouts[6]

    with tempfile.TemporaryDirectory() as td:
        imgs = render_pages(src, dpi, td)
        assert len(imgs) == n, 'rendered %d != %d pages' % (len(imgs), n)
        for i, img in enumerate(imgs, start=1):
            s = prs.slides.add_slide(blank)
            s.shapes.add_picture(img, 0, 0, width=prs.slide_width, height=prs.slide_height)
            txt = page_text(src, i)
            if txt:
                tf = s.notes_slide.notes_text_frame
                tf.text = txt
                for p in tf.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(11)

    prs.save(dst)
    print('PPTX 已写出: %s  %d 页  %d bytes' % (dst, n, os.path.getsize(dst)))


if __name__ == '__main__':
    main()
