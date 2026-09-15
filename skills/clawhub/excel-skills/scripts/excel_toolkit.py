# -*- coding: utf-8 -*-
"""
excel_toolkit.py - Excel一站式工具集

整合自 workspace/skills 六个技能的Excel能力，配合 excel-skills/SKILL.md 使用。
永久保留脚本，清理临时文件时不得删除（2026-07-24教训）。

用法:
    python excel_toolkit.py info <file>                          查看文件概况
    python excel_toolkit.py convert <src> <dst>                  xlsx/csv/json互转
    python excel_toolkit.py report <src> <dst> [sheet]           生成规范样式报表
    python excel_toolkit.py pdf2xlsx <pdf> <dst> [page]          PDF表格转Excel
    python excel_toolkit.py merge <file> <dst>                   合并多Sheet到一个Sheet
"""

import sys
import os

sys.stdout.reconfigure(encoding='utf-8')

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def read_data(file_path):
    """读取xlsx/csv/json为DataFrame"""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == '.csv':
        return pd.read_csv(file_path)
    if ext in ('.xlsx', '.xls'):
        return pd.read_excel(file_path, engine='openpyxl')
    if ext == '.json':
        return pd.read_json(file_path)
    raise ValueError(f"不支持的格式: {file_path}")


def write_report(df, dst, sheet_name='数据'):
    """生成规范样式报表：表头样式 + 冻结首行 + 列宽自适应

    样式铁律（2026-07-24教训）：每种样式独立构造、独立赋值，
    禁止封装混合类型返回值，禁止跨单元格共享样式对象后修改。
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    # 写入表头和数据
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append(list(row))

    # 表头样式：每种类型独立构造（铁律1）
    header_font = Font(name='微软雅黑', bold=True, size=11, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    center_align = Alignment(horizontal='center', vertical='center')
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    # 冻结首行
    ws.freeze_panes = 'A2'

    # 列宽自适应（按内容长度估算，中文字符按2倍计）
    from openpyxl.utils import get_column_letter
    for idx, col in enumerate(df.columns, 1):
        header_len = sum(2 if ord(c) > 127 else 1 for c in str(col))
        sample_len = 0
        if len(df) > 0:
            sample = df[col].astype(str).head(100)
            sample_len = max((sum(2 if ord(c) > 127 else 1 for c in s) for s in sample), default=0)
        width = min(max(header_len, sample_len) + 4, 50)
        ws.column_dimensions[get_column_letter(idx)].width = width

    wb.save(dst)
    return dst


def pdf_to_excel(pdf_path, dst, page=0):
    """PDF表格提取转Excel"""
    import pdfplumber
    with pdfplumber.open(pdf_path) as pdf:
        if page >= len(pdf.pages):
            raise ValueError(f"页码超出范围: 共{len(pdf.pages)}页")
        tables = pdf.pages[page].extract_tables()
    if not tables:
        raise ValueError(f"第{page}页未检测到表格")
    merged = []
    for t in tables:
        merged.extend(t)
    df = pd.DataFrame(merged[1:], columns=merged[0])
    return write_report(df, dst, sheet_name=f'第{page + 1}页')


def merge_sheets(src, dst, add_source_col=True):
    """合并工作簿内所有Sheet到一个Sheet"""
    wb = openpyxl.load_workbook(src, data_only=True)
    frames = []
    for name in wb.sheetnames:
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        df = pd.DataFrame(rows[1:], columns=rows[0])
        if add_source_col:
            df.insert(0, '_来源Sheet', name)
        frames.append(df)
    if not frames:
        raise ValueError("工作簿内无数据")
    merged = pd.concat(frames, ignore_index=True)
    return write_report(merged, dst, sheet_name='合并结果')


def show_info(file_path):
    """查看文件概况"""
    if file_path.endswith(('.xlsx', '.xls')):
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        print(f"文件: {file_path}")
        print(f"Sheet数: {len(wb.sheetnames)}")
        for name in wb.sheetnames:
            ws = wb[name]
            print(f"  - {name}: {ws.max_row}行 x {ws.max_column}列")
        wb.close()
    else:
        df = read_data(file_path)
        print(f"文件: {file_path}")
        print(f"行数: {len(df)}, 列数: {len(df.columns)}")
        print(f"列名: {', '.join(map(str, df.columns))}")
        print(f"缺失值: {df.isnull().sum().sum()}")


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        return 1
    cmd, src = sys.argv[1], sys.argv[2]

    if cmd == 'info':
        show_info(src)
    elif cmd == 'convert':
        if len(sys.argv) < 4:
            print("用法: convert <src> <dst>")
            return 1
        dst = sys.argv[3]
        df = read_data(src)
        if dst.endswith('.csv'):
            df.to_csv(dst, index=False, encoding='utf-8-sig')
        elif dst.endswith('.xlsx'):
            write_report(df, dst)
        elif dst.endswith('.json'):
            df.to_json(dst, orient='records', force_ascii=False, indent=2)
        else:
            raise ValueError(f"不支持的目标格式: {dst}")
        print(f"转换完成: {src} -> {dst} ({len(df)}行)")
    elif cmd == 'report':
        if len(sys.argv) < 4:
            print("用法: report <src> <dst> [sheet]")
            return 1
        dst = sys.argv[3]
        sheet = sys.argv[4] if len(sys.argv) > 4 else None
        df = read_data(src) if not src.endswith('.xlsx') or sheet is None else pd.read_excel(src, sheet_name=sheet, engine='openpyxl')
        write_report(df, dst)
        print(f"报表生成: {dst} ({len(df)}行 x {len(df.columns)}列)")
    elif cmd == 'pdf2xlsx':
        if len(sys.argv) < 4:
            print("用法: pdf2xlsx <pdf> <dst> [page]")
            return 1
        dst = sys.argv[3]
        page = int(sys.argv[4]) if len(sys.argv) > 4 else 0
        pdf_to_excel(src, dst, page)
        print(f"PDF转Excel完成: {dst}")
    elif cmd == 'merge':
        if len(sys.argv) < 4:
            print("用法: merge <file> <dst>")
            return 1
        merge_sheets(src, sys.argv[3])
        print(f"Sheet合并完成: {sys.argv[3]}")
    else:
        print(f"未知命令: {cmd}")
        print(__doc__)
        return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())
