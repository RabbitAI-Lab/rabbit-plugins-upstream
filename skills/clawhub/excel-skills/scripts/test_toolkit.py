# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding='utf-8')
import os, tempfile
sys.path.insert(0, r'C:\Users\Administrator\.openclaw\workspace\linlan-skills\excel-skills\scripts')
from excel_toolkit import read_data, write_report, merge_sheets
import pandas as pd
import openpyxl

tmp = tempfile.mkdtemp()
src = os.path.join(tmp, 'sample.xlsx')

# 构造多Sheet测试数据
with pd.ExcelWriter(src, engine='openpyxl') as w:
    pd.DataFrame({'姓名': ['张三', '李四'], '销售额': [1200.5, 3400.75]}).to_excel(w, sheet_name='一月', index=False)
    pd.DataFrame({'姓名': ['王五'], '销售额': [900]}).to_excel(w, sheet_name='二月', index=False)

# 1. report: 规范报表
df = pd.read_excel(src, sheet_name='一月')
r1 = write_report(df, os.path.join(tmp, 'report.xlsx'))
wb = openpyxl.load_workbook(r1)
ws = wb.active
c = ws['A1']
print('report: 冻结 =', ws.freeze_panes, '| 表头填充 =', c.fill.start_color.rgb, '| bold =', c.font.bold)
print('report: 列宽A =', round(ws.column_dimensions['A'].width))

# 2. merge: 多Sheet合并
r2 = merge_sheets(src, os.path.join(tmp, 'merged.xlsx'))
md = pd.read_excel(r2)
print('merge: 行数 =', len(md), '| 列 =', list(md.columns))
print('merged中文列宽 =', round(openpyxl.load_workbook(r2).active.column_dimensions['B'].width))
print('merge数据:')
print(md.to_string(index=False))

# 3. convert: xlsx -> csv (utf-8-sig)
csv_dst = os.path.join(tmp, 'out.csv')
md.to_csv(csv_dst, index=False, encoding='utf-8-sig')
back = pd.read_csv(csv_dst)
assert len(back) == 3, 'csv行数不符'
print('convert: csv往返行数 =', len(back))

print('ALL PASS')
