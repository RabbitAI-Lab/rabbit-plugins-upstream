"""Batching must reduce opens without relaxing source or context verification."""
import copy
import hashlib
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import openpyxl

from claim_source_guard import validate_sources, source_value
from source_read_batch import SourceReadBatch
from finish_agent_handoff import verify_packet


class BatchTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory(); self.addCleanup(temp.cleanup)
        self.root = Path(temp.name); self.path = self.root/'values.xlsx'
        book = openpyxl.Workbook(); sheet = book.active; sheet.title='data'
        sheet.append(['金额','期间','平台','店铺','单位']); sheet.append([100,'2026-W36','tmall','测试店','元'])
        book.save(self.path); book.close()
        self.ref = dict(source_file=self.path.name,source_sha256=hashlib.sha256(self.path.read_bytes()).hexdigest(),sheet='data',range='A2',field='金额',contexts={
            'period':{'column':'B','field':'期间'},'platform':{'column':'C','field':'平台'},
            'store':{'column':'D','field':'店铺'},'unit':{'column':'E','field':'单位'}})
        self.payload = dict(source_context_policy='required',scope=dict(platform='tmall',store='测试店'),claims=[
            dict(claim_id='C'+str(n),metric='gmv',value=100,unit='元',period='2026-W36',status='verified',source_ref=copy.deepcopy(self.ref)) for n in range(8)])

    def test_one_excel_open_including_contexts(self):
        original = openpyxl.load_workbook
        with patch('openpyxl.load_workbook', wraps=original) as opened:
            checks=validate_sources(self.payload,self.root)
        self.assertEqual(opened.call_count,1)
        self.assertEqual(len(checks),8)
        self.assertTrue(all(c['value']=='100' and c['source_context']['status']=='source_context_verified' for c in checks))

    def test_wrong_value_remains_blocked(self):
        self.payload['claims'][2]['value']=101
        with self.assertRaisesRegex(ValueError,'claim_value_not_reproducible'): validate_sources(self.payload,self.root)

    def test_context_mismatch_remains_blocked(self):
        self.payload['scope']['store']='另一店'
        with self.assertRaisesRegex(ValueError,'source_context_value_mismatch'): validate_sources(self.payload,self.root)

    def test_different_hash_for_same_file_is_not_reused(self):
        self.payload['claims'][3]['source_ref']['source_sha256']='0'*64
        with self.assertRaisesRegex(ValueError,'claim_source_hash_mismatch'): validate_sources(self.payload,self.root)

    def test_mutation_after_snapshot_is_blocked(self):
        batch=SourceReadBatch(self.payload['claims'],self.root)
        self.assertEqual(source_value(self.ref,self.root,batch)[0],100)
        self.path.write_bytes(b'changed')
        with self.assertRaisesRegex(ValueError,'changed_during_read'): batch.verify_unchanged()

    def test_small_memory_budget_falls_back_to_streaming(self):
        batch=SourceReadBatch(self.payload['claims'],self.root,cell_limit=1)
        self.assertEqual(source_value(self.ref,self.root,batch)[0],100)
        batch.verify_unchanged()

    def test_new_validation_does_not_reuse_old_snapshot(self):
        validate_sources(self.payload,self.root)
        self.path.write_bytes(b'changed')
        with self.assertRaisesRegex(ValueError,'claim_source_hash_mismatch'): validate_sources(self.payload,self.root)

    def test_memory_budget_applies_across_multiple_files(self):
        second = self.root/'second.xlsx'; second.write_bytes(self.path.read_bytes())
        ref = copy.deepcopy(self.ref); ref['source_file']=second.name
        claims = self.payload['claims']+[dict(self.payload['claims'][0],source_ref=ref)]
        batch = SourceReadBatch(claims,self.root,cell_limit=15)
        self.assertIsNotNone(batch.rows(self.path,self.ref))
        self.assertIsNone(batch.rows(second,ref))
        self.assertEqual(batch.cached_cells,10)
        self.assertEqual(source_value(ref,self.root,batch)[0],100)
        batch.verify_unchanged()

    def test_unknown_cost_stays_unknown_without_read(self):
        row=self.payload['claims'][0]; row.update(status='unknown',value=None)
        self.payload['claims']=[row]; self.path.write_bytes(b'not Excel')
        self.assertEqual(validate_sources(self.payload,self.root)[0]['status'],'unknown_not_recomputed')

    def test_packet_content_cannot_diverge_from_file(self):
        file=self.root/'baseline.json'; file.write_text('{"value":1}',encoding='utf-8')
        card={'read_packet':{'baseline':{'path':str(file),'sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'content':{'value':1}}}}
        verify_packet(card); card['read_packet']['baseline']['content']['value']=2
        with self.assertRaisesRegex(ValueError,'content_changed'): verify_packet(card)


if __name__=='__main__': unittest.main()
