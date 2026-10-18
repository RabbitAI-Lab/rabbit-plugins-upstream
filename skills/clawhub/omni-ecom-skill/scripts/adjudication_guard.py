"""Verify conflict dispositions by reading source cells and recomputing arithmetic.

This verifies reproducibility against supplied sources, not source authenticity
or causal validity. Those remain separate acquisition and review checks.
"""
import ast
import csv
import hashlib
import json
from decimal import Decimal, InvalidOperation, localcontext
from pathlib import Path

from artifact_contract import validate_analysis_identity


def require(ok, reason):
    if not ok:
        raise ValueError(reason)


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def fingerprint(conflicts):
    return hashlib.sha256(json.dumps(conflicts,ensure_ascii=False,sort_keys=True,
                                    separators=(',',':'),allow_nan=False).encode()).hexdigest()


def number(value):
    require(not isinstance(value,bool), 'recompute_number_invalid')
    try:
        result=Decimal(str(value))
    except (InvalidOperation,ValueError):
        raise ValueError('recompute_number_invalid')
    require(result.is_finite(), 'recompute_number_invalid')
    return result


def calculate(expression, inputs):
    require(isinstance(expression,str) and len(expression)<=2000, 'recompute_formula_invalid')
    try:
        tree=ast.parse(expression,mode='eval')
    except (SyntaxError,ValueError,RecursionError):
        raise ValueError('recompute_formula_invalid')
    require(sum(1 for _ in ast.walk(tree))<=100, 'recompute_formula_too_complex')
    def visit(node):
        if isinstance(node,ast.Name):
            require(node.id in inputs, 'recompute_input_missing')
            return inputs[node.id]
        if isinstance(node,ast.Constant) and type(node.value) in (int,float):
            return number(node.value)
        if isinstance(node,ast.UnaryOp) and isinstance(node.op,(ast.USub,ast.UAdd)):
            return -visit(node.operand) if isinstance(node.op,ast.USub) else visit(node.operand)
        if isinstance(node,ast.BinOp) and isinstance(node.op,(ast.Add,ast.Sub,ast.Mult,ast.Div)):
            a,b=visit(node.left),visit(node.right)
            if isinstance(node.op,ast.Add): return a+b
            if isinstance(node.op,ast.Sub): return a-b
            if isinstance(node.op,ast.Mult): return a*b
            require(b!=0, 'recompute_zero_denominator')
            return a/b
        raise ValueError('recompute_operator_not_allowed')
    require(any(isinstance(n,ast.Name) for n in ast.walk(tree)), 'recompute_source_required')
    with localcontext() as ctx:
        ctx.prec=28
        return visit(tree.body)


def cell_value(spec, resolve, source_names):
    path=resolve(spec.get('file'))
    require(path.name in source_names, 'recompute_source_not_in_inventory')
    original=sha(path)
    require(spec.get('sha256')==original, 'recompute_source_changed')
    if path.suffix.lower() in ('.csv','.xlsx','.xls'):
        selected=spec.get('rows')
        require(isinstance(selected,list) and bool(selected) and
                all(type(i) is int and i>=0 for i in selected) and len(selected)==len(set(selected)), 'recompute_rows_invalid')
        positions=set(selected)
    if path.suffix.lower()=='.csv':
        with path.open(encoding=spec.get('encoding','utf-8-sig'),newline='') as f:
            reader=csv.DictReader(f)
            require(spec.get('column') in (reader.fieldnames or []), 'recompute_column_missing')
            require(len(reader.fieldnames)==len(set(reader.fieldnames)), 'recompute_duplicate_header')
            values=[number(row[spec['column']]) for i,row in enumerate(reader) if i in positions]
            require(len(values)==len(positions), 'recompute_rows_missing')
    elif path.suffix.lower() in ('.xlsx','.xls'):
        header=spec.get('header_row')
        require(type(header) is int and header>=1 and isinstance(spec.get('sheet'),str), 'recompute_sheet_selection_missing')
        book=None
        try:
            if path.suffix.lower()=='.xlsx':
                try:
                    import openpyxl
                except ImportError:
                    raise ValueError('recompute_dependency_missing:openpyxl')
                book=openpyxl.load_workbook(path,read_only=True,data_only=True)
                require(spec['sheet'] in book.sheetnames, 'recompute_sheet_missing')
                iterator=book[spec['sheet']].iter_rows(min_row=header,values_only=True)
            else:
                try:
                    import xlrd
                except ImportError:
                    raise ValueError('recompute_dependency_missing:xlrd')
                book=xlrd.open_workbook(str(path),on_demand=True)
                require(spec['sheet'] in book.sheet_names(), 'recompute_sheet_missing')
                sheet=book.sheet_by_name(spec['sheet'])
                iterator=(sheet.row_values(i) for i in range(header-1,sheet.nrows))
            headings=next(iterator,())
            require(spec.get('column') in headings, 'recompute_column_missing')
            require(headings.count(spec['column'])==1, 'recompute_duplicate_header')
            column=headings.index(spec['column'])
            values=[number(row[column]) for i,row in enumerate(iterator) if i in positions]
            require(len(values)==len(positions), 'recompute_rows_missing')
        finally:
            if book is not None:
                if hasattr(book,'close'): book.close()
                elif hasattr(book,'release_resources'): book.release_resources()
    elif path.suffix.lower()=='.json':
        value=json.loads(path.read_text(encoding='utf-8-sig'))
        pointer=spec.get('pointer')
        require(isinstance(pointer,str) and pointer.startswith('/'), 'recompute_pointer_invalid')
        try:
            for part in pointer[1:].split('/'):
                part=part.replace('~1','/').replace('~0','~')
                value=value[int(part)] if isinstance(value,list) else value[part]
        except (KeyError,IndexError,ValueError,TypeError):
            raise ValueError('recompute_cell_missing')
        values=[number(value)]
    else:
        raise ValueError('recompute_source_format_unsupported')
    op=spec.get('aggregation','value')
    require(op in ('value','sum','mean'), 'recompute_aggregation_invalid')
    require(op!='value' or len(values)==1, 'recompute_value_not_single')
    total=sum(values,Decimal(0))
    result=total/len(values) if op=='mean' else total
    require(sha(path)==original, 'recompute_source_changed')
    return result,path,original


def validate_downstream(row,report,values,actual):
    require(isinstance(row.get('downstream'),list), 'adjudication_downstream_review_missing')
    if not row['downstream']:
        require(bool(row.get('no_downstream_reason')), 'adjudication_downstream_review_missing')
    ids=[]
    for dependent in row['downstream']:
        require(isinstance(dependent,dict) and isinstance(dependent.get('claim_id'),str) and
                bool(dependent['claim_id']) and dependent['claim_id'] not in ids,'adjudication_downstream_id_invalid')
        ids.append(dependent['claim_id'])
        result=calculate(dependent.get('formula'),dict(values,adjudicated=actual))
        require(result==number(dependent.get('result')), 'adjudication_downstream_mismatch')
        dc=report.get('adjudicated_downstream',{}).get(dependent['claim_id'],{})
        require(number(dc.get('value'))==result, 'adjudication_downstream_report_stale')


def validate_adjudication(found, verdict, report, resolve, source_names,baseline=None,inventory=None,baseline_sha256=None):
    validate_analysis_identity(verdict,'conflict_verdict',report.get('team_version'))
    for key in ('run_id','client_scope','team_version','report_revision'):
        require(verdict.get(key)==report.get(key), 'adjudication_identity_mismatch')
    require(verdict.get('conflicts_sha256')==fingerprint(found), 'adjudication_conflict_changed')
    hard={c['conflict_id']:c for c in found if c['severity']=='hard'}
    rows=verdict.get('verdicts')
    require(isinstance(rows,list) and all(isinstance(r,dict) for r in rows) and len(rows)==len(hard), 'adjudication_incomplete')
    require({r.get('conflict_id') for r in rows}==set(hard), 'adjudication_incomplete')
    results,sources=[],{}
    claims=report.get('adjudicated_claims',{})
    for row in rows:
        cid=row['conflict_id']
        require(row.get('decision') and row.get('not_adopted_disposition') and row.get('residual_risk'), 'adjudication_explanation_missing')
        method=row.get('method')
        if method in ('scope_degradation','pending_verification'):
            require(report.get('task_type')=='data_quality_audit', 'unresolved_conflict_requires_data_quality_report')
            require(any(x.get('conflict_id')==cid and x.get('reason') and x.get('remedy')
                        for x in report.get('unresolved_conflicts',[])), 'unresolved_conflict_not_disclosed')
            require(not report.get('actions'), 'unresolved_conflict_action_blocked')
            results.append(dict(conflict_id=cid,status='disclosed_unresolved',method=method))
            continue
        if method in ('caliber_recomputation','trend_recomputation'):
            from semantic_adjudication import validate_semantic
            result,used,values,actual=validate_semantic(hard[cid],row,report,resolve,baseline,inventory,baseline_sha256)
            validate_downstream(row,report,values,actual)
            results.append(result); sources.update(used)
            continue
        require(method in ('lead_recomputation','authority_source'), 'adjudication_method_invalid')
        require(hard[cid]['type']=='numeric', 'causal_conflict_requires_scope_review')
        specs=row.get('inputs')
        require(isinstance(specs,dict) and bool(specs), 'recompute_inputs_missing')
        values={}
        for name,spec in specs.items():
            require(isinstance(name,str) and name.isidentifier(), 'recompute_variable_invalid')
            value,path,original=cell_value(spec,resolve,source_names)
            values[name]=value
            sources[str(path)]=original
        actual=calculate(row.get('formula'),values)
        require(actual==number(row.get('result')), 'recompute_result_mismatch')
        claim=claims.get(cid,{})
        require(claim.get('metric') and claim.get('period') and claim.get('unit') and
                number(claim.get('value'))==actual, 'adjudication_report_not_updated')
        require(claim['metric']==hard[cid]['metric'] and claim['period']==hard[cid]['period'],
                'adjudication_metric_scope_mismatch')
        validate_downstream(row,report,values,actual)
        results.append(dict(conflict_id=cid,status='recomputed',method=method,result=str(actual)))
    return results,sources
