"""Reproduce ledger values from local sources, not authenticate their origin."""
import ast
import csv
import json
import re
from decimal import Decimal, ROUND_HALF_UP, localcontext
from pathlib import Path

from adjudication_guard import calculate, number, require, sha, fingerprint


def source_value(source, root, batch=None):
    name=source.get('source_file')
    require(isinstance(name,str) and bool(name), 'claim_source_file_missing')
    relative=Path(name)
    require(not relative.is_absolute() and '..' not in relative.parts, 'claim_source_path_outside')
    path=(root/relative).resolve()
    require(path.is_relative_to(root) and path.is_file(), 'claim_source_outside_or_missing')
    digest=batch.digest(path) if batch is not None else sha(path)
    require(digest==str(source.get('source_sha256','')).lower(), 'claim_source_hash_mismatch')
    suffix=path.suffix.lower()
    if suffix=='.json':
        pointer=source.get('pointer')
        require(isinstance(pointer,str) and pointer.startswith('/'), 'claim_source_pointer_required')
        value=batch.json_value(path) if batch is not None else json.loads(path.read_text(encoding='utf-8-sig'))
        for part in pointer[1:].split('/'):
            key=part.replace('~1','/').replace('~0','~')
            if isinstance(value,list):
                require(bool(re.fullmatch(r'0|[1-9][0-9]*',key)), 'claim_source_pointer_invalid')
                value=value[int(key)]
            else:
                value=value[key]
        values=[number(value)]
    else:
        require(suffix in ('.csv','.xlsx','.xls'), 'claim_source_format_unsupported')
        selection=re.fullmatch(r'([A-Z]+)([1-9][0-9]*)(?::([A-Z]+)([1-9][0-9]*))?', str(source.get('range','')))
        require(selection is not None, 'claim_source_range_required')
        col,start,endcol,end=selection.groups()
        require(endcol is None or endcol==col, 'claim_source_single_column_required')
        start,end=int(start),int(end or start)
        header=source.get('header_row',1)
        require(type(header) is int and 1<=header<start<=end<=1048576, 'claim_source_rows_invalid')
        column=0
        for letter in col:
            column=column*26+ord(letter)-64
        require(column<=16384, 'claim_source_column_invalid')
        column-=1
        book=None
        try:
            snapshot=batch.rows(path,source) if batch is not None else None
            if snapshot is not None:
                rows=iter(snapshot)
            elif suffix=='.csv':
                book=path.open(encoding=source.get('encoding','utf-8-sig'),newline='')
                rows=csv.reader(book)
            elif suffix=='.xlsx':
                import openpyxl
                book=openpyxl.load_workbook(path,read_only=True,data_only=True)
                require(source.get('sheet') in book.sheetnames, 'claim_source_sheet_missing')
                rows=book[source['sheet']].iter_rows(values_only=True)
            else:
                import xlrd
                book=xlrd.open_workbook(str(path),on_demand=True)
                require(source.get('sheet') in book.sheet_names(), 'claim_source_sheet_missing')
                sheet=book.sheet_by_name(source['sheet'])
                rows=(sheet.row_values(i) for i in range(sheet.nrows))
            values=[]
            for index,row in enumerate(rows,1):
                if index==header:
                    require(column<len(row) and row[column]==source.get('field'), 'claim_source_field_mismatch')
                    require(row.count(source['field'])==1, 'claim_source_field_ambiguous')
                if start<=index<=end:
                    require(column<len(row), 'claim_source_cell_missing')
                    values.append(number(row[column]))
                if index>=end:
                    break
            require(len(values)==end-start+1, 'claim_source_rows_missing')
        finally:
            if book is not None:
                if hasattr(book,'close'): book.close()
                elif hasattr(book,'release_resources'): book.release_resources()
    op=source.get('aggregation','value')
    require(op in ('value','sum','mean'), 'claim_source_aggregation_invalid')
    require(op!='value' or len(values)==1, 'claim_source_aggregation_required')
    with localcontext() as ctx:
        ctx.prec=28
        result=sum(values,Decimal(0))
        if op=='mean': result/=len(values)
    if batch is None:
        require(sha(path)==digest, 'claim_source_changed_during_read')
    return result,dict(file=name,sha256=digest)


def validate_sources(payload, root):
    """Call only after structural validation. Unknowns are never read as zeros."""
    root=Path(root).resolve()
    require(payload.get('source_context_policy','legacy_unbound') in ('legacy_unbound','required'),
            'source_context_policy_invalid')
    claims={row['claim_id']:row for row in payload['claims']}
    from source_read_batch import SourceReadBatch
    batch=SourceReadBatch(payload['claims'],root)
    done,visiting,checks={},set(),[]

    def evaluate(key):
        if key in done: return done[key]
        require(key not in visiting, 'claim_dependency_cycle')
        visiting.add(key)
        row=claims[key]
        status=row['status']
        if status=='unknown':
            result=None
            check=dict(claim_id=key,status='unknown_not_recomputed')
        elif status=='verified':
            result,source=source_value(row['source_ref'],root,batch)
            check=dict(claim_id=key,status='source_recomputed',source=source)
            from source_context_guard import validate_context
            check['source_context']=validate_context(payload,row,root,batch)
        else:
            formula=row['formula']
            inputs=row['inputs']
            require(len(inputs)==len(set(inputs)), 'claim_inputs_duplicate')
            bindings=formula.get('bindings')
            if bindings is None:
                metrics=[claims[item]['metric'] for item in inputs]
                require(len(metrics)==len(set(metrics)), 'claim_input_metric_ambiguous')
                bindings=dict(zip(metrics,inputs))
            require(isinstance(bindings,dict) and set(bindings.values())==set(inputs)
                    and len(bindings)==len(inputs), 'claim_formula_bindings_invalid')
            values={name:evaluate(item) for name,item in bindings.items()}
            require(all(value is not None for value in values.values()), 'claim_unknown_input')
            expression=formula['expression']
            result=calculate(expression,values)
            tree=ast.parse(expression,mode='eval')
            names={node.id for node in ast.walk(tree) if isinstance(node,ast.Name)}
            require(names==set(values), 'claim_unused_formula_input')
            # Ratio declarations must actually divide the declared numerator by
            # denominator; a constant plus zero times genuine inputs is not ROI.
            numerator,denominator=formula['numerator'],formula['denominator']
            ratio_names={'roas','conversion_rate','conversion','转化率','uv_value','gmv_per_visitor','访客价值','aov','refund_rate','roi','net_roas','contribution_roas'}
            from metric_semantics import validate_expression
            normalized = validate_expression(row,claims,bindings,values,tree,result)
            if not normalized and str(formula['name']).lower() in ratio_names:
                body=tree.body
                require(isinstance(body,ast.BinOp) and isinstance(body.op,ast.Div)
                        and isinstance(body.left,ast.Name) and isinstance(body.right,ast.Name)
                        and body.left.id==numerator and body.right.id==denominator,
                        'claim_ratio_expression_mismatch')
                require(claims[bindings[numerator]]['metric']==numerator
                        and claims[bindings[denominator]]['metric']==denominator,
                        'claim_ratio_binding_mismatch')
            check=dict(claim_id=key,status='formula_recomputed',inputs=inputs)
            from net_return_guard import validate as validate_net_return
            net_return = validate_net_return(payload,row,claims,bindings,values,tree,result)
            if net_return is not None:
                check['net_return'] = net_return
            from accounting_guard import validate_accounting
            accounting = validate_accounting(payload,row,claims,evaluate,root)
            if accounting is not None:
                check['accounting'] = accounting
        if result is not None:
            places=row.get('decimal_places')
            if places is not None:
                require(type(places) is int and 0<=places<=12, 'claim_rounding_invalid')
                with localcontext() as ctx:
                    ctx.prec=max(28,len(result.as_tuple().digits)+abs(result.adjusted())+places+2)
                    result=result.quantize(Decimal(1).scaleb(-places),rounding=ROUND_HALF_UP)
            require(result==number(row['value']), 'claim_value_not_reproducible:'+key)
            check['value']=str(result)
        done[key]=result
        visiting.remove(key)
        checks.append(check)
        return result

    for key in claims:
        evaluate(key)
    batch.verify_unchanged()
    return checks


def validate_report_sources(report, report_dir, report_text, *, require_context=False):
    """Bind the displayed ledger to exact files, returning all freeze inputs."""
    from claim_guard import read_json, validate_claims
    root=Path(report_dir).resolve()
    ledger=report.get('claim_ledger')
    require(isinstance(ledger,dict), 'claim_report_ledger_missing')
    name=ledger.get('file')
    require(isinstance(name,str) and bool(name) and not Path(name).is_absolute()
            and '..' not in Path(name).parts, 'claim_report_ledger_path_invalid')
    path=(root/name).resolve()
    require(path.is_relative_to(root) and path.is_file(), 'claim_report_ledger_outside_or_missing')
    original=sha(path)
    require(ledger.get('sha256')==original, 'claim_report_ledger_hash_mismatch')
    payload=read_json(path)
    for field in ('run_id','client_scope','period'):
        require(bool(report.get(field)) and payload.get(field)==report[field], 'claim_report_identity_mismatch:'+field)
    errors,counts=validate_claims(payload,report_text)
    require(not errors, 'claim_report_structure_blocked:'+(';'.join(errors)))
    require(fingerprint(ledger.get('claims'))==fingerprint(payload['claims']), 'claim_report_ledger_projection_mismatch')
    require(all(type(ledger.get(key)) is int and ledger[key]==value for key,value in counts.items()),
            'claim_report_counts_mismatch')
    # Formal completion chooses this policy in code, never from a report's
    # version/mode flag. Legacy ledgers remain readable but cannot downgrade it.
    effective=dict(payload,source_context_policy='required') if require_context else payload
    checks=validate_sources(effective,path.parent)
    if require_context:
        from source_context_guard import summary
        require(summary(checks)['status']=='all_verified_source_labels_checked',
                'formal_source_context_evidence_missing')
    require(ledger.get('source_checks')==checks, 'claim_report_source_checks_mismatch')
    from source_context_guard import expected_context
    by_id={row['claim_id']:row for row in payload['claims']}
    for check in checks:
        if check.get('source_context',{}).get('status')=='source_context_verified':
            row=by_id[check['claim_id']]
            require(expected_context(payload,row)==expected_context(report,row),
                    'claim_report_source_context_mismatch')
    files={str(path.relative_to(root)):original}
    from previous_action_review import validate_feedback
    for item in validate_feedback(report,root):
        files[item['file']]=item['sha256']
    for check in checks:
        if 'source' in check:
            source=check['source']
            source_path=(path.parent/source['file']).resolve()
            require(source_path.is_relative_to(root), 'claim_report_source_outside')
            relative=str(source_path.relative_to(root))
            require(relative not in files or files[relative]==source['sha256'], 'claim_report_source_snapshot_conflict')
            files[relative]=source['sha256']
    for name,digest in files.items():
        require(sha(root/name)==digest, 'claim_report_source_changed_during_check')
    return dict(status='claim_sources_recomputed',run_id=payload['run_id'],client_scope=payload['client_scope'],
                period=payload['period'],claims_file=str(path.relative_to(root)),claims_sha256=original,
                **counts,source_checks=checks,
                inputs=[dict(file=name,sha256=digest) for name,digest in sorted(files.items())])


def validate_pdf_source_binding(receipt, report_dir, report):
    """A rendered PDF must bind the same source snapshot as its current report."""
    root=Path(report_dir).resolve()
    markdown=root/'report.md'
    expected={'report.json':sha(root/'report.json'),'report.md':sha(markdown)}
    actual=validate_report_sources(report,root,markdown.read_text(encoding='utf-8-sig'))
    expected.update({item['file']:item['sha256'] for item in actual['inputs']})
    require(receipt.get('input_fingerprints')==expected, 'pdf_source_fingerprints_missing_or_stale')
    require(receipt.get('claim_source_gate')==actual, 'pdf_claim_source_gate_stale')
    for field in ('run_id','client_scope','team_version','report_revision'):
        require(receipt.get(field)==report.get(field), 'pdf_source_identity_mismatch')
    for name,digest in expected.items():
        require(sha(root/name)==digest, 'pdf_source_changed_during_check')
    return actual
