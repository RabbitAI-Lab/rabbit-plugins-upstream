"""Short-lived source snapshots for ONE verification call, never a trusted receipt."""
import csv
import json
import re
from itertools import islice
from pathlib import Path

from adjudication_guard import sha, require


class SourceReadBatch:
    def __init__(self, claims, root, cell_limit=2_000_000):
        self.root = Path(root).resolve()
        self.cell_limit = cell_limit
        self.limits = {}
        self.digests = {}
        self.snapshots = {}
        self.cached_cells = 0
        for claim in claims:
            if claim.get('status') != 'verified':
                continue
            ref = claim.get('source_ref', {})
            match = re.fullmatch(r'([A-Z]+)([1-9][0-9]*)(?::([A-Z]+)([1-9][0-9]*))?', str(ref.get('range', '')))
            if match:
                key = self.key(ref)
                self.limits[key] = max(self.limits.get(key, 0), int(match[4] or match[2]))

    @staticmethod
    def key(ref):
        return (str(ref.get('source_file')), str(ref.get('sheet', '')), str(ref.get('encoding', 'utf-8-sig')))

    def digest(self, path):
        path = path.resolve()
        if path not in self.digests:
            self.digests[path] = sha(path)
        return self.digests[path]

    def rows(self, path, ref):
        key = self.key(ref)
        if key in self.snapshots:
            return self.snapshots[key]
        end = self.limits.get(key)
        if end is None:
            return None
        book = None
        digest = self.digest(path)
        try:
            if path.suffix.lower() == '.csv':
                book = path.open(encoding=ref.get('encoding', 'utf-8-sig'), newline='')
                rows = csv.reader(book)
            elif path.suffix.lower() == '.xlsx':
                import openpyxl
                book = openpyxl.load_workbook(path, read_only=True, data_only=True)
                require(ref.get('sheet') in book.sheetnames, 'claim_source_sheet_missing')
                rows = book[ref['sheet']].iter_rows(values_only=True)
            elif path.suffix.lower() == '.xls':
                import xlrd
                book = xlrd.open_workbook(str(path), on_demand=True)
                require(ref.get('sheet') in book.sheet_names(), 'claim_source_sheet_missing')
                sheet = book.sheet_by_name(ref['sheet'])
                rows = (sheet.row_values(i) for i in range(sheet.nrows))
            else:
                return None
            snapshot, cells = [], 0
            for row in islice(rows, end):
                cells += len(row)
                if self.cached_cells + cells > self.cell_limit:
                    # Large tables retain the bounded-memory streaming verifier.
                    self.snapshots[key] = None
                    return None
                snapshot.append(tuple(row))
            require(sha(path) == digest, 'claim_source_changed_during_read')
            self.snapshots[key] = tuple(snapshot)
            self.cached_cells += cells
            return self.snapshots[key]
        finally:
            if book is not None:
                if hasattr(book, 'close'): book.close()
                elif hasattr(book, 'release_resources'): book.release_resources()

    def json_value(self, path):
        key = ('json', str(path))
        if key not in self.snapshots:
            digest = self.digest(path)
            self.snapshots[key] = json.loads(path.read_text(encoding='utf-8-sig'))
            require(sha(path) == digest, 'claim_source_changed_during_read')
        return self.snapshots[key]

    def verify_unchanged(self):
        for path, digest in self.digests.items():
            require(sha(path) == digest, 'claim_source_changed_during_read')
