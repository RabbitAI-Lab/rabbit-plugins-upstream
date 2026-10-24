#!/usr/bin/env python3
"""
Extract core instances from Excel (关键资源标记表).

Single-customer mode (CES): no customer grouping, output flat instances dict.
Output format: {source, instances: {svc: [inst, ...]}}
"""

import json
import logging
import sys
from pathlib import Path
from collections import defaultdict

from resource_type import SERVICE_METRIC_DIMENSION_MAP
from workspace import INTERMEDIATE_DIR

try:
    import openpyxl
except ImportError:
    print("Error: openpyxl is required. Install with: pip install openpyxl")

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("core-instance")

DEFAULT_INPUT = str(Path(__file__).parent.parent / "assets" / "关键资源标记表.xlsx")
DEFAULT_OUTPUT = str(Path(INTERMEDIATE_DIR) / "core_instances.json")

SHEET_NAME = "关键资源标记"

SERVICE_NAME_MAPPING = {
    "ECS-tidb": "ECS", "ECS-openresty": "ECS", "RDS-PG": "RDS",
    "公网NAT": "NAT", "DC-物理连接": "DC", "CC-域间带宽": "CC",
    "云原生防护": "DDOS", "WAF": "WAF", "ELB": "ELB", "CCE": "CCE",
    "DCS": "DCS", "DMS": "DMS", "GeminiDB": "GeminiDB", "DDOS": "DDOS",
    "DDS": "DDS", "GEIP": "GEIP", "EIP": "EIP",
}


def normalize_service_name(raw_name: str) -> str:
    if raw_name in SERVICE_NAME_MAPPING:
        return SERVICE_NAME_MAPPING[raw_name]
    prefix = raw_name.split("-")[0].strip()
    if prefix in SERVICE_NAME_MAPPING:
        return SERVICE_NAME_MAPPING[prefix]
    return raw_name


def _parse_sheet(ws) -> dict:
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))

    col_region_idx = None
    col_svc_idx = None
    col_inst_idx = None
    col_core_idx = None
    for i, val in enumerate(header_row):
        if val and str(val).strip().lower() == "region":
            col_region_idx = i
        if val and str(val).strip() == "resource_type":
            col_svc_idx = i
        if val and str(val).strip() == "resource_id":
            col_inst_idx = i
        if val and str(val).strip() == "是否核心业务":
            col_core_idx = i

    required_cols = {
        "region": col_region_idx,
        "resource_type": col_svc_idx,
        "resource_id": col_inst_idx,
        "是否核心业务": col_core_idx,
    }
    missing = [name for name, idx in required_cols.items() if idx is None]
    if missing:
        raise ValueError(f"Sheet '{ws.title}' 缺少必要列: {missing}")

    # Validate sheet has data rows beyond header
    if ws.max_row < 2:
        raise ValueError(
            f"Sheet '{ws.title}' 没有数据行（仅有表头，无实际数据）。"
            f"请在关键资源标记表.xlsx 中填写核心实例数据，"
            f"至少包含 region、resource_type、resource_id、是否核心业务 四列。"
        )

    max_col = max(col_svc_idx, col_inst_idx,
                  col_region_idx if col_region_idx is not None else 0,
                  col_core_idx if col_core_idx is not None else 0) + 1

    instances = defaultdict(list)
    skipped = 0
    for row in ws.iter_rows(min_row=2, max_col=max_col, values_only=True):
        if col_core_idx is not None and row[col_core_idx]:
            if str(row[col_core_idx]).strip() != "是":
                skipped += 1
                continue

        raw_service = str(row[col_svc_idx]).strip().split("\n")[0].strip() if row[col_svc_idx] else ""
        instance_id = str(row[col_inst_idx]).strip() if row[col_inst_idx] else ""
        region = str(row[col_region_idx]).strip() if col_region_idx is not None and row[col_region_idx] else ""
        if not raw_service or not instance_id:
            continue
        service = normalize_service_name(raw_service)
        if service == "VPC":
            continue
        inst = {
            "instance_id": instance_id,
            "raw_service": raw_service,
            "region": region,
        }
        dim = SERVICE_METRIC_DIMENSION_MAP.get(service)
        if dim:
            inst["metric_dimension"] = dim
        instances[service].append(inst)

    if skipped:
        logger.info(f"  Skipped {skipped} rows (是否核心业务 != 是)")

    if not instances:
        raise ValueError(
            f"Sheet '{ws.title}' 没有提取到任何核心实例。\n"
            f"请检查以下条件是否满足：\n"
            f"  1. 「是否核心业务」列至少有一行值为「是」\n"
            f"  2. 对应行的 resource_type 和 resource_id 列不为空\n"
            f"请在关键资源标记表.xlsx 中增加核心实例清单后重新运行。"
        )

    return dict(instances)


def get_core_instances(input_path: str = DEFAULT_INPUT, output_path: str = DEFAULT_OUTPUT) -> dict:
    input_p = Path(input_path)
    if not input_p.exists():
        raise FileNotFoundError(f"关键资源标记表不存在: {input_p}")

    logger.info(f"Reading: {input_p}")
    wb = openpyxl.load_workbook(str(input_p), data_only=True)

    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Sheet '{SHEET_NAME}' 不存在，可用sheet: {wb.sheetnames}")

    ws = wb[SHEET_NAME]
    logger.info(f"Processing sheet: {SHEET_NAME}")

    instances = _parse_sheet(ws)

    result = {
        "source": input_p.name,
        "instances": instances,
    }

    output_p = Path(output_path)
    output_p.parent.mkdir(parents=True, exist_ok=True)
    with open(output_p, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    total_instances = sum(len(v) for v in instances.values())
    logger.info(f"Saved to {output_p}")
    logger.info(f"Total: {len(instances)} services, {total_instances} instances")
    for svc, insts in sorted(instances.items()):
        logger.info(f"  {svc}: {len(insts)} instances (raw: {set(i['raw_service'] for i in insts)})")

    # List all extracted instances
    logger.info(f"  Instance details ({total_instances} total):")
    for svc, insts in sorted(instances.items()):
        for inst in insts:
            dim = inst.get('metric_dimension', '')
            dim_str = f" | dim={dim}" if dim else ""
            logger.info(f"    {svc} | {inst['instance_id']} | region={inst['region']}{dim_str}")

    return result


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="Extract core instances from Excel (CES single-customer)")
    parser.add_argument("--input", default=DEFAULT_INPUT, help="Input Excel file path")
    parser.add_argument("--output", default=DEFAULT_OUTPUT, help="Output JSON file path")
    args = parser.parse_args()
    try:
        get_core_instances(input_path=args.input, output_path=args.output)
    except (ValueError, FileNotFoundError) as e:
        print(f"\n[ERROR] 核心实例提取失败，技能流程终止。\n{e}\n", file=sys.stderr)
        sys.exit(1)
