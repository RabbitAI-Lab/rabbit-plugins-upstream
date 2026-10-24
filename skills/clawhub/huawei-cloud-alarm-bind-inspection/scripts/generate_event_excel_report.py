#!/usr/bin/env python3
import json
import argparse
import logging
from pathlib import Path
from collections import defaultdict
from workspace import WORKSPACE_DIR, INTERMEDIATE_DIR

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
except ImportError:
    print("Error: openpyxl is required. Install with: pip install openpyxl")
    raise

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("event-excel-report")

DEFAULT_INPUT = str(Path(INTERMEDIATE_DIR) / "event_bind_report.json")
DEFAULT_OUTPUT = str(Path(WORKSPACE_DIR) / "event_bind_report.xlsx")
DEFAULT_NOTIFICATION_RULES = str(Path(INTERMEDIATE_DIR) / "online_notification_rules.json")
DEFAULT_ECS_SPECIAL_EVENTS = str(Path(INTERMEDIATE_DIR) / "ecs_special_events.json")
DEFAULT_AOM_EVENT_INPUT = str(Path(INTERMEDIATE_DIR) / "aom_event_bind_report.json")

ECS_SPECIAL_EVENT_IDS = {
    "RoCELinkStatusDown",
    "gpuDRAMUncorrectableEccError",
    "DAVPSMICardNotFound",
    "DAVPLspciCardNotFound",
    "TemperatureOverSdLimit",
    "DeviceCoreAbnormal",
    "gpuDriverLinkFailureAlarm",
    "gpuPcieLinkFailureAlarm",
    "vmLostGpuAlarm",
    "FPGALinkFault",
    "UncorrectableEccErrorCount",
}

HEADER_FONT = Font(bold=True, size=11)
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT_WHITE = Font(bold=True, size=11, color="FFFFFF")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

HEADERS = [
    "监控体系",
    "云服务-维度",
    "实例ID",
    "Region",
    "底线事件配置百分比",
    "覆盖的事件",
    "未覆盖的事件",
    "通知规则名称",
    "优化建议",
]


BASELINE_TO_RULE_SVC = {
    "RDS": "RDS", "GeminiDB": "NoSQL", "ECS": "ECS", "BMS": "BMS",
    "DDS": "DDS", "GAUSSDB": "GAUSSDB", "DCS": "DCS", "TaurusDB": "GAUSSDBV5",
    "DDM": "DMS", "OBS": "OBS", "EIP": "EIP", "GEIP": "EIP",
    "CBR": "CBR", "GA": "GA", "HSS": "HSS", "ELB": "ELB",
    "SecMaster": "SecMaster", "CSMS": "CSMS", "modelarts": "ModelArts",
    "SFSTurbo": "EFS",
}


def load_json(path):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def _load_ecs_special_event_names(ecs_special_path=None):
    if not ecs_special_path:
        return []
    try:
        data = load_json(ecs_special_path)
        return data.get("records", [])
    except Exception as e:
        logger.warning(f"Failed to load ECS special events: {e}")
        return []


def _build_event_to_rule_names(notification_rules_path):
    rules_data = load_json(notification_rules_path)
    mapping = defaultdict(set)
    for rule in rules_data.get("rules", []):
        rule_name = rule.get("name", "")
        if not rule_name:
            continue
        for ar in rule.get("alarmRules", []):
            svc = ar.get("service", "")
            event_name = ar.get("eventName", "")
            if svc and event_name:
                mapping[(svc, event_name)].add(rule_name)
    return {k: sorted(v) for k, v in mapping.items()}



AOM_EVENT_HEADERS = [
    "云服务类型", "底线事件数", "已覆盖数",
    "绑定率", "覆盖的事件", "未覆盖的事件",
    "核心实例数", "状态", "优化建议",
]


def _write_aom_event_sheet(ws, aom_report):
    """Write AOM event bind analysis as a separate sheet."""
    items = aom_report.get("items", [])
    summary = aom_report.get("summary", {})

    # Summary row
    ws.cell(row=1, column=1, value=f"AOM事件基线绑定率汇总："
            f"总底线事件 {summary.get('total_baseline_events', 0)}，"
            f"已覆盖 {summary.get('total_covered', 0)}，"
            f"覆盖率 {summary.get('overall_coverage', 'N/A')}").font = Font(bold=True, size=12)

    # Headers
    for col, h in enumerate(AOM_EVENT_HEADERS, 1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.font = HEADER_FONT_WHITE
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    for row_idx, item in enumerate(items, 4):
        covered_names = "、".join(
            item.get("event_names", {}).get(eid, eid)
            for eid in item.get("covered_events", [])
        )
        uncovered_names = "、".join(
            item.get("event_names", {}).get(eid, eid)
            for eid in item.get("uncovered_events", [])
        )

        status = item.get("status", "")
        if status == "no_core_instances":
            suggestion = "该服务无核心实例，暂不评估"
        elif status == "no_online_rules":
            suggestion = "需先创建AOM事件告警规则"
        elif item.get("bind_rate") == "100.0%":
            suggestion = "已达标"
        else:
            suggestion = "人工复核"

        values = [
            item.get("云服务简称", ""),
            item.get("baseline_count", 0),
            item.get("covered_count", 0),
            item.get("bind_rate", "N/A"),
            covered_names,
            uncovered_names,
            item.get("instance_count", 0),
            status,
            suggestion,
        ]

        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=(col >= 5))

    widths = [14, 12, 10, 12, 50, 50, 12, 18, 20]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    ws.auto_filter.ref = f"A3:{openpyxl.utils.get_column_letter(len(AOM_EVENT_HEADERS))}{3 + len(items)}"




def _write_aom_event_rows(ws, aom_report, start_row):
    """Append AOM event rows to the main sheet, skipping instance_count=0."""
    items = aom_report.get("items", [])
    aom_items = [item for item in items if item.get("instance_count", 0) > 0]

    for offset, item in enumerate(aom_items):
        row_idx = start_row + offset
        covered_names = "、".join(
            item.get("event_names", {}).get(eid, eid)
            for eid in item.get("covered_events", [])
        )
        uncovered_names = "、".join(
            item.get("event_names", {}).get(eid, eid)
            for eid in item.get("uncovered_events", [])
        )

        status = item.get("status", "")
        if status == "no_online_rules":
            suggestion = "需先创建AOM事件告警规则"
        elif item.get("bind_rate") == "100.0%":
            suggestion = "已达标"
        else:
            suggestion = "人工复核"

        values = [
            "AOM",
            item.get("云服务简称", ""),
            "(服务级)",
            "",
            item.get("bind_rate", "N/A"),
            covered_names,
            uncovered_names,
            "",
            suggestion,
        ]

        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=(col >= 5))

        rate_str = item.get("bind_rate", "")
        rate_cell = ws.cell(row=row_idx, column=5)
        if rate_str and rate_str != "N/A":
            try:
                rate_val = float(rate_str.rstrip("%"))
                if rate_val < 100:
                    rate_cell.fill = RED_FILL
                    rate_cell.font = RED_FONT
                else:
                    rate_cell.fill = GREEN_FILL
                    rate_cell.font = GREEN_FONT
            except ValueError:
                pass

    return len(aom_items)

def generate(report_path, output_path, notification_rules_path=None, ecs_special_events_path=None):
    report = load_json(report_path)
    by_instance_all = report.get("by_instance", [])

    event_to_rules = {}
    if notification_rules_path:
        try:
            event_to_rules = _build_event_to_rule_names(notification_rules_path)
        except Exception as e:
            logger.warning(f"Failed to load notification rules for rule name mapping: {e}")

    ecs_special_records = _load_ecs_special_event_names(ecs_special_events_path)
    ecs_special_display = ""
    if ecs_special_records:
        names = "、".join(r.get("事件ID", "") for r in ecs_special_records)
        ecs_special_display = f"人工复核：是否需要ECS智算场景的事件：【{names}】"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "事件基线绑定率"

    for col_idx, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT_WHITE
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    row_idx = 2
    for inst in by_instance_all:
            covered_names = ", ".join(ev.get("事件ID", "") for ev in inst.get("covered_events", []))
            missing_names = ", ".join(ev.get("事件ID", "") for ev in inst.get("missing_events", []))

            rule_names_set = set()
            for ev in inst.get("covered_events", []):
                svc = ev.get("云服务简称", "")
                eid = ev.get("事件ID", "")
                if svc and eid:
                    rule_svc = BASELINE_TO_RULE_SVC.get(svc, svc)
                    matched = event_to_rules.get((rule_svc, eid), [])
                    rule_names_set.update(matched)
            rule_names_str = ", ".join(sorted(rule_names_set))

            if inst.get("bound_status") == "no_baseline":
                suggestion = "该服务无事件基线"
            elif inst.get("service") == "ECS" and ecs_special_display:
                suggestion = ecs_special_display
            elif inst.get("bind_rate") == "100.0%":
                suggestion = "已达标"
            else:
                suggestion = "人工复核"

            row_data = [
                "CES",
                inst.get("svc_dim", inst.get("service", "")),
                inst.get("instance_id", ""),
                inst.get("region", ""),
                inst.get("bind_rate", "N/A"),
                covered_names,
                missing_names,
                rule_names_str,
                suggestion,
            ]

            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.border = THIN_BORDER
                cell.alignment = Alignment(vertical="center", wrap_text=(col_idx >= 5))

            row_idx += 1

    col_widths = [10, 14, 40, 18, 18, 50, 50, 30, 20]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

    ws.auto_filter.ref = ws.dimensions

    # Merge AOM event rows into the main sheet
    # Derive AOM input path from output_path (same workspace)
    aom_event_input_path = str(Path(output_path).parent / "intermediate" / "aom_event_bind_report.json")
    try:
        aom_event_report = load_json(aom_event_input_path)
        if aom_event_report and aom_event_report.get("items"):
            aom_count = _write_aom_event_rows(ws, aom_event_report, row_idx)
            logger.info(f"AOM event rows merged: {aom_count} rows")
    except Exception as e:
        logger.warning(f"Failed to merge AOM event rows: {e}")

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    logger.info(f"Event Excel report saved to {output_path}")
    logger.info(f"Total rows: {row_idx - 2}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generate event baseline binding rate Excel report")
    parser.add_argument("--input", default=DEFAULT_INPUT)
    parser.add_argument("--output", default=DEFAULT_OUTPUT)
    parser.add_argument("--notification-rules", default=DEFAULT_NOTIFICATION_RULES, help="Path to online_notification_rules.json for rule name mapping")
    parser.add_argument("--ecs-special-events", default=DEFAULT_ECS_SPECIAL_EVENTS, help="Path to ecs_special_events.json")
    args = parser.parse_args()
    generate(args.input, args.output, notification_rules_path=args.notification_rules, ecs_special_events_path=args.ecs_special_events)
