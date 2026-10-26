#!/usr/bin/env python3
import json
import re
import argparse
import logging
from pathlib import Path
from collections import defaultdict
from workspace import WORKSPACE_DIR, INTERMEDIATE_DIR

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
except ImportError:
    print("Error: openpyxl is required. Install with: pip install openpyxl")
    raise

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("excel-report")

DEFAULT_INPUT = str(Path(INTERMEDIATE_DIR) / "indicator_bind_report.json")
DEFAULT_RULES = str(Path(INTERMEDIATE_DIR) / "online_alarm_rules.json")
DEFAULT_BASELINE = str(Path(INTERMEDIATE_DIR) / "indicator_baseline.json")
DEFAULT_OUTPUT = str(Path(WORKSPACE_DIR) / "indicator_bind_report.xlsx")
DEFAULT_AOM_INPUT = str(Path(INTERMEDIATE_DIR) / "aom_indicator_bind_report.json")

HEADER_FONT = Font(bold=True, size=11)
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT_WHITE = Font(bold=True, size=11, color="FFFFFF")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
RED_FONT = Font(color="9C0006")
GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
GREEN_FONT = Font(color="006100")
BLUE_FONT = Font(color="0064FA")
ORANGE_FONT = Font(color="FC8800")

BASELINE_TO_RULE_SVC = {
    "ECS": "ECS",
    "RDS": "RDS",
    "NAT": "NAT",
    "DCS": "DCS",
    "DMS": "DMS",
    "ELB": "ELB",
    "WAF": "WAF",
    "CCE": "CCE",
    "GeminiDB": "GAUSSDBFORNOSQL",
    "DC": "DCAAS",
    "CC": "CC",
    "DDOS": "DDOS",
    "VPC": "VPC",
    "EVS": "EVS",
    "GEIP": "EIP",
    "CSS": "ES",
    "SFSTurbo": "EFS",
}

CORE_TO_BASELINE_SVC = {
    "EIP": "VPC",
}


def load_json(path):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def extract_uuids(text):
    if not text:
        return set()
    return set(re.findall(
        r'[a-f0-9]{8}-[a-f0-9]{4}-[a-f0-9]{4}-[a-f0-9]{4}-[a-f0-9]{12}',
        text,
    ))


def extract_dim_key(dim_value):
    if not dim_value:
        return ""
    return ",".join(kv.split("=")[0] for kv in dim_value.split(",") if "=" in kv)


def _inst_matches_resource(inst_id, rd):
    parent = rd.get("parent") or {}
    pid = parent.get("resource_id", "")
    rid = rd.get("resource_id", "")

    if pid:
        uuids_p = extract_uuids(pid)
        if uuids_p and inst_id in uuids_p:
            return "parent", parent
        if pid == inst_id:
            return "parent", parent
        if len(pid) in (32, 36) and pid == inst_id:
            return "parent", parent

    if rid and rid != "*":
        uuids_r = extract_uuids(rid)
        if uuids_r and inst_id in uuids_r:
            return "self", rd
        if rid == inst_id:
            return "self", rd
        if len(rid) in (32, 36) and rid == inst_id:
            return "self", rd

    return None, None


def _get_dim_key_for_match(match_type, matched_obj):
    if match_type == "parent":
        return extract_dim_key(matched_obj.get("dim_value", ""))
    else:
        return extract_dim_key(matched_obj.get("dim_value", ""))


def _dim_key_matches_baseline_dim(dim_key, baseline_dim):
    if not dim_key or not baseline_dim:
        return False
    return set(dim_key.split(",")) == set(baseline_dim.split(","))


def _extract_service_from_namespace(namespace):
    if not namespace:
        return ""
    parts = namespace.split(".")
    return parts[1] if len(parts) >= 2 else namespace


def get_rule_template_info(rule):
    """CES: no template concept — metrics are directly in policies[].metric_name."""
    name = rule.get("name", "CES告警规则")
    template_type = rule.get("alarm_template_id", "")
    if template_type:
        type_label = "告警模板"
        is_system = False
    else:
        type_label = "自定义规则"
        is_system = False
    metrics = set()
    for policy in rule.get("policies", []):
        mname = policy.get("metric_name", "")
        if mname:
            metrics.add(mname.lstrip("#"))
    return {
        "name": name,
        "type_label": type_label,
        "is_system": is_system,
        "metrics": metrics,
    }


def compare_with_baseline(template_name, template_metrics, baseline_metric_ids, metric_id_to_name):
    if not template_metrics and not baseline_metric_ids:
        return f"{template_name}：与基线一致"
    extra = template_metrics - baseline_metric_ids
    missing = baseline_metric_ids - template_metrics
    if not extra and not missing:
        return f"{template_name}：与基线一致"
    parts = []
    if extra:
        names = [metric_id_to_name.get(mid, mid) for mid in sorted(extra)]
        parts.append(f"{template_name}指标溢出：" + "、".join(names))
    if missing:
        names = [metric_id_to_name.get(mid, mid) for mid in sorted(missing)]
        parts.append(f"{template_name}指标缺失：" + "、".join(names))
    return "\n".join(parts)


def _is_ces_wildcard_rule(rule):
    resources = rule.get("resources", [])
    if not resources:
        return True
    for res in resources:
        for dim in res.get("dimensions", []):
            if dim.get("value", "") == "*":
                return True
    return False


def _ces_match_instance(inst_id, resources):
    """Check if inst_id matches any resource dimension. Return dim_key or None."""
    for res in resources:
        dimensions = res.get("dimensions", [])
        for dim in dimensions:
            value = dim.get("value", "")
            if not value or value == "*":
                continue
            if value == inst_id or inst_id in extract_uuids(value):
                dim_names = [d.get("name", "") for d in dimensions if d.get("name")]
                return ",".join(dim_names)
    return None


def _ces_wildcard_dim_key(rule):
    for policy in rule.get("policies", []):
        dn = policy.get("dimension_name", "")
        if dn:
            return dn
    return ""


def _collect_instance_dim_rule_info(inst_id, svc_rules, rule_svc, baseline_metric_ids, metric_id_to_name, baseline_dim, baseline_svc=""):
    rule_names = []
    template_names = []
    template_types = []
    match_results = []
    has_system = False

    for rule in svc_rules:
        if _is_ces_wildcard_rule(rule):
            dim_key = _ces_wildcard_dim_key(rule)
            if baseline_svc == "VPC" and dim_key == "bandwidth_id,publicip_id":
                dim_key = "publicip_id"
            rule_covers = _dim_key_matches_baseline_dim(dim_key, baseline_dim) if dim_key else False
        else:
            dim_key = _ces_match_instance(inst_id, rule.get("resources", []))
            if baseline_svc == "VPC" and dim_key == "bandwidth_id,publicip_id":
                dim_key = "publicip_id"
            rule_covers = _dim_key_matches_baseline_dim(dim_key, baseline_dim) if dim_key else False

        if not rule_covers:
            continue

        rule_names.append(rule.get("name", ""))
        tpl_info = get_rule_template_info(rule)
        template_names.append(tpl_info["name"])
        template_types.append(tpl_info["type_label"])
        if tpl_info["is_system"]:
            has_system = True
        match_results.append(compare_with_baseline(tpl_info["name"], tpl_info["metrics"], baseline_metric_ids, metric_id_to_name))

    return rule_names, template_names, template_types, match_results, has_system


def _format_rule_values(rule_names, template_names, template_types, match_results, has_system):
    if not rule_names:
        return "", "", "", "", ""

    rule_name_val = "\n".join(rule_names)
    tpl_name_val = "\n".join(template_names)
    tpl_type_val = "\n".join(template_types)
    match_val = "\n".join(match_results)
    suggestion_val = "优化系统模板" if has_system else "人工复核"
    return rule_name_val, tpl_name_val, tpl_type_val, match_val, suggestion_val


def _apply_cell_styles(ws, row_idx, bind_rate_value, match_results, has_system, suggestion_val):
    rate_cell = ws.cell(row=row_idx, column=5)
    if bind_rate_value < 1.0:
        rate_cell.fill = RED_FILL
        rate_cell.font = RED_FONT

    match_cell = ws.cell(row=row_idx, column=11)
    all_match = all("与基线一致" in r for r in match_results) if match_results else False
    if all_match:
        match_cell.fill = GREEN_FILL
        match_cell.font = GREEN_FONT
    elif match_results:
        match_cell.fill = RED_FILL
        match_cell.font = RED_FONT

    suggestion_cell = ws.cell(row=row_idx, column=12)
    if has_system:
        suggestion_cell.font = BLUE_FONT
    elif suggestion_val:
        suggestion_cell.font = ORANGE_FONT


def _set_column_widths(ws):
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 25
    ws.column_dimensions["F"].width = 20
    ws.column_dimensions["G"].width = 50
    ws.column_dimensions["H"].width = 50
    ws.column_dimensions["I"].width = 30
    ws.column_dimensions["J"].width = 30
    ws.column_dimensions["K"].width = 14
    ws.column_dimensions["L"].width = 45
    ws.column_dimensions["M"].width = 16


def _set_column_widths_summary(ws):
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 25
    ws.column_dimensions["F"].width = 20
    ws.column_dimensions["G"].width = 50
    ws.column_dimensions["H"].width = 50
    ws.column_dimensions["I"].width = 30
    ws.column_dimensions["J"].width = 30
    ws.column_dimensions["K"].width = 14
    ws.column_dimensions["L"].width = 45
    ws.column_dimensions["M"].width = 16


HEADERS = [
    "用户名", "云服务类型", "核心实例ID", "Region", "指标维度", "底线指标配置百分比",
    "覆盖的指标", "未覆盖的指标",
    "告警规则名称", "模板名称", "模板类型", "模板是否与基线匹配", "优化建议",
]

HEADERS_NO_CUST = [
    "监控体系", "云服务类型", "核心实例ID", "Region", "指标维度", "底线指标配置百分比",
    "覆盖的指标", "未覆盖的指标",
    "告警规则名称", "模板名称", "模板类型", "模板是否与基线匹配", "优化建议",
]


def _write_report_sheet(ws, instances, rules_by_svc, baseline_by_svc_dim, metric_id_to_name):
    headers = HEADERS_NO_CUST
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = HEADER_FONT_WHITE
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    for row_idx, inst in enumerate(instances, 2):
        core_svc = inst["service"]
        is_no_rules = inst.get("bound_status") == "no_online_rules"

        is_no_baseline = inst.get("bound_status") == "no_baseline"

        if is_no_rules:
            dim_cn = inst.get("metric_dimension_cn", "")
            baseline_dim = inst.get("metric_dimension", "")
            baseline_svc = CORE_TO_BASELINE_SVC.get(core_svc, core_svc)
            uncovered_ids = sorted(baseline_by_svc_dim.get((baseline_svc, baseline_dim), set()))
            uncovered_names = [metric_id_to_name.get(mid, mid) for mid in uncovered_ids]
            uncovered_str = "、".join(uncovered_names) if uncovered_names else ""
            values = [
                "CES",
                inst["service"],
                inst["instance_id"],
                inst.get("region", ""),
                dim_cn,
                inst["bind_rate"],
                "",
                uncovered_str,
                "",
                "",
                "",
                "",
                "需先创建告警规则",
            ]
            for col, val in enumerate(values, 1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                cell.border = THIN_BORDER
                cell.alignment = Alignment(vertical="center", wrap_text=(col >= 6))
            suggestion_col = 12
            suggestion_cell = ws.cell(row=row_idx, column=suggestion_col)
            suggestion_cell.font = ORANGE_FONT
            continue

        if is_no_baseline:
            dim_cn = inst.get("metric_dimension_cn", "")
            values = [
                "CES",
                inst["service"],
                inst["instance_id"],
                inst.get("region", ""),
                dim_cn,
                inst["bind_rate"],
                "",
                "",
                "",
                "",
                "",
                "",
                "该服务无指标基线定义",
            ]
            for col, val in enumerate(values, 1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                cell.border = THIN_BORDER
                cell.alignment = Alignment(vertical="center", wrap_text=(col >= 6))
            suggestion_col = 12
            suggestion_cell = ws.cell(row=row_idx, column=suggestion_col)
            suggestion_cell.font = BLUE_FONT
            continue

        coverage = inst.get("metrics_coverage", {})
        covered = [metric_id_to_name.get(mid, mid) for mid, v in coverage.items() if v]
        uncovered = [metric_id_to_name.get(mid, mid) for mid, v in coverage.items() if not v]

        baseline_svc = CORE_TO_BASELINE_SVC.get(core_svc, core_svc)
        rule_svc = BASELINE_TO_RULE_SVC.get(baseline_svc, baseline_svc)
        if rule_svc == "VPC" and core_svc == "EIP":
            rule_svc = "EIP"
        svc_rules = rules_by_svc.get(rule_svc, [])
        baseline_dim = inst.get("metric_dimension", "")
        baseline_metric_ids = baseline_by_svc_dim.get((baseline_svc, baseline_dim), set())

        inst_id = inst["instance_id"]
        rule_names, template_names, template_types, match_results, has_system = \
            _collect_instance_dim_rule_info(inst_id, svc_rules, rule_svc, baseline_metric_ids, metric_id_to_name, baseline_dim, baseline_svc)

        rule_name_val, tpl_name_val, tpl_type_val, match_val, suggestion_val = \
            _format_rule_values(rule_names, template_names, template_types, match_results, has_system)

        dim_cn = inst.get("metric_dimension_cn", "")
        dim_display = dim_cn if dim_cn else baseline_dim

        values = [
            "CES",
            inst["service"],
            inst["instance_id"],
            inst.get("region", ""),
            dim_display,
            inst["bind_rate"],
            "、".join(covered) if covered else "",
            "、".join(uncovered) if uncovered else "",
            rule_name_val,
            tpl_name_val,
            tpl_type_val,
            match_val,
            suggestion_val,
        ]

        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=(col >= 6))

        _apply_cell_styles(ws, row_idx, inst["bind_rate_value"], match_results, has_system, suggestion_val)

    _set_column_widths(ws)


def _build_rules_by_svc(rules):
    rules_by_svc = defaultdict(list)
    for rule in rules:
        svc = _extract_service_from_namespace(rule.get("namespace", ""))
        # Distinguish EIP from VPC
        if svc == "VPC":
            for res in rule.get("resources", []):
                for dim in res.get("dimensions", []):
                    if dim.get("name", "") == "publicip_id":
                        svc = "EIP"
                        break
        rules_by_svc[svc].append(rule)
    return dict(rules_by_svc)


def _build_baseline_lookup(baseline_data):
    baseline_by_svc_dim = defaultdict(set)
    metric_id_to_name = {}
    for rec in baseline_data.get("records", []):
        svc = rec["云服务简称"]
        dim = rec.get("指标维度", "")
        baseline_by_svc_dim[(svc, dim)].add(rec["指标ID"])
        metric_id_to_name[rec["指标ID"]] = rec["指标名称"]
    return dict(baseline_by_svc_dim), metric_id_to_name


def _enrich_metric_id_to_name(metric_id_to_name, rules):
    for rule in rules:
        for policy in rule.get("policies", []):
            mname = policy.get("metric_name", "").lstrip("#")
            if mname and mname not in metric_id_to_name:
                metric_id_to_name[mname] = mname

AOM_HEADERS = [
    "云服务类型", "指标维度", "底线指标数", "已覆盖数",
    "绑定率", "覆盖的指标", "未覆盖的指标",
    "核心实例数", "状态", "优化建议",
]


def _write_aom_sheet(ws, aom_report):
    """Write AOM indicator bind analysis as a separate sheet."""
    items = aom_report.get("items", [])
    summary = aom_report.get("summary", {})

    # Summary row
    ws.cell(row=1, column=1, value=f"AOM指标基线绑定率汇总："
            f"总底线指标 {summary.get('total_baseline_metrics', 0)}，"
            f"已覆盖 {summary.get('total_covered', 0)}，"
            f"覆盖率 {summary.get('overall_coverage', 'N/A')}").font = Font(bold=True, size=12)

    # Headers
    for col, h in enumerate(AOM_HEADERS, 1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.font = HEADER_FONT_WHITE
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    for row_idx, item in enumerate(items, 4):
        covered_names = "、".join(
            item.get("metric_names", {}).get(mid, mid)
            for mid in item.get("covered_metrics", [])
        )
        uncovered_names = "、".join(
            item.get("metric_names", {}).get(mid, mid)
            for mid in item.get("uncovered_metrics", [])
        )

        status = item.get("status", "")
        if status == "no_core_instances":
            suggestion = "该服务无核心实例，暂不评估"
        elif status == "no_online_rules":
            suggestion = "需先创建AOM告警规则"
        elif item.get("bind_rate") == "100.0%":
            suggestion = "已达标"
        else:
            suggestion = "人工复核"

        values = [
            item.get("云服务简称", ""),
            item.get("指标维度-中文", item.get("指标维度", "")),
            item.get("baseline_count", 0),
            item.get("covered_count", 0),
            item.get("bind_rate", "N/A"),
            covered_names,
            uncovered_names,
            item.get("instance_count", 0),
            status,
            suggestion,
        ]

        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=(col >= 6))

        # Color bind rate
        rate_str = item.get("bind_rate", "")
        rate_cell = ws.cell(row=row_idx, column=5)
        if rate_str and rate_str != "N/A":
            try:
                rate_val = float(rate_str.rstrip("%"))
                if rate_val < 100:
                    rate_cell.fill = RED_FILL
                    rate_cell.font = RED_FONT
                else:
                    rate_cell.fill = GREEN_FILL
                    rate_cell.font = GREEN_FONT
            except ValueError:
                pass

    # Column widths
    widths = [14, 30, 12, 10, 12, 50, 50, 12, 18, 20]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    ws.auto_filter.ref = f"A3:{openpyxl.utils.get_column_letter(len(AOM_HEADERS))}{3 + len(items)}"




def _write_aom_rows(ws, aom_report, start_row, svc_to_instance_ids):
    """Append AOM indicator rows to the main sheet, skipping instance_count=0."""
    items = aom_report.get("items", [])
    aom_items = [item for item in items if item.get("instance_count", 0) > 0]

    for offset, item in enumerate(aom_items):
        row_idx = start_row + offset
        covered_names = "、".join(
            item.get("metric_names", {}).get(mid, mid)
            for mid in item.get("covered_metrics", [])
        )
        uncovered_names = "、".join(
            item.get("metric_names", {}).get(mid, mid)
            for mid in item.get("uncovered_metrics", [])
        )

        status = item.get("status", "")
        if status == "no_online_rules":
            suggestion = "需先创建AOM告警规则"
        elif item.get("bind_rate") == "100.0%":
            suggestion = "已达标"
        else:
            suggestion = "人工复核"

        svc_name = item.get("云服务简称", "")
        inst_ids = svc_to_instance_ids.get(svc_name, [])
        inst_id_str = "、".join(inst_ids) if inst_ids else ""

        values = [
            "AOM",
            svc_name,
            inst_id_str,
            "",
            item.get("指标维度-中文", item.get("指标维度", "")),
            item.get("bind_rate", "N/A"),
            covered_names,
            uncovered_names,
            "",
            "",
            "",
            "",
            suggestion,
        ]

        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=(col >= 6))

        rate_str = item.get("bind_rate", "")
        rate_cell = ws.cell(row=row_idx, column=6)
        if rate_str and rate_str != "N/A":
            try:
                rate_val = float(rate_str.rstrip("%"))
                if rate_val < 100:
                    rate_cell.fill = RED_FILL
                    rate_cell.font = RED_FONT
                else:
                    rate_cell.fill = GREEN_FILL
                    rate_cell.font = GREEN_FONT
            except ValueError:
                pass

    return len(aom_items)

def generate(input_path, output_path, rules_path, baseline_path):
    report = load_json(input_path)
    rules_data = load_json(rules_path)
    baseline_data = load_json(baseline_path)

    baseline_by_svc_dim, metric_id_to_name = _build_baseline_lookup(baseline_data)

    rules = rules_data.get("rules", [])
    rules_by_svc = _build_rules_by_svc(rules)
    _enrich_metric_id_to_name(metric_id_to_name, rules)

    instances = report.get("by_instance", [])

    # Exclude AOM-monitored services from CES no_baseline instances
    aom_input_path = str(Path(output_path).parent / "intermediate" / "aom_indicator_bind_report.json")
    aom_services = set()
    try:
        aom_report_for_svcs = load_json(aom_input_path)
        for item in aom_report_for_svcs.get("items", []):
            aom_services.add(item.get("云服务简称", ""))
    except Exception:
        pass
    instances = [inst for inst in instances
                 if not (inst.get("bound_status") == "no_baseline" and inst.get("service") in aom_services)]

    instances.sort(key=lambda x: (x["service"], x.get("metric_dimension", ""), x["bind_rate_value"]))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "指标基线绑定率"
    _write_report_sheet(ws, instances, rules_by_svc, baseline_by_svc_dim, metric_id_to_name)

    # Merge AOM indicator rows into the main sheet
    # Derive AOM input path from output_path (same workspace)
    aom_input_path = str(Path(output_path).parent / "intermediate" / "aom_indicator_bind_report.json")
    # Load core instances for AOM instance IDs
    core_instances_path = str(Path(output_path).parent / "intermediate" / "core_instances.json")
    svc_to_instance_ids = {}
    try:
        core_data = load_json(core_instances_path)
        for svc, insts in core_data.get("instances", {}).items():
            svc_to_instance_ids[svc] = [i.get("instance_id", "") for i in insts]
    except Exception:
        pass

    try:
        aom_report = load_json(aom_input_path)
        if aom_report and aom_report.get("items"):
            aom_count = _write_aom_rows(ws, aom_report, len(instances) + 2, svc_to_instance_ids)
            logger.info(f"AOM indicator rows merged: {aom_count} rows")
    except Exception as e:
        logger.warning(f"Failed to merge AOM rows: {e}")

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    logger.info(f"Excel report saved to {output_path}")
    logger.info(f"Total rows: {len(instances)}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generate Excel report for bind rate")
    parser.add_argument("--input", default=DEFAULT_INPUT)
    parser.add_argument("--output", default=DEFAULT_OUTPUT)
    parser.add_argument("--rules", default=DEFAULT_RULES)
    parser.add_argument("--baseline", default=DEFAULT_BASELINE)
    args = parser.parse_args()
    generate(args.input, args.output, args.rules, args.baseline)
