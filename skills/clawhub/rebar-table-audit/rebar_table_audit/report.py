"""读 DXF、汇总成 evidence.json / 钢筋数量表.xlsx / report.md。"""
from __future__ import annotations

import datetime
import hashlib
import json
from pathlib import Path
from typing import Optional

import ezdxf
from ezdxf import recover

SCHEMA = "rebar-table-audit/1"


def _sanitize_text(s: str, counter: list[int]) -> str:
    """兜底：正常情况下 extract.normalize_text 已经把孤立代理项归一化掉了，text 字段不该
    再触发这里；但 raw 字段按设计保留归一化前的原文，理论上仍可能残留无法编码成 utf-8
    的孤立代理项，写 JSON/xlsx 前要有个安全网。用 backslashreplace（如 ``\\udccc``）而不是
    destructive 的 '?'，保留可读的字节信息，不静默丢数据；替换次数记进 counter[0]。"""
    try:
        s.encode("utf-8")
        return s
    except UnicodeEncodeError:
        counter[0] += 1
        return s.encode("utf-8", "backslashreplace").decode("utf-8")


def sanitize_evidence(evidence: dict) -> tuple[dict, int]:
    """深度遍历 evidence，清掉所有字符串里无法编码成 utf-8 的孤立代理项。返回 (清洗后的 evidence, 替换次数)。"""
    counter = [0]

    def walk(obj):
        if isinstance(obj, str):
            return _sanitize_text(obj, counter)
        if isinstance(obj, dict):
            return {k: walk(v) for k, v in obj.items()}
        if isinstance(obj, list):
            return [walk(v) for v in obj]
        return obj

    cleaned = walk(evidence)
    return cleaned, counter[0]


def sha256_of(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def read_dxf(path: Path):
    """ezdxf.readfile 优先；失败则 ezdxf.recover.readfile，并返回 auditor 错误数。"""
    try:
        doc = ezdxf.readfile(str(path))
        return doc, None
    except Exception:
        doc, auditor = recover.readfile(str(path))
        return doc, len(auditor.errors)


def build_source_info(original_path: Path, dxf_path: Path, doc) -> dict:
    return {
        "basename": original_path.name,
        "sha256": sha256_of(original_path),
        "bytes": original_path.stat().st_size,
        "dwg_version": doc.dxfversion,
    }


def build_evidence(
    source_path: Path,
    dxf_path: Path,
    doc,
    recover_error_count: Optional[int],
    conversion_info: Optional[dict],
    tables: list[dict],
    findings: list[dict],
    gaps: list[dict],
    symbol_mappings_used: list[dict],
) -> dict:
    findings_sorted = sorted(findings, key=lambda f: (f["table_id"], f["rule_id"], f["row_mark"] or ""))
    gaps_sorted = sorted(gaps, key=lambda g: (g.get("table_id") or "", g["rule_id"], g["message"]))
    tables_sorted = sorted(tables, key=lambda t: t["id"])

    conversion = dict(conversion_info) if conversion_info else {"tool": None, "tool_version": None, "dxf_sha256": sha256_of(dxf_path)}
    conversion["recover_used"] = recover_error_count is not None
    conversion["recover_error_count"] = recover_error_count

    evidence = {
        "schema": SCHEMA,
        "generated_at": datetime.datetime.now(datetime.timezone.utc).isoformat(),
        "source": build_source_info(source_path, dxf_path, doc),
        "conversion": conversion,
        "tables": tables_sorted,
        "findings": findings_sorted,
        "gaps": gaps_sorted,
        "symbol_mappings_used": symbol_mappings_used,
    }
    return evidence


def write_evidence_json(evidence: dict, path: Path) -> None:
    path.write_text(json.dumps(evidence, ensure_ascii=False, indent=2, sort_keys=False), encoding="utf-8")


def _grade_summary(tables: list[dict]) -> list[dict]:
    rows = []
    for t in tables:
        sums: dict[str, float] = {}
        for r in t["rows"]:
            p = r["parsed"]
            if p["grade"] and p["weight_kg"] is not None:
                sums[p["grade"]] = sums.get(p["grade"], 0.0) + p["weight_kg"]
        declared = {d["grade"]: d["value_kg"] for d in t["declared_totals"]}
        grades = sorted(set(sums) | set(declared))
        for g in grades:
            actual = sums.get(g)
            expect = declared.get(g)
            diff = (actual - expect) if (actual is not None and expect is not None) else None
            rows.append({"table_id": t["id"], "title": t["title"], "grade": g, "row_sum_kg": actual, "declared_kg": expect, "diff_kg": diff})
    return rows


def write_xlsx(evidence: dict, path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "汇总"
    ws_sum.append(["表 ID", "表标题", "牌号", "行重量之和(kg)", "声明合计(kg)", "差值(kg)"])
    for row in _grade_summary(evidence["tables"]):
        ws_sum.append([row["table_id"], row["title"], row["grade"], row["row_sum_kg"], row["declared_kg"], row["diff_kg"]])

    ws_detail = wb.create_sheet("明细")
    ws_detail.append(["表 ID", "编号", "位置", "牌号", "牌号来源", "直径(mm)", "单根长(m)", "根数", "总长(m)", "重量(kg)", "单位重(kg/m)", "行 handles"])
    for t in evidence["tables"]:
        for r in t["rows"]:
            p = r["parsed"]
            handles = ";".join(sorted({h for c in r["cells"] for h in c["handles"]}))
            ws_detail.append([
                t["id"], p["mark"], p["position"], p["grade"], p["grade_source"], p["diameter_mm"],
                p["single_length_m"], p["count"], p["total_length_m"], p["weight_kg"], p["unit_weight_kg_per_m"], handles,
            ])

    ws_find = wb.create_sheet("问题清单")
    ws_find.append(["规则", "级别", "表 ID", "编号", "疑点说明", "期望值", "实际值", "证据 handles"])
    for f in evidence["findings"]:
        ws_find.append([f["rule_id"], f["severity"], f["table_id"], f["row_mark"], f["message"], f["expected"], f["actual"], ";".join(f["handles"])])

    ws_gap = wb.create_sheet("证据缺口")
    ws_gap.append(["规则", "表 ID", "说明", "证据 handles"])
    for g in evidence["gaps"]:
        ws_gap.append([g["rule_id"], g.get("table_id"), g["message"], ";".join(g["handles"])])

    ws_sym = wb.create_sheet("符号映射")
    ws_sym.append(["符号", "牌号", "来源", "使用次数"])
    for s in evidence["symbol_mappings_used"]:
        ws_sym.append([s["code"], s["grade"], s["source"], s["count"]])

    wb.save(path)


def write_report_md(evidence: dict, path: Path) -> None:
    lines = []
    lines.append(f"# 钢筋数量表核对报告\n")
    lines.append(f"- 源文件：`{evidence['source']['basename']}`（sha256 `{evidence['source']['sha256'][:12]}…`，{evidence['source']['bytes']} 字节，DXF 版本 {evidence['source']['dwg_version']}）")
    conv = evidence["conversion"]
    if conv.get("tool"):
        lines.append(f"- 转换器：{conv['tool']}（版本 {conv.get('tool_version') or '未知'}），生成 DXF sha256 `{conv['dxf_sha256'][:12]}…`")
    else:
        lines.append(f"- 输入本身即 DXF，未经转换；sha256 `{conv['dxf_sha256'][:12]}…`")
    if conv.get("recover_used"):
        lines.append(f"- 注意：正常读取失败，走了 ezdxf.recover 兜底，auditor 报告 {conv.get('recover_error_count')} 处结构错误")
    lines.append(f"- 找到数量表：{len(evidence['tables'])} 张\n")

    for t in evidence["tables"]:
        lines.append(f"## {t['id']}：{t['title']}")
        lines.append(f"- 标题 handle：`{t['title_handle']}`；所在布局：{t['layout']}；行数：{len(t['rows'])}；列数：{len(t['columns'])}")
        if t["scope_notes"]:
            notes = "；".join(n["text"] for n in t["scope_notes"])
            lines.append(f"- 范围注释：{notes}")
        if t["declared_totals"]:
            decl = "；".join(f"{d['grade']} {d['value_kg']}kg" for d in t["declared_totals"])
            lines.append(f"- 表内声明合计：{decl}")
        lines.append("")

    sev_count = {}
    for f in evidence["findings"]:
        sev_count[f["severity"]] = sev_count.get(f["severity"], 0) + 1
    lines.append("## 疑点汇总")
    lines.append(f"- error: {sev_count.get('error', 0)}，warning: {sev_count.get('warning', 0)}，证据缺口: {len(evidence['gaps'])}\n")
    if evidence["findings"]:
        lines.append("| 规则 | 级别 | 表 | 编号 | 说明 | 期望 | 实际 |")
        lines.append("|---|---|---|---|---|---|---|")
        for f in evidence["findings"]:
            lines.append(f"| {f['rule_id']} | {f['severity']} | {f['table_id']} | {f['row_mark'] or ''} | {f['message']} | {f['expected']} | {f['actual']} |")
        lines.append("")

    if evidence["gaps"]:
        lines.append("## 证据缺口")
        lines.append("| 规则 | 表 | 说明 |")
        lines.append("|---|---|---|")
        for g in evidence["gaps"]:
            lines.append(f"| {g['rule_id']} | {g.get('table_id') or ''} | {g['message']} |")
        lines.append("")

    if evidence["symbol_mappings_used"]:
        lines.append("## 本次用到的符号映射（需人工确认 observed-in-sample 来源的条目）")
        lines.append("| 符号 | 牌号 | 来源 | 次数 |")
        lines.append("|---|---|---|---|")
        for s in evidence["symbol_mappings_used"]:
            lines.append(f"| {s['code']} | {s['grade']} | {s['source']} | {s['count']} |")
        lines.append("")

    lines.append("---")
    lines.append("算术不一致只是疑点，不代表设计错误；是否有误由设计/造价人员依据证据 handle 回图核实。")
    path.write_text("\n".join(lines), encoding="utf-8")
