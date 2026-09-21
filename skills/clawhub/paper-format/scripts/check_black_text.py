#!/usr/bin/env python3
"""Report explicit non-black text colors in visible DOCX content and used styles."""

from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


BLACK_VALUES = {None, "000000", "auto", "AUTO"}


def color_of_rpr(rpr):
    if rpr is None:
        return None
    color = rpr.find(qn("w:color"))
    if color is None:
        return None
    theme = color.get(qn("w:themeColor"))
    value = color.get(qn("w:val"))
    return f"theme:{theme}" if theme else value


def paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in doc.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs
        for table in section.header.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from cell.paragraphs
        for table in section.footer.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from cell.paragraphs


def main():
    if len(sys.argv) != 2:
        print(json.dumps({"status": "error", "message": "用法: check_black_text.py <成稿.docx>"}, ensure_ascii=False))
        raise SystemExit(2)
    path = Path(sys.argv[1])
    try:
        doc = Document(path)
    except Exception as exc:
        print(json.dumps({"status": "error", "message": f"无法解析 DOCX: {exc}"}, ensure_ascii=False))
        raise SystemExit(2)

    issues = []
    used_styles = set()
    for index, paragraph in enumerate(paragraphs(doc), 1):
        used_styles.add(paragraph.style.style_id)
        for run in paragraph.runs:
            value = color_of_rpr(run._r.rPr)
            if value not in BLACK_VALUES:
                issues.append({"paragraph": index, "text": run.text[:40], "color": value, "source": "direct"})

    for style in doc.styles:
        if style.style_id not in used_styles:
            continue
        value = color_of_rpr(style.element.rPr)
        if value not in BLACK_VALUES:
            issues.append({"style": style.name, "color": value, "source": "used_style"})

    payload = {
        "status": "pass" if not issues else "fail",
        "issues": issues,
        "note": "仅检查可解析文字及使用中的段落样式；图片、图表和复杂绘图对象仍需渲染检查。",
    }
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    raise SystemExit(0 if not issues else 1)


if __name__ == "__main__":
    main()
