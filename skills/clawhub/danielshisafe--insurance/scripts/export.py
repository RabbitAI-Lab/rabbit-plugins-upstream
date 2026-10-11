#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
export.py —— 分析报告「二次交付格式」导出器（HTML → PDF / Word / Excel）

★ 定位（SKILL §9.2）
  本技能的**主交付是 HTML**：保色、保版式、图表（SVG）原生可看，且任何环境都能出。
  HTML 出完以后，由模型**主动问询**使用人要不要再要 PDF / Word / Excel（可多选），
  使用人点选哪个，就用本脚本转哪个；选「不需要」就到此为止。

★ 三种格式各自的角色（别搞混）
  · PDF  = 「排版快照」：与 HTML 视觉一致，适合发人、打印、存档。
           走**无头浏览器打印**，不依赖任何 Python 第三方库。
  · Word = 「可编辑文档」：给使用人自己改字、加批注、走公司审批流。
           走 HTML → DOCX 结构化转换（标题/段落/表格/列表/**图片**）。
  · Excel= 「数据表 + 图」：把报告里的**表格**拆成一个个 sheet，
           图表按所属小节挂到对应 sheet 里，方便使用人自己再算、对着看。
           ⚠️ 导的是**表 + 图**，不是散文；报告正文不进 Excel（塞进去只会变成一坨）。

★ 图片：SVG → PNG（★ v1.5.1 起，Word / Excel 也能带图）
  Word 装不下内联 SVG，但装得下 PNG。故先把 SVG 栅格化成 PNG 再插进去：
  `chrome --headless --screenshot` 直接对 .svg 截图，零 Python 依赖。
  · 关键一步：必须把 SVG 根标签的 width/height **改成分辨率尺寸**（保留 viewBox），
    否则图只渲染在放大后窗口的左上角，四周全是空白（实测踩到）。
  · `--svg-scale 2`（默认）出 2 倍图，Word/Excel 里按 1 倍尺寸显示 → 打印更清晰。
  · 没浏览器 / 没 Pillow 时自动退回占位说明，并**明说原因**。

用法：
  python export.py --html 报告.html --to pdf
  python export.py --html 报告.html --to pdf docx xlsx
  python export.py --html 报告.html --to docx xlsx --svg-scale 3
  python export.py --html 报告.html --to xlsx --tables tables.json
  python export.py --html 报告.html --to pdf --browser edge

依赖：
  PDF  ：无（只需本机有 Chrome 或 Edge）
  DOCX ：python-docx + lxml（缺失时脚本明说）；插图还需浏览器（转 PNG）
  XLSX ：openpyxl（缺失时脚本明说）；插图还需 Pillow + 浏览器
"""

import argparse
import io
import os
import re
import subprocess
import sys

# ★ 用 reconfigure 而不是新建 TextIOWrapper：后者在本文件被 import 时
#   会接管并（GC 时）关闭调用方的 stdout.buffer，导致外层 print 报
#   "I/O operation on closed file"（实测踩到）。
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

HERE = os.path.dirname(os.path.abspath(__file__))

# ---------------------------------------------------------------- 打印样式
# HTML 没写 @media print 时自动注入：不注入则背景色全白、图超一页被腰斩。
PRINT_CSS = """
<style id="__export_print__">
*{-webkit-print-color-adjust:exact;print-color-adjust:exact}
@media print{
  @page{size:A4;margin:11mm 9mm}
  html,body{background:#fff}
  body{padding:0;font-size:10pt;line-height:1.5}
  section{padding:13px 15px;margin-bottom:11px;page-break-inside:auto}
  h1,h2,h3{page-break-after:avoid}
  figure,.note,tr,li{page-break-inside:avoid}
  table{font-size:9pt}
  figure svg{max-height:118mm}
}
</style>
"""

BROWSERS = [
    ("chrome", r"C:\Program Files\Google\Chrome\Application\chrome.exe"),
    ("chrome", r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe"),
    ("edge", r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
    ("edge", r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
]

BLOCK_TAGS = ("h1", "h2", "h3", "h4", "h5", "h6", "table", "figure",
              "ul", "ol", "p", "blockquote", "pre")


def _local(tag):
    """去掉 lxml 的 {namespace} 前缀并小写。

    ★ 内联 <svg> 在 lxml 里是 `{http://www.w3.org/2000/svg}svg`，
      直接 `tag.lower() == 'svg'` 永远不成立。
    """
    return tag.split("}")[-1].lower() if isinstance(tag, str) else ""


# ================================================================ PDF
def find_browser(prefer=None, user_path=None):
    if user_path:
        return user_path if os.path.exists(user_path) else None
    cands = BROWSERS
    if prefer in ("chrome", "edge"):
        cands = [c for c in BROWSERS if c[0] == prefer] + \
                [c for c in BROWSERS if c[0] != prefer]
    for _, p in cands:
        if os.path.exists(p):
            return p
    from shutil import which
    for exe in ("chrome", "msedge", "google-chrome", "chromium", "microsoft-edge"):
        w = which(exe)
        if w:
            return w
    return None


def inject_print_css(html, path_out):
    """没有 @media print 就注入一段；直接改原文件太粗暴，另存一个临时副本。"""
    if "@media print" in html:
        return path_out, False
    if "</head>" in html:
        html = html.replace("</head>", PRINT_CSS + "</head>", 1)
    else:
        html = PRINT_CSS + html
    tmp = os.path.splitext(path_out)[0] + ".__print__.html"
    with io.open(tmp, "w", encoding="utf-8") as f:
        f.write(html)
    return tmp, True


def to_pdf(html_path, out_path, browser=None, prefer=None, keep_tmp=False,
           virtual_time=8000):
    with io.open(html_path, encoding="utf-8", errors="replace") as f:
        html = f.read()
    src, injected = inject_print_css(html, html_path)

    exe = find_browser(prefer, browser)
    if not exe:
        return None, ("未找到 Chrome / Edge，无法导出 PDF。"
                      "可手动用浏览器打开 HTML → 打印 → 另存为 PDF；"
                      "或用 --browser 指定浏览器路径。")

    cmd = [
        exe, "--headless=new", "--disable-gpu", "--no-sandbox",
        "--no-pdf-header-footer", "--run-all-compositor-stages-before-draw",
        "--virtual-time-budget=%d" % virtual_time,
        "--print-to-pdf=%s" % os.path.abspath(out_path),
        os.path.abspath(src),
    ]
    try:
        r = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                           timeout=180)
    except Exception as e:
        return None, "调用浏览器失败：%r" % (e,)
    finally:
        if injected and not keep_tmp and os.path.exists(src):
            try:
                os.remove(src)
            except OSError:
                pass

    if not os.path.exists(out_path) or os.path.getsize(out_path) == 0:
        msg = (r.stderr or b"").decode("utf-8", "replace")[-400:]
        return None, "PDF 未生成（浏览器退出码 %s）：%s" % (r.returncode, msg)

    n = None
    try:
        import pypdfium2 as pypdf
        n = len(pypdf.PdfDocument(out_path))
    except Exception:
        pass
    return out_path, ("已注入打印样式" if injected else "HTML 已自带打印样式") + \
        ("，共 %d 页" % n if n else "")


# ================================================================ SVG → PNG
def prepare_svg(svg, scale=2):
    """把 SVG 根标签的 width/height 改成「原始尺寸 × scale」，并补齐 viewBox / xmlns。

    ★ 不改尺寸的话，Chrome 会把 980×580 的图画在 1960×1160 窗口的左上角，
      四周全白（实测四个象限只有左上 6.98% 有内容，其余 0%）。
    """
    m = re.match(r"<\s*svg\b[^>]*>", svg, re.S | re.I)
    if not m:
        return None
    tag = m.group(0)

    def attr(name):
        mm = re.search(name + r"\s*=\s*[\"']([^\"']*)[\"']", tag, re.I)
        return mm.group(1).strip() if mm else None

    def num(v):
        if not v:
            return None
        mm = re.search(r"(-?\d+(?:\.\d+)?)", str(v).replace("px", "").replace("pt", ""))
        return float(mm.group(1)) if mm else None

    w, h = num(attr("width")), num(attr("height"))
    vb = attr("viewBox") or attr("viewbox")
    if (not w or not h) and vb:
        p = vb.replace(",", " ").split()
        if len(p) == 4:
            w = w or float(p[2])
            h = h or float(p[3])
    if not w or not h:
        w, h = 900.0, 560.0
    W, H = max(int(round(w * scale)), 1), max(int(round(h * scale)), 1)

    new = re.sub(r"\s(?:width|height)\s*=\s*[\"'][^\"']*[\"']", "", tag).rstrip()
    if not new.endswith(">"):
        new += ">"
    new = new[:-1].rstrip() + ' width="%d" height="%d">' % (W, H)
    if not re.search(r"\bviewBox\s*=", new):
        new = new[:-1].rstrip() + ' viewBox="0 0 %g %g">' % (w, h)
    if not re.search(r"\bxmlns\s*=", new):
        new = '<svg xmlns="http://www.w3.org/2000/svg" ' + new[5:].lstrip()
    return svg.replace(tag, new, 1), W, H, w, h


def svg_to_png(svg_text, out_png, scale=2, browser=None, prefer=None,
               virtual_time=5000):
    """SVG 文本 → PNG。返回 (png路径, 原始宽, 原始高, 错误信息)。"""
    prep = prepare_svg(svg_text, scale)
    if not prep:
        return None, None, None, "SVG 结构解析失败（没找到 <svg> 根标签）"
    svg2, W, H, w, h = prep

    exe = find_browser(prefer, browser)
    if not exe:
        return None, None, None, ("未找到 Chrome / Edge，无法把 SVG 转成 PNG"
                                  "（可用 --browser 指定路径）")

    tmp = out_png + ".__tmp.svg"
    try:
        with io.open(tmp, "w", encoding="utf-8") as f:
            f.write(svg2)
        cmd = [
            exe, "--headless=new", "--disable-gpu", "--no-sandbox",
            "--hide-scrollbars", "--default-background-color=FFFFFFFF",
            "--virtual-time-budget=%d" % virtual_time,
            "--screenshot=%s" % os.path.abspath(out_png),
            "--window-size=%d,%d" % (W, H),
            os.path.abspath(tmp),
        ]
        subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                       timeout=180)
    except Exception as e:
        return None, None, None, "SVG 转 PNG 失败：%r" % (e,)
    finally:
        if os.path.exists(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass

    if not os.path.exists(out_png) or os.path.getsize(out_png) == 0:
        return None, None, None, "SVG 转 PNG 未产出文件（浏览器截图失败）"
    return out_png, w, h, None


# ================================================================ 块解析
def parse_blocks(html_path):
    """按**文档顺序**把 HTML 拆成块：heading / table / figure / list / para。

    ★ 只处理「叶子级」块元素，容器（div/section）不处理，让它的子元素被访问到；
      消费掉的块会把整棵子树标记为已消费，避免重复。
    ★ SVG 从**原始 HTML 文本**里正则取，不走 lxml 序列化——lxml 的 HTML 解析器
      会把大小写敏感的属性名（viewBox）全小写，导出的 SVG 就废了。
    """
    try:
        from lxml import html as LH
    except ImportError:
        return _parse_blocks_fallback(html_path)

    with io.open(html_path, encoding="utf-8", errors="replace") as f:
        raw = f.read()
    raw_svgs = re.findall(r"<\s*svg\b.*?<\s*/\s*svg\s*>", raw, re.S | re.I)

    doc = LH.parse(html_path)
    root = doc.getroot()
    blocks = []
    state = {"svg": 0}          # 递归里要改，用 dict 兜着

    def table_rows(el):
        rows = []
        for tr in el.iter("tr"):
            cells = [(c.text_content() or "").strip()
                     for c in tr.iter("th", "td")]
            if any(cells):
                rows.append(cells)
        return rows

    def rec(node):
        """★ 递归下降，而不是「全量迭代 + 已消费集合」。

        早期版本用 `id(el)` 做 consumed 集合，但 lxml 的元素代理对象随时会被 GC、
        **内存地址会被复用**，于是大量无关元素被误判成「已消费」——
        实测整篇只剩 2 个 figure 块，标题 / 表格 / 段落全被吞掉（0 段正文 0 张表）。
        递归下降天然不会重复处理，根本不需要那个集合。
        """
        for el in node:
            ln = _local(el.tag)
            if not ln:
                continue
            if ln in ("script", "style", "head", "nav", "noscript", "title",
                      "meta", "link"):
                continue

            if re.fullmatch(r"h[1-6]", ln):
                blocks.append({"type": "heading", "level": min(int(ln[1]), 4),
                               "el": el, "text": _norm((el.text_content() or ""))})

            elif ln == "table":
                rows = table_rows(el)
                if rows:
                    blocks.append({"type": "table", "rows": rows})

            elif ln == "figure":
                subs = list(el.iter())
                has_svg = any(_local(d.tag) == "svg" for d in subs)
                img_src = None
                for d in subs:
                    if _local(d.tag) == "img":
                        img_src = d.get("src")
                if not (has_svg or img_src):
                    rec(el)       # 没有图的 figure 当容器，让里面的 <p> 正常出来
                    continue
                cap = ""
                for d in subs:
                    if _local(d.tag) == "figcaption":
                        cap = _norm((d.text_content() or ""))
                        break
                blk = {"type": "figure", "caption": cap, "img_src": img_src,
                       "svg": None}
                if has_svg and state["svg"] < len(raw_svgs):
                    blk["svg"] = raw_svgs[state["svg"]]
                    state["svg"] += 1
                blocks.append(blk)

            elif ln in ("ul", "ol"):
                items = [d for d in el.iter() if _local(d.tag) == "li"]
                if items:
                    blocks.append({"type": "list", "ordered": ln == "ol",
                                   "items": items})

            elif ln in ("p", "blockquote", "pre"):
                if (el.text_content() or "").strip():
                    blocks.append({"type": "para", "el": el})

            elif ln in ("div", "section", "article", "main", "header", "footer"):
                # 没有块级子元素、且自己有文字 → 当成一个段落；否则继续往下走
                kids = [_local(d.tag) for d in el.iter() if d is not el]
                if not any(k in BLOCK_TAGS for k in kids) and \
                   (el.text_content() or "").strip():
                    blocks.append({"type": "para", "el": el})
                else:
                    rec(el)

            else:
                rec(el)

    rec(root)
    # 不在 <figure> 里的裸 <svg>，补成 figure 块，别丢图
    while state["svg"] < len(raw_svgs):
        blocks.append({"type": "figure", "caption": "", "img_src": None,
                       "svg": raw_svgs[state["svg"]]})
        state["svg"] += 1
    return blocks


def _parse_blocks_fallback(html_path):
    """lxml 缺失时的极简兜底：只认标题 / 表格 / SVG。"""
    with io.open(html_path, encoding="utf-8", errors="replace") as f:
        raw = f.read()
    blocks = []
    for m in re.finditer(r"<h([1-6])[^>]*>(.*?)</h\1>|<table.*?</table>|<svg\b.*?</svg>",
                         raw, re.S | re.I):
        seg = m.group(0)
        if seg[:3].lower() == "<sv":
            blocks.append({"type": "figure", "caption": "", "img_src": None,
                           "svg": seg})
        elif seg[:6].lower() == "<table":
            rows = []
            for rm in re.finditer(r"<tr.*?</tr>", seg, re.S | re.I):
                cells = [re.sub(r"<.*?>", "", c).replace("&nbsp;", " ").strip()
                         for c in re.findall(r"<t[hd].*?</t[hd]>", rm.group(0),
                                             re.S | re.I)]
                if any(cells):
                    rows.append(cells)
            if rows:
                blocks.append({"type": "table", "rows": rows})
        else:
            txt = _norm(re.sub(r"<.*?>", "", m.group(2)))
            if txt:
                blocks.append({"type": "heading",
                               "level": min(int(m.group(1)), 4),
                               "el": None, "text": txt, "raw_text": txt})
    return blocks


def prepare_images(blocks, img_dir, scale=2, browser=None, prefer=None,
                   use=True):
    """把 blocks 里的 SVG 统一栅格化成 PNG（重跑时复用同名文件）。

    直接改写 block：成功写 `png` / `pw` / `ph`，失败写 `png_err`。
    """
    n_ok = n_fail = fig_no = 0
    for i, b in enumerate(blocks):
        if b.get("type") != "figure":
            continue
        fig_no += 1     # ★ 文件名按「第几张图」编号，不要用块序号（否则会出现「图19.png」）
        if not use:
            b["png_err"] = "已用 --no-images 关闭插图"
            n_fail += 1
            continue
        src = b.get("svg")
        if not src:
            if b.get("img_src") and not b["img_src"].startswith(("http", "data:")):
                b["png"] = b["img_src"]       # 外链图片直接用
                b["pw"] = b["ph"] = None
                n_ok += 1
            else:
                b["png_err"] = "该图位没有可栅格化的 SVG"
                n_fail += 1
            continue
        os.makedirs(img_dir, exist_ok=True)
        png = os.path.join(img_dir, "图%d.png" % fig_no)
        if os.path.exists(png) and os.path.getsize(png) > 0:
            try:
                from PIL import Image
                w0, h0 = Image.open(png).size
                b["png"], b["pw"], b["ph"] = png, w0 / float(scale), h0 / float(scale)
                n_ok += 1
                continue
            except Exception:
                pass
        path, w, h, err = svg_to_png(src, png, scale, browser, prefer)
        if path:
            b["png"], b["pw"], b["ph"] = path, w, h
            n_ok += 1
        else:
            b["png_err"] = err
            n_fail += 1
    return n_ok, n_fail


# ================================================================ XLSX
SHEET_BAD = re.compile(r"[\[\]:*?/\\]")


def _sheet_name(raw, used):
    name = (raw or "表").strip() or "表"
    name = SHEET_BAD.sub("_", name)[:31] or "表"
    base, i = name, 2
    while name in used:
        suffix = "_%d" % i
        name = base[:31 - len(suffix)] + suffix
        i += 1
    used.add(name)
    return name


def _cell_value(s):
    """能转成数字就转数字（去千分位），否则留文本。

    ★ 不要转「万元」「倍」「%」这类带单位的：转了就把单位吃掉，反而误导。
      所以只接受**纯数字**与「带千分位的纯数字」两类。
    """
    t = (s or "").strip()
    if not t or t in ("-", "—", "–", "/", "N/A"):
        return t
    t2 = t.replace(",", "").replace("，", "")
    if re.fullmatch(r"-?\d+(\.\d+)?", t2):
        return float(t2) if "." in t2 else int(t2)
    return t


def _clean_sheet_title(s):
    """标题 → sheet 名：去掉编号前缀与 ★，长度留给 _sheet_name 统一截断。"""
    s = re.sub(r"^\s*\d+\s*[.、]?\s*", "", s or "")   # "2杠杆率对比表" / "2. …" / "2、…"
    s = s.replace("★", "").strip()
    return s


def _text_w(s):
    w = 0.0
    for ch in (s or ""):
        w += 2.0 if ord(ch) > 0x2E80 else 1.0
    return min(w + 2, 60)


def _norm(s):
    """HTML 源码里的换行/缩进 → 无（分段由块级标签决定）；连续空格 → 一个。"""
    s = re.sub(r"[ \t\r\f\v]*\n[ \t\r\f\v\n]*", "", s or "")
    return re.sub(r"[ \t]{2,}", " ", s)


def _group_sections(blocks):
    """按标题把 blocks 分组 → [{title, tables:[[rows]], figs:[block]}]。

    ★ 只有「含表或含图」的小节才会出 sheet——纯文字小节（如"计算口径"）不进 Excel。
    """
    secs = []
    cur = {"title": None, "tables": [], "figs": []}
    for b in blocks:
        t = b.get("type")
        if t == "heading":
            if cur["tables"] or cur["figs"]:
                secs.append(cur)
            cur = {"title": b.get("text", ""), "tables": [], "figs": []}
        elif t == "table":
            cur["tables"].append(b["rows"])
        elif t == "figure":
            cur["figs"].append(b)
    if cur["tables"] or cur["figs"]:
        secs.append(cur)
    return secs


def to_xlsx(html_path, out_path, blocks=None, tables_json=None, scale=2):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        return None, ("缺少 openpyxl，无法导出 Excel。"
                      "装：`pip install openpyxl`；或改用 CSV 交付（表格文本可直接粘进 Excel）。")

    if tables_json:
        import json
        with io.open(tables_json, encoding="utf-8") as f:
            data = json.load(f)
        tables = [{"name": k, "rows": v} for k, v in data.items()] \
            if isinstance(data, dict) else data
        secs = [{"title": t.get("name"), "tables": [t.get("rows") or []],
                 "figs": []} for t in tables if t.get("rows")]
    else:
        blocks = blocks if blocks is not None else parse_blocks(html_path)
        secs = _group_sections(blocks)

    if not secs:
        return None, "报告里没找到任何表格或图表，Excel 无内容可导（可改用 PDF/Word）。"

    have_pil = True
    try:
        from openpyxl.drawing.image import Image as XLImage
        import PIL  # noqa: F401
    except ImportError:
        have_pil = False

    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    n_sheet = n_img = n_noimg = 0
    for s in secs:
        ws = wb.create_sheet(_sheet_name(_clean_sheet_title(s["title"]), used))
        n_sheet += 1
        first = True
        for rows in s["tables"]:
            if not rows:
                continue
            if not first:
                ws.append([])
            for r in rows:
                ws.append([_cell_value(c) for c in r])
            first = False
        # 表头（每个表的第一行）加样式：row_i 逐表累加，中间空一行故 +1
        row_i = 1
        for rows in s["tables"]:
            if not rows:
                continue
            for c in ws[row_i]:
                if c.value is None:
                    continue
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="2F5597")
                c.alignment = Alignment(horizontal="center", vertical="center",
                                        wrap_text=True)
            row_i += len(rows) + 1
        if s["tables"]:
            ws.freeze_panes = "A2"
        # 列宽（CJK 记 2）
        width = {}
        for rows in s["tables"]:
            for r in rows:
                for j, c in enumerate(r, 1):
                    width[j] = max(width.get(j, 8), _text_w(str(c)))
        for j, w in width.items():
            ws.column_dimensions[get_column_letter(j)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, (int, float)):
                    c.alignment = Alignment(horizontal="right")
        # 图挂在��� sheet 末尾
        # ★ 锚点行必须自己累加：`ws.max_row` 不会因为 add_image 而增加，
        #   多张图会全部叠在同一个格子（实测两张图都锚在 A3，完全重合）。
        anchor_row = ws.max_row + 2
        for b in s["figs"]:
            png = b.get("png")
            if not png or not os.path.exists(png):
                n_noimg += 1
                ws.append([])
                ws.append(["［图表：%s］" % (b.get("png_err") or "未能栅格化，请看 HTML 或 PDF")])
                continue
            if not have_pil:
                n_noimg += 1
                ws.append([])
                ws.append(["［图表：缺少 Pillow，无法在 Excel 里插图；PNG 在 %s］" % png])
                continue
            try:
                img = XLImage(png)
                w = b.get("pw") or 900.0
                h = b.get("ph") or w * 0.6
                if w > 900:                       # 别撑爆一屏
                    h = h * 900.0 / w
                    w = 900.0
                img.width, img.height = int(w), int(h)
                ws.add_image(img, "A%d" % anchor_row)
                anchor_row += int(h / 20.0) + 3   # 默认行高 ≈20px，给图留出位置
                n_img += 1
            except Exception as e:
                n_noimg += 1
                ws.append([])
                ws.append(["［图表插图失败：%r］" % (e,)])

    if n_sheet == 0:
        return None, "解析到的内容都是空的，Excel 未生成。"
    wb.save(out_path)
    summary = "；".join(s.title for s in wb.worksheets)
    extra = "" if not n_img else "，含 %d 张图" % n_img
    bad = "" if not n_noimg else "，%d 张图未能插入" % n_noimg
    return out_path, "%d 个 sheet（%s）%s%s" % (n_sheet, summary, extra, bad)


# ================================================================ DOCX
def to_docx(html_path, out_path, blocks=None, title=None):
    try:
        from docx import Document
        from docx.shared import Pt, Inches, Cm, RGBColor
        from docx.oxml.ns import qn
    except ImportError as e:
        return None, ("缺少 python-docx，无法导出 Word：%s。"
                      "装：`pip install python-docx`。"
                      "或在本机用 Office/Word 直接打开 HTML 另存为 .docx。" % e)

    if blocks is None:
        try:
            blocks = parse_blocks(html_path)
        except ImportError as e:
            return None, ("缺少 lxml，无法解析 HTML：%s。装：`pip install lxml`。" % e)

    doc = Document()
    # 中文字体：西文与东亚都要设，否则 Word 里中文回落到宋体/方框
    st = doc.styles["Normal"]
    st.font.name = "Microsoft YaHei"
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    # 纸张改成 A4：python-docx 默认模板是 Letter + 1.25in 页边距，
    # 与 PDF 的 A4 不一致，图会按 6.0in 宽插入、打印时再缩放，白白糊一层。
    sec0 = doc.sections[0]
    try:
        sec0.page_width, sec0.page_height = Cm(21.0), Cm(29.7)
        sec0.left_margin = sec0.right_margin = Cm(2.2)
        sec0.top_margin = sec0.bottom_margin = Cm(2.0)
    except Exception:
        pass
    usable = sec0.page_width - sec0.left_margin - sec0.right_margin

    n_p = n_tb = n_img = n_noimg = 0
    skip_tags = {"script", "style", "nav", "noscript"}

    def add_runs(par, node):
        """把 <b>/<i> 之类内联标签转成 run，保留加粗斜体。

        ★ HTML 源码里的换行/缩进不是真的分段（分段由 <p> 决定），
          不归一化会在 Word 里变成一堆空行，故把「含换行的空白」整段抹掉。
        """
        if node.text:
            par.add_run(_norm(node.text))
        for ch in node:
            if not isinstance(ch.tag, str):
                continue
            if _local(ch.tag) in skip_tags:
                continue
            sub, tail = (ch.text or ""), (ch.tail or "")
            if _local(ch.tag) in ("b", "strong"):
                par.add_run(_norm(sub)).bold = True
            elif _local(ch.tag) in ("i", "em"):
                par.add_run(_norm(sub)).italic = True
            elif _local(ch.tag) == "br":
                par.add_run("\n")
            else:
                par.add_run(_norm(sub))
            for gc in ch:                       # <b><i>x</i></b> 之类嵌套
                if isinstance(gc.tag, str) and _local(gc.tag) not in skip_tags:
                    add_runs(par, gc)
            if tail:
                par.add_run(_norm(tail))

    def note(text, color=0x996600):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.font.size = Pt(9)
        r.italic = True
        r.font.color.rgb = RGBColor((color >> 16) & 255, (color >> 8) & 255, color & 255)
        return p

    if title:
        doc.add_heading(title, level=0)

    for b in blocks:
        t = b.get("type")

        if t == "heading":
            if b.get("el") is not None:
                h = doc.add_heading(level=b["level"])
                add_runs(h, b["el"])
            else:
                doc.add_heading(b.get("raw_text") or b.get("text") or "",
                                level=b["level"])
            n_p += 1

        elif t == "table":
            rows = b["rows"]
            ncol = max(len(r) for r in rows)
            tb = doc.add_table(rows=0, cols=ncol)
            try:
                tb.style = "Light Grid Accent 1"
            except KeyError:
                tb.style = "Table Grid"
            for i, r in enumerate(rows):
                cells = tb.add_row().cells
                for j in range(ncol):
                    v = _norm(r[j] if j < len(r) else "")
                    cells[j].text = ""
                    run = cells[j].paragraphs[0].add_run(v)
                    if i == 0:
                        run.bold = True
                    for cc in cells[j].paragraphs[0].runs:
                        cc.font.size = Pt(9)
            n_tb += 1
            doc.add_paragraph()

        elif t == "list":
            for li in b["items"]:
                p = doc.add_paragraph(style="List Number" if b["ordered"]
                                      else "List Bullet")
                add_runs(p, li)
                n_p += 1

        elif t == "para":
            if not (b["el"].text_content() or "").strip():
                continue
            p = doc.add_paragraph()
            add_runs(p, b["el"])
            n_p += 1

        elif t == "figure":
            if b.get("caption"):
                cp = doc.add_paragraph()
                cr = cp.add_run(b["caption"])
                cr.font.size = Pt(9)
                cr.italic = True
                cr.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
                n_p += 1
            png = b.get("png")
            if png and os.path.exists(png):
                try:
                    w_in = Inches((b.get("pw") or 6.5) / 96.0)
                    if w_in > usable:
                        w_in = usable
                    doc.add_picture(png, width=w_in)
                    n_img += 1
                    continue
                except Exception as e:
                    note("［图表插入失败：%r］" % (e,))
                    n_noimg += 1
                    continue
            note("［图表：%s］" % (b.get("png_err") or
                                 "Word 不支持内嵌 SVG，请看 HTML 或 PDF 版本"))
            n_noimg += 1

    if n_p == 0 and n_tb == 0 and n_img == 0:
        return None, "HTML 里没解析出任何正文内容，Word 未生成。"
    doc.save(out_path)
    parts = ["%d 段正文 / %d 张表" % (n_p, n_tb)]
    parts.append("%d 张图已插入" % n_img)
    if n_noimg:
        parts.append("%d 张图未插入" % n_noimg)
    return out_path, "，".join(parts)


# ================================================================ CLI
def main():
    ap = argparse.ArgumentParser(
        description="分析报告二次交付格式导出：HTML → PDF / Word / Excel",
        formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--html", required=True, help="主交付 HTML 报告")
    ap.add_argument("--to", nargs="+", required=True,
                    choices=["pdf", "docx", "xlsx"], help="要导出的格式（可多选）")
    ap.add_argument("--out-dir", default=None, help="输出目录，默认与 HTML 同目录")
    ap.add_argument("--name", default=None, help="输出文件名前缀，默认取 HTML 主名")
    ap.add_argument("--title", default=None,
                    help="Word 用：强制指定大标题；不给则用正文 <h1>，不拿文件名凑")
    ap.add_argument("--tables", default=None,
                    help="Excel 用：表格 JSON（{sheet名:[[行]]} 或 [{name,rows}]）；"
                         "不给则自动从 HTML 抽表 + 图")
    ap.add_argument("--browser", default=None,
                    help="PDF / SVG 转图 用：指定浏览器 exe 路径")
    ap.add_argument("--prefer", default=None, choices=["chrome", "edge"],
                    help="PDF / SVG 转图 用：两者都有时偏好哪个")
    ap.add_argument("--svg-scale", type=float, default=2.0,
                    help="SVG 转 PNG 的倍率（默认 2，越大越清晰、文件越大）")
    ap.add_argument("--img-dir", default=None,
                    help="PNG 存放目录，默认 <out-dir>/<name>_图")
    ap.add_argument("--no-images", action="store_true",
                    help="不栅格化插图，Word/Excel 里一律留占位说明")
    ap.add_argument("--keep-tmp", action="store_true",
                    help="保留注入打印样式后的临时 HTML，便于排查")
    a = ap.parse_args()

    html = os.path.abspath(a.html)
    if not os.path.exists(html):
        print("✗ 找不到 HTML：%s" % html)
        return 2
    out_dir = os.path.abspath(a.out_dir) if a.out_dir else os.path.dirname(html)
    if not os.path.isdir(out_dir):
        os.makedirs(out_dir, exist_ok=True)
    base = a.name or os.path.splitext(os.path.basename(html))[0]
    img_dir = os.path.abspath(a.img_dir) if a.img_dir else \
        os.path.join(out_dir, base + "_图")

    print("源报告：%s" % html)
    print("目标格式：%s" % " / ".join(a.to))
    print("-" * 62)

    need_img = (not a.no_images) and ("docx" in a.to or "xlsx" in a.to)
    blocks = None
    if "docx" in a.to or "xlsx" in a.to:
        blocks = parse_blocks(html)
        n_fig = sum(1 for b in blocks if b.get("type") == "figure")
        if n_fig and need_img:
            ok, bad = prepare_images(blocks, img_dir, a.svg_scale,
                                     a.browser, a.prefer, use=True)
            print("图表栅格化：成功 %d / 失败 %d -> %s" % (ok, bad, img_dir))
        elif n_fig:
            prepare_images(blocks, img_dir, a.svg_scale, a.browser, a.prefer,
                           use=False)
            print("图表：%d 张（--no-images，仅占位说明）" % n_fig)
        else:
            print("图表：HTML 里没有 <svg>/<img>")
        print("-" * 62)

    ok, fail = [], []
    for fmt in a.to:
        out = os.path.join(out_dir, "%s.%s" % (base, fmt))
        if fmt == "pdf":
            path, msg = to_pdf(html, out, a.browser, a.prefer, a.keep_tmp)
        elif fmt == "xlsx":
            path, msg = to_xlsx(html, out, blocks=blocks, tables_json=a.tables,
                                scale=a.svg_scale)
        else:
            path, msg = to_docx(html, out, blocks=blocks, title=a.title)
        if path:
            ok.append((fmt, out, msg))
            print("✓ %-5s -> %s  (%s, %s)" %
                  (fmt.upper(), out, "{:,} 字节".format(os.path.getsize(out)), msg))
        else:
            fail.append((fmt, msg))
            print("✗ %-5s 失败：%s" % (fmt.upper(), msg))

    print("-" * 62)
    print("成功 %d / 失败 %d" % (len(ok), len(fail)))
    if fail:
        print("\n⚠ 未导出的格式请改用其它方式交付，并在报告里说明原因，不要省略不提。")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
