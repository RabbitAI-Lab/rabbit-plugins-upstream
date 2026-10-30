#!/usr/bin/env python
"""
台账导出 Excel 服务

从台账生成 .xlsx 产物（openpyxl 可用，缺失时自动 pip 安装兜底，仍失败降级 .csv），
并把成功任务的本地产物图（按 task_id_标题 重命名）与表格一起压成 zip 包，
落盘到 1688-item-batch-opt-apply/ 下，返回产物绝对路径与摘要，供主 Agent以
present_files 非交互式展示/下载（钉钉文档「Excel（Skill）、文件下载支持，
非交互式，在产物中展示」口径）。
"""

import csv
import os
import re
import subprocess
import sys
import zipfile

from _const import OP_LABELS, store_file_prefix
from _errors import ServiceError
from _ledger import load_ledger, summary_stats, ledger_dir

try:
    import openpyxl
    _HAS_OPENPYXL = True
except Exception:
    _HAS_OPENPYXL = False


def _ensure_openpyxl():
    """懒加载 openpyxl，缺失时自动 pip 安装兜底（沿用 white_bg Pillow 先例）。

安装失败或超时返回 None，调用方降级 csv；不阻断导出主流程。
    """
    global openpyxl, _HAS_OPENPYXL
    if _HAS_OPENPYXL:
        return openpyxl
    try:
        subprocess.run(
            [sys.executable, "-m", "pip", "install", "--quiet", "openpyxl"],
            check=True, timeout=120,
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
        )
        import openpyxl as _opx
        openpyxl = _opx
        _HAS_OPENPYXL = True
        return openpyxl
    except Exception:
        return None

COLUMNS = [
    ("task_id", "任务ID", 10),
    ("item_id", "商品ID", 16),
    ("item_title", "商品标题", 28),
    ("op_label", "优化项", 12),
    ("requirements", "要求", 20),
    ("status", "执行状态", 10),
    ("result_summary", "执行结果", 30),
    ("apply_status", "应用状态", 10),
    ("apply_detail", "应用回执", 24),
    ("error", "失败原因", 24),
]

_STATUS_ZH = {
    "parsed": "待确认", "confirmed": "待执行", "skipped": "已跳过",
    "running": "执行中", "success": "成功", "failed": "失败", "cancelled": "已取消",
}
_APPLY_ZH = {
    "pending": "未应用", "requested": "应用中", "applied": "已应用",
    "failed": "应用失败", "unknown": "结果未知", "cancelled": "已取消应用",
}


def _cell_value(task: dict, key: str) -> str:
    if key == "op_label":
        return task.get("op_label") or OP_LABELS.get(task.get("op"), task.get("op") or "")
    if key == "status":
        return _STATUS_ZH.get(str(task.get("status")), str(task.get("status")))
    if key == "apply_status":
        return _APPLY_ZH.get(str(task.get("apply_status")), str(task.get("apply_status")))
    v = task.get(key)
    return "" if v is None else str(v)


def _safe_name(value: str) -> str:
    """文件名安全化：非法字符与空白转下划线，限长 40。"""
    cleaned = re.sub(r"[\\/:*?\"<>|\s]+", "_", str(value or "").strip())
    return cleaned[:40] or "untitled"


def _package_results(tasks: list, out_dir: str, base_name: str, table_path: str):
    """把台账 result_images 的本地产物图（按 task_id_标题 重命名）与表格压成 zip。

只收本地文件（线上 https 链接无法也不必入包）；没有可入包图片时返回 None，
不产出空包。zip 内结构：表格文件在根，图片在 images/ 下。
    """
    package_files = []
    used = set()
    for task in tasks:
        if str(task.get("status")) != "success":
            continue
        stem = _safe_name("{}_{}".format(
            task.get("task_id"), task.get("item_title") or task.get("op_label") or ""))
        idx = 0
        for img in task.get("result_images") or []:
            img = str(img)
            if not img or img.startswith(("http://", "https://", "data:")):
                continue
            if not os.path.isfile(img):
                continue
            ext = os.path.splitext(img)[1] or ".png"
            arc = "images/{}{}".format(stem, ext)
            while arc in used:
                idx += 1
                arc = "images/{}_{}{}".format(stem, idx, ext)
            used.add(arc)
            package_files.append((img, arc))
    if not package_files:
        return None
    zip_path = os.path.join(out_dir, base_name + ".zip")
    try:
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.write(table_path, os.path.basename(table_path))
            for src, arc in package_files:
                zf.write(src, arc)
    except OSError as exc:
        raise ServiceError("产物打包失败：{}".format(exc))
    return zip_path


def _unique_base(out_dir: str, base: str, exts: tuple) -> str:
    """产物文件名冲突兜底：目标表格/zip 已存在（历史批次撞名、重复导出、
    多店铺多批次同会话运行）时自增序号（_2、_3…），表格与 zip 共用同一
    变体名，绝不静默覆盖既有产物。"""
    candidate = base
    seq = 1
    while any(os.path.exists(os.path.join(out_dir, candidate + ext)) for ext in exts):
        seq += 1
        candidate = "{}_{}".format(base, seq)
    return candidate


def export_results(base_dir: str = None, only_done: bool = False) -> dict:
    ledger = load_ledger(base_dir)
    tasks = ledger.get("tasks", [])
    if only_done:
        tasks = [t for t in tasks if str(t.get("status")) in ("success", "failed", "skipped", "cancelled")]
    if not tasks:
        raise ServiceError("台账中没有可导出的任务；请先运行 parse_tasks 并落账结果")

    opx = _ensure_openpyxl()
    stats = summary_stats(ledger)
    out_dir = ledger_dir(base_dir)
    os.makedirs(out_dir, exist_ok=True)
    batch_id = ledger.get("batch_id") or "batch"
    # 店铺前缀（台账 store_name 字段，parse_tasks --store_name 写入）：多店铺同
    # 会话运行时产物文件名不重叠；无店铺名时退化为旧版无前缀命名。
    store_prefix = store_file_prefix(ledger.get("store_name"))
    ident = "{}_{}".format(store_prefix, batch_id) if store_prefix else str(batch_id)
    base_name = "批量任务结果_{}".format(ident)
    base_name = _unique_base(
        out_dir, base_name, (".xlsx", ".zip") if opx is not None else (".csv", ".zip")
    )

    if opx is not None:
        out_path = os.path.join(out_dir, base_name + ".xlsx")
        wb = opx.Workbook()
        ws = wb.active
        ws.title = "任务结果"
        for col, (_k, header, width) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = opx.styles.Font(bold=True)
            ws.column_dimensions[opx.utils.get_column_letter(col)].width = width
        for r, task in enumerate(tasks, start=2):
            for c, (key, _h, _w) in enumerate(COLUMNS, start=1):
                ws.cell(row=r, column=c, value=_cell_value(task, key))
        try:
            wb.save(out_path)
        except OSError as exc:
            raise ServiceError("Excel 产物写入失败：{}".format(exc))
    else:
        out_path = os.path.join(out_dir, base_name + ".csv")
        try:
            with open(out_path, "w", encoding="utf-8-sig", newline="") as f:
                writer = csv.writer(f)
                writer.writerow([h for _k, h, _w in COLUMNS])
                for task in tasks:
                    writer.writerow([_cell_value(task, k) for k, _h, _w in COLUMNS])
        except OSError as exc:
            raise ServiceError("CSV 产物写入失败：{}".format(exc))

    if not (os.path.isfile(out_path) and os.path.getsize(out_path) > 0):
        raise ServiceError("导出产物异常（文件为空），请重试")

    package_path = _package_results(tasks, out_dir, base_name, out_path)

    result = {
        "file_path": os.path.abspath(out_path),
        "file_format": "xlsx" if opx is not None else "csv",
        "row_count": len(tasks),
        "batch_id": batch_id,
        "store_name": ledger.get("store_name") or "",
        "stats": stats,
    }
    if package_path and os.path.isfile(package_path) and os.path.getsize(package_path) > 0:
        result["package_path"] = os.path.abspath(package_path)
        result["package_file_count"] = len([1 for _f in zipfile.ZipFile(package_path).infolist()])
    return result
