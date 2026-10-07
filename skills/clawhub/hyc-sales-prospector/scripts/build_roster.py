#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
SALES-PROSPECTOR 标准名录生成器
输入：JSON（企业数据）
输出：CSV(utf-8-sig) + XLSX（带 A 级高亮）

用法：
    python build_roster.py <input.json> <output_dir>
"""

import csv
import json
import sys
import os
from datetime import date

HEADERS = [
    "序号", "企业名称", "招牌名/门店名", "所在行政区", "详细地址",
    "联系电话", "联系人", "联系人来源", "主营/经营范围", "规模信号",
    "关联企业", "潜力评级", "评级理由", "建议切入话术", "数据来源",
]

KEYS = [
    "name", "shop_name", "district", "address", "phone", "contact",
    "contact_source", "business", "scale", "related", "grade",
    "reason", "script", "source",
]

FILL = "未公开"

# 中文列宽（近似字符数）
WIDTHS = [5, 34, 20, 16, 40, 20, 14, 12, 40, 30, 34, 10, 26, 40, 26]


def normalize(row, idx):
    out = [idx]
    for k in KEYS:
        v = row.get(k, "")
        if v is None or str(v).strip() == "":
            v = FILL
        out.append(str(v).strip())
    return out


def sort_rows(rows):
    """A 级优先，其次按评级、再按行政区聚类"""
    order = {"A": 0, "B": 1, "C": 2}
    return sorted(
        rows,
        key=lambda r: (
            order.get(str(r.get("grade", "C")).strip().upper()[:1], 3),
            str(r.get("district", "")),
        ),
    )


def write_csv(path, data):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(HEADERS)
        w.writerows(data)


def write_xlsx(path, data):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        return False

    wb = Workbook()
    ws = wb.active
    ws.title = "扫街名单"

    head_fill = PatternFill("solid", fgColor="1F3864")
    head_font = Font(color="FFFFFF", bold=True, size=11)
    a_fill = PatternFill("solid", fgColor="FFF2CC")
    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(HEADERS)
    for c in ws[1]:
        c.fill = head_fill
        c.font = head_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border

    for r in data:
        ws.append(r)

    for i, row in enumerate(ws.iter_rows(min_row=2), start=0):
        grade = ""
        if len(data[i]) > 11:
            grade = str(data[i][11]).strip().upper()[:1]
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = border
            if grade == "A":
                c.fill = a_fill
        if grade == "A":
            row[1].font = Font(bold=True)

    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "B2"
    wb.save(path)
    return True


def main():
    if len(sys.argv) < 3:
        print("usage: build_roster.py <input.json> <output_dir>")
        sys.exit(1)

    src, outdir = sys.argv[1], sys.argv[2]
    with open(src, "r", encoding="utf-8") as f:
        payload = json.load(f)

    region = payload.get("region", "区域")
    industry = payload.get("industry", "行业")
    day = payload.get("date", date.today().strftime("%Y%m%d"))
    rows = sort_rows(payload.get("rows", []))

    data = [normalize(r, i + 1) for i, r in enumerate(rows)]

    os.makedirs(outdir, exist_ok=True)
    base = os.path.join(outdir, f"{region}_{industry}_扫街名单_{day}")
    csv_path = base + ".csv"
    write_csv(csv_path, data)
    print("CSV:", csv_path)

    xlsx_path = base + ".xlsx"
    if write_xlsx(xlsx_path, data):
        print("XLSX:", xlsx_path)
    else:
        print("XLSX: skipped (openpyxl 未安装)")

    # 统计（"未公开"/"—" 视为缺失）
    def valid(v):
        s = str(v or "").strip()
        return bool(s) and s not in ("未公开", "—", "-", "待补充")

    total = len(data)
    graded = {}
    with_phone = 0
    with_contact = 0
    with_addr = 0
    for r in rows:
        g = str(r.get("grade", "")).strip().upper()[:1]
        if g:
            graded[g] = graded.get(g, 0) + 1
        if valid(r.get("phone")):
            with_phone += 1
        if valid(r.get("contact")):
            with_contact += 1
        if valid(r.get("address")) and "未公开" not in str(r.get("address")):
            with_addr += 1
    print("总数:", total)
    print("分级:", graded)
    if total:
        print("电话覆盖率: %.0f%% (%d/%d)" % (with_phone * 100.0 / total, with_phone, total))
        print("联系人覆盖率: %.0f%% (%d/%d)" % (with_contact * 100.0 / total, with_contact, total))
        print("地址覆盖率: %.0f%% (%d/%d)" % (with_addr * 100.0 / total, with_addr, total))


if __name__ == "__main__":
    main()
