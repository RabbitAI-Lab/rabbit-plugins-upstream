#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
水利计价 Excel 整理器
清理青山大禹导出的"分页重复表头"，输出规范化的 Excel：
  表头加粗 / 冻结首行 / 列宽自适应 / 金额千分位 / 层级缩进 / 合计行高亮
用法: python3 tidy_excel.py <源文件.xlsx> <输出.xlsx>
"""
import sys, re
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

HDR_FILL = PatternFill('solid', fgColor='305496')
HDR_FONT = Font(name='等线', size=11, bold=True, color='FFFFFF')
BODY = Font(name='等线', size=10)
SUB_FILL = PatternFill('solid', fgColor='DDEBF7')
SUB_FONT = Font(name='等线', size=10, bold=True)
TOT_FILL = PatternFill('solid', fgColor='FFF2CC')
TOT_FONT = Font(name='等线', size=10, bold=True, color='C00000')
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = '#,##0.00'
NUM = '#,##0.00'

TITLE_KW = ('合同编号', '工程名称', '组    号', '单位：')
HDR_KW = ('序号', '名称', '单位', '数量', '单价', '合价', '项目', '金额', '规格',
          '型号', '编号', '费率', '价格', '型号规格', '备注', '分组', '定额外')


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def tidy_sheet(ws):
    rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
    # 1) 丢弃全空行
    rows = [r for r in rows if any(c is not None and str(c).strip() != '' for c in r)]
    if not rows:
        return None
    # 2) 找首行表头（含>=2 个表头关键词的行）
    hdr_i = None
    for i, r in enumerate(rows[:8]):
        s = ' '.join(str(c) for c in r if c)
        if sum(1 for k in HDR_KW if k in s) >= 2:
            hdr_i = i; break
    if hdr_i is None:
        hdr_i = 0
    title = ' '.join(str(c) for c in rows[0] if c)[:40]
    hdr = rows[hdr_i]
    # 3) 去掉表头之后重复出现的 标题行/表头行
    data = []
    seen_hdr = False
    for r in rows[hdr_i + 1:]:
        s = ' '.join(str(c) for c in r if c)
        if sum(1 for k in HDR_KW if k in s) >= 3 and not any(  # 重复表头行
                is_num(c) for c in r) and any(str(c).strip() == '' for c in r):
            if seen_hdr:
                continue
            seen_hdr = True
            continue
        # 分页标题行：含合同编号/工程名称等，且本行无任何数字单元格
        if any(k in s for k in TITLE_KW) and not any(is_num(c) for c in r):
            continue
        data.append(r)
    return title, hdr, data


def write(ws, title, hdr, data):
    ws.append([str(c).replace('\n', '') if c is not None else '' for c in hdr])
    for c in range(1, len(hdr) + 1):
        cell = ws.cell(1, c)
        cell.fill = HDR_FILL; cell.font = HDR_FONT
        cell.alignment = Alignment('center', 'center', wrap_text=True)
        cell.border = BORDER
    for r in data:
        ws.append(list(r))
    for r in range(2, ws.max_row + 1):
        first = str(ws.cell(r, 1).value or '')
        # 层级判断：编号为 1.1 / 1.1.1 形式
        depth = first.count('.') if re.match(r'^[\d.]+$', first.strip()) else -1
        is_tot = '合计' in ' '.join(str(ws.cell(r, c).value or '') for c in range(1, len(hdr) + 1))
        is_sub = depth == 0
        for c in range(1, len(hdr) + 1):
            cell = ws.cell(r, c)
            cell.border = BORDER
            if is_tot:
                cell.font = TOT_FONT; cell.fill = TOT_FILL
            elif is_sub:
                cell.font = SUB_FONT; cell.fill = SUB_FILL
            else:
                cell.font = BODY
            if is_num(cell.value):
                cell.number_format = MONEY
                cell.alignment = Alignment('right', 'center')
            else:
                al = Alignment('left', 'center', wrap_text=len(str(cell.value or '')) > 30)
                if depth > 0:
                    al = Alignment('left', 'center', wrap_text=True, indent=depth)
                cell.alignment = al
    for i in range(1, len(hdr) + 1):
        vals = [len(str(ws.cell(r, i).value or '')) for r in range(1, min(ws.max_row, 300) + 1)]
        w = max(vals) if vals else 10
        ws.column_dimensions[get_column_letter(i)].width = min(max(w * 1.3 + 2, 9), 46)
    ws.freeze_panes = 'A2'
    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(hdr))}{ws.max_row}"
    ws.row_dimensions[1].height = 30


def main():
    src, out = sys.argv[1], sys.argv[2]
    wb = load_workbook(src, data_only=True)
    nwb = Workbook(); nwb.remove(nwb.active)
    stat = []
    for nm in wb.sheetnames:
        r = tidy_sheet(wb[nm])
        if not r:
            continue
        title, hdr, data = r
        if not data:
            continue
        ws = nwb.create_sheet(nm[:31])
        write(ws, title, hdr, data)
        stat.append((nm, len(data)))
    nwb.save(out)
    print('生成:', out)
    for nm, n in stat:
        print(f"  [{nm}] {n} 行")


if __name__ == '__main__':
    main()
