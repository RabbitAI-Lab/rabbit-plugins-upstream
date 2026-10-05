# -*- coding: utf-8 -*-
"""填写《混合式教学设计课程教学日历模版.docx》，输出《混合式教学教学日历.docx》。

用法：
    python fill_calendar.py <design_data.json> [输出目录]

字段结构见 python scripts/ref_reader.py design_plan_guide（calendar 数组）。
"""
import os
import sys
from copy import deepcopy

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from _common import ensure_deps, load_json, add_word_footer, open_template  # noqa: E402

CAL_KEYS = ["seq", "date", "week", "lecture", "minutes", "location",
            "students", "content", "form", "activity"]


def main():
    if not ensure_deps():
        print("依赖未就绪，请先运行 scripts/_bootstrap_deps.py 或请 AI 工具协助。")
        return 1

    from docx import Document  # 延迟导入

    if len(sys.argv) < 2:
        print("用法：python fill_calendar.py <design_data.json> [输出目录]")
        return 1

    data = load_json(sys.argv[1])
    out_dir = sys.argv[2] if len(sys.argv) > 2 else os.path.dirname(os.path.abspath(sys.argv[1]))
    os.makedirs(out_dir, exist_ok=True)

    doc = open_template("混合式教学设计课程教学日历模版.docx")
    t = doc.tables[0]

    # ---- 1. 课程名称（文档正文"课程名称："后追加） ----
    bi = data.get("basic_info", {}) or {}
    course_name = (bi.get("course_name") or "").strip()
    if course_name:
        for p in doc.paragraphs:
            if p.text.strip().startswith("课程名称"):
                p.add_run(course_name)
                break

    # ---- 2. 日历表格：删除示例行，按数据行重填 ----
    rows = data.get("calendar", []) or []
    if len(t.rows) > 3:
        tpl_tr = deepcopy(t.rows[3]._tr)  # 以"……"行为样式模板
        # 删除第 1、2、3 行（示例行）
        for ri in (3, 2, 1):
            t._tbl.remove(t.rows[ri]._tr)
        for item in rows:
            new_tr = deepcopy(tpl_tr)
            t._tbl.append(new_tr)
            r = t.rows[-1]
            vals = [str(item.get(k, "") or "") for k in CAL_KEYS]
            for ci, v in enumerate(vals):
                if ci < len(r.cells):
                    r.cells[ci].text = v
    else:
        print("提示：日历模版结构异常，未填写表格。")

    # ---- 3. 页脚：空 5 行后插入页脚说明 ----
    add_word_footer(doc)

    out = os.path.join(out_dir, "混合式教学教学日历.docx")
    doc.save(out)
    print("已生成：%s" % out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
