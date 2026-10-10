# -*- coding: utf-8 -*-
"""共享工具：路径、依赖准备、JSON 读取、参考资料读取、页脚读取与 Word 页脚插入。"""
import os
import sys
import json
import base64
import tempfile

try:
    import _refs  # 参考资料读取入口（同目录）
except ImportError:  # 兜底：直接以脚本方式加载时补上模块搜索路径
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import _refs

PKG_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ASSETS_DIR = os.path.join(PKG_ROOT, "assets")
REFERENCES_DIR = os.path.join(PKG_ROOT, "references")
TEMPLATES_DIR = os.path.join(ASSETS_DIR, "templates")
VENDOR_DIR = os.path.join(ASSETS_DIR, "vendor")
SCRIPTS_DIR = os.path.dirname(os.path.abspath(__file__))

# references/ 下应存在的参考资料（以专用格式存放，须经 scripts/ref_reader.py 读取）
REFERENCE_FILES = _refs.REF_NAMES
REFERENCE_REQUIRED_KEYS = ("kind", "name", "title", "version")
# 每个参考资料必须非空的关键内容字段（用于自检时校验内容是否完整）
REFERENCE_CONTENT_KEYS = {
    "methodology": ("toc", "sections"),
    "bloom_reference": ("blocks",),
    "design_plan_guide": ("sections",),
    "faq": ("error_handling", "questions"),
    "output_preview": ("sections",),
    "footer_note": ("entries", "version_line"),
}


def ensure_deps():
    """确保 python-docx / python-pptx 可用；缺失时调用 _bootstrap_deps.py 自动安装。"""
    try:
        import docx  # noqa: F401
        import pptx  # noqa: F401
        return True
    except ImportError:
        pass
    try:
        import subprocess
        r = subprocess.run(
            [sys.executable, os.path.join(SCRIPTS_DIR, "_bootstrap_deps.py")],
            cwd=PKG_ROOT,
        )
        if r.returncode != 0:
            return False
        import docx  # noqa: F401
        import pptx  # noqa: F401
        return True
    except Exception as e:  # pragma: no cover
        print("依赖准备失败：%s" % e)
        return False


# 必需字段 -> 人类可读的栏目名（用于字段级校验提示）
REQUIRED_FIELDS = (
    ("basic_info.course_name", ("basic_info", "course_name"), "课程名称"),
    ("basic_info.major", ("basic_info", "major"), "面向专业"),
    ("basic_info.content", ("basic_info", "content"), "授课内容"),
    ("basic_info.section", ("basic_info", "section"), "所属课程节次"),
    ("basic_info.hours", ("basic_info", "hours"), "学时安排"),
    ("basic_info.students", ("basic_info", "students"), "授课人数"),
    ("whole_course.design_idea", ("whole_course", "design_idea"), "整门课程 · 混合式教学设计思路"),
    ("whole_course.innovation", ("whole_course", "innovation"), "整门课程 · 创新举措"),
    ("whole_course.effect", ("whole_course", "effect"), "整门课程 · 效果达成情况"),
    ("single_lesson.learning_analysis", ("single_lesson", "learning_analysis"), "一课次 · 学情分析"),
    ("single_lesson.learning_goals", ("single_lesson", "learning_goals"), "一课次 · 学习目标"),
    ("single_lesson.content_resources", ("single_lesson", "content_resources"), "一课次 · 内容与资源"),
    ("single_lesson.process_method", ("single_lesson", "process_method"), "一课次 · 过程与方法"),
    ("single_lesson.evaluation", ("single_lesson", "evaluation"), "一课次 · 评价与反馈"),
    ("single_lesson.effect", ("single_lesson", "effect"), "一课次 · 教学效果达成情况"),
    ("calendar", ("calendar",), "课程教学日历（每讲次一行）"),
)

# 模版占位符文本（出现即说明对应栏目未填写）
PLACEHOLDERS = (
    "（整门课程的混合式教学设计思路）",
    "（整门课程的创新举措）",
    "（整门课程的效果达成情况）",
    "（本课次学生的学情分析）",
    "（本课次学生能够掌握和展现的具体知识、技能和能力等，对课程目标的贡献等情况）",
    "（本课次内容与资源的选取、制作、使用情况）",
    "（本课次混合教学的实施过程与方法）",
    "（本课次学习评价与反馈方式）",
    "（本课次教学效果与特色）",
)


def _get_path(data, path):
    cur = data
    for key in path:
        if not isinstance(cur, dict):
            return None
        cur = cur.get(key)
    return cur


def _has_text(val):
    """判断字段是否"有实质内容"（支持字符串 / 字典 / 列表）。"""
    if val is None:
        return False
    if isinstance(val, str):
        return bool(val.strip())
    if isinstance(val, dict):
        return any(_has_text(v) for v in val.values())
    if isinstance(val, (list, tuple)):
        return len(val) > 0
    return True


def validate_json(data):
    """字段级校验：返回 (错误列表, 警告列表)，均为人类可读中文提示。

    只做"提醒"，不阻断生成——缺失的栏目会保持模版原样，由 AI 补齐后重跑更好。
    """
    errors, warnings = [], []
    if not isinstance(data, dict):
        return ["数据文件顶层不是 JSON 对象（应为一对大括号 {}）。"], []

    for _, path, label in REQUIRED_FIELDS:
        if not _has_text(_get_path(data, path)):
            errors.append("缺少「%s」的内容（%s）" % (label, ".".join(path)))

    # 日历逐行检查
    cal = data.get("calendar")
    if isinstance(cal, list) and cal:
        seqs = [str((row or {}).get("seq", "")).strip() for row in cal]
        if len(set(seqs)) != len(seqs):
            warnings.append("教学日期的「序号」存在重复，请改为连续且不重复。")
        for row in cal:
            row = row or {}
            loc = str(row.get("location", "") or "")
            if any(k in loc for k in ("大学", "学院", "校区", "校园", "学校")):
                warnings.append("日历第 %s 讲的授课地点疑似出现学校/校区名称：「%s」，"
                                "请改为「教室」，非集中在线学习填「在线」。" % (row.get("seq", "?"), loc))
                break

    # 三测题检查
    quizzes = (data.get("interactions", {}) or {}).get("quizzes") or []
    if quizzes:
        for qz in quizzes:
            for q in (qz.get("questions") or []):
                opts = q.get("options") or []
                if len(opts) != 4:
                    warnings.append("测试题「%s」的选项不是 4 个（应为 1 正确 + 3 干扰）。"
                                    % str(q.get("q", ""))[:20])
                    break
                a = q.get("answer")
                if not isinstance(a, int) or not (0 <= a < 4):
                    warnings.append("测试题「%s」的 answer 应为 0–3 的整数。"
                                    % str(q.get("q", ""))[:20])
                    break
                if not str(q.get("explain", "") or "").strip():
                    warnings.append("测试题「%s」缺少解析（explain），学生作答后将得不到反馈。"
                                    % str(q.get("q", ""))[:20])
                    break

    # 残留占位符检查（内容字段里出现模版占位符，说明照抄了模版）
    flat = json.dumps(data, ensure_ascii=False)
    for ph in PLACEHOLDERS:
        if ph in flat:
            warnings.append("内容中仍含模版占位符「%s」，请替换为实质内容。" % ph)

    return errors, warnings


def report_validation(data, strict=False):
    """打印校验结果。strict=True 时存在错误即返回 False。"""
    errors, warnings = validate_json(data)
    if errors:
        print("【数据自检】发现 %d 处必需内容缺失：" % len(errors))
        for e in errors:
            print("  ✗ %s" % e)
        print("  提示：可先按以上清单补齐后重跑；缺失栏目将保持模版原样。")
    if warnings:
        print("【数据自检】%d 处建议改进：" % len(warnings))
        for w in warnings:
            print("  ! %s" % w)
    if not errors and not warnings:
        print("【数据自检】通过：必需栏目齐全，未发现明显问题。")
    return not (errors and strict)


def load_json(path, validate=True):
    """读取 UTF-8 JSON 文件，出错时给出友好提示；可顺带做字段级自检。

    当由 build.py 统一调度时（命令行带 --from-build 标记），本处的自检转为静默，
    避免同一条校验信息在每个子脚本里重复打印。
    """
    if "--from-build" in sys.argv:
        validate = False
    try:
        with open(path, "r", encoding="utf-8-sig") as f:
            data = json.load(f)
    except FileNotFoundError:
        print("【错误 E01】找不到数据文件：%s" % path)
        print("  原因：路径写错，或 design_data.json 尚未生成。")
        print("  修复：确认路径正确；若文件不存在，请让 AI 重新撰写并写出该 JSON。")
        sys.exit(1)
    except json.JSONDecodeError as e:
        print("【错误 E02】数据文件不是有效的 JSON：%s" % path)
        print("  原因：%s" % e)
        print("  修复：常见于多余逗号、中文引号（“”‘’）或单引号；请让 AI 重新生成该文件，"
              "字符串务必用英文双引号。")
        sys.exit(1)
    except UnicodeDecodeError:
        print("【错误 E03】数据文件编码不是 UTF-8：%s" % path)
        print("  修复：请让 AI 以 UTF-8 编码重新写出该文件。")
        sys.exit(1)
    if validate:
        report_validation(data)
    return data


def read_footer():
    """读取页脚说明（references/footer_note.json），返回逐行列表。

    返回结构为 [标题, "", "1.…", "", "2.…", … , "", "Ver：…"]，
    与历史上的纯文本页脚逐行一一对应，供 Word/HTML/PPT 原样插入。
    """
    try:
        data = _refs.load("footer_note")
    except _refs.RefMissing:
        print("【错误 E06】找不到页脚说明文件：%s" % _refs.ref_path("footer_note"))
        print("  原因：技能包不完整，或该文件被移动、改名。")
        print("  修复：重新解压技能包；references/ 下的文件请勿改名或挪动。")
        sys.exit(1)
    except _refs.RefError as e:
        print("【错误 E07】页脚说明无法读取：%s" % e)
        print("  原因：技能包不完整，或该文件被改动、损坏。")
        print("  修复：重新解压技能包；references/ 下的文件请勿手工编辑。")
        sys.exit(1)
    except Exception as e:
        print("【错误 E07】页脚说明无法读取：%s" % e)
        print("  修复：重新解压技能包。")
        sys.exit(1)

    lines = [str(data.get("title", "") or "").replace("\r\n", "\n").replace("\r", "\n")]
    for e in (data.get("entries") or []):
        lines.append("")
        text = "%s.%s" % (e.get("no", ""), e.get("text", ""))
        lines.extend(text.replace("\r\n", "\n").replace("\r", "\n").split("\n"))
    ver = str(data.get("version_line", "") or "").strip()
    if ver:
        lines.append("")
        lines.extend(ver.split("\n"))
    return lines


def check_references():
    """校验 references/ 下参考资料是否齐全、可读取、结构完整。

    返回有问题的条目说明列表（空列表 = 全部正常）。
    """
    bad = []
    for name in REFERENCE_FILES:
        try:
            data = _refs.load(name)
        except _refs.RefMissing:
            bad.append("%s.json（缺失）" % name)
            continue
        except _refs.RefError:
            bad.append("%s.json（无法读取）" % name)
            continue
        except Exception:
            bad.append("%s.json（无法读取）" % name)
            continue
        if not isinstance(data, dict):
            bad.append("%s.json（结构异常）" % name)
            continue
        missing = [k for k in REFERENCE_REQUIRED_KEYS
                   if not str(data.get(k, "") or "").strip()]
        if missing:
            bad.append("%s.json（缺字段 %s）" % (name, "/".join(missing)))
            continue
        lack = [k for k in REFERENCE_CONTENT_KEYS.get(name, ()) if not data.get(k)]
        if lack:
            bad.append("%s.json（内容为空：%s）" % (name, "/".join(lack)))
    return bad


def add_word_footer(doc):
    """在 Word 文档末尾：空 5 行，再插入页脚说明（逐行原样插入，左对齐）。"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    for _ in range(5):
        doc.add_paragraph("")
    for line in read_footer():
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def open_template(name):
    """打开 Word 模板文档，返回 python-docx Document 对象。

    SkillHub 不支持 zip 内携带 .docx 文件，故模板以 base64 形式存于
    assets/templates/templates.json（键为模板文件名）。加载顺序：
    1) 若同目录存在 <name>.docx 则直接打开（本地调试兼容）；
    2) 否则从 templates.json 解码该模板的 base64 为临时 docx 后打开。
    """
    from docx import Document
    docx_path = os.path.join(TEMPLATES_DIR, name)
    if os.path.exists(docx_path):
        return Document(docx_path)
    json_path = os.path.join(TEMPLATES_DIR, "templates.json")
    with open(json_path, "r", encoding="utf-8") as f:
        b64 = json.load(f)[name]
    raw = base64.b64decode(b64)
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
    try:
        tmp.write(raw)
    finally:
        tmp.close()
    return Document(tmp.name)


def strip_lines(text):
    """把多行文本拆成去掉首尾空白的行列表，并去掉完全空白的行。"""
    if not text:
        return []
    return [ln.strip() for ln in text.split("\n") if ln.strip()]
