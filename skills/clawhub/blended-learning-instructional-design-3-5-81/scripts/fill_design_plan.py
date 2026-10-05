# -*- coding: utf-8 -*-
"""填写《混合式教学设计方案模版.docx》，输出《混合式教学设计方案.docx》。

用法：
    python fill_design_plan.py <design_data.json> [输出目录]

字段结构见 python scripts/ref_reader.py design_plan_guide。
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from _common import ensure_deps, load_json, add_word_footer, open_template  # noqa: E402


def set_cell_text(cell, text):
    """把文本写入单元格：按换行拆分为多段落、左对齐、保留缩进，便于阅读。

    约定：文本中每行一段；行首用 2 个空格缩进表示子项（· 教师活动等），
    脚本按比例设置左缩进；以【】开头的行作为阶段标题加粗显示。
    """
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    text = (text or "").strip()
    cell.text = ""  # 清空原内容
    if not cell.paragraphs:
        cell.add_paragraph()
    lines = text.split("\n")
    for i, raw in enumerate(lines):
        line = raw.rstrip()
        n_indent = len(raw) - len(raw.lstrip(" "))
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.text = line.lstrip(" ") or ""
        if n_indent:
            p.paragraph_format.left_indent = Pt(11 * (n_indent // 2))
        for r in p.runs:
            r.font.size = Pt(10.5)
            if line.startswith("【"):
                r.font.bold = True
                r.font.size = Pt(11)


def build_quiz_text(data):
    """将前测/中测/后测 渲染为『题目 + 参考答案 + 解析』纯文本（写入 Word 评价与反馈）。"""
    quizzes = (data.get("interactions", {}) or {}).get("quizzes")
    if not quizzes:
        try:
            from gen_html_courseware import auto_quizzes
            quizzes = auto_quizzes(data)
        except Exception:
            quizzes = []
    if not quizzes:
        return ""
    blocks = []
    for qz in quizzes:
        lines = ["【%s】" % qz.get("title", "测试")]
        for qi, q in enumerate(qz.get("questions", []), 1):
            opts = q.get("options") or []
            a = q.get("answer", 0)
            ans = opts[a] if 0 <= a < len(opts) else ""
            lines.append("Q%d. %s" % (qi, q.get("q", "")))
            lines.append("  参考答案：%s" % ans)
            fb = q.get("explain", "")
            if fb:
                lines.append("  解析：%s" % fb)
        blocks.append("\n".join(lines))
    return "\n\n".join(blocks)


def build_goals_text(lg):
    """将学习目标渲染为 知识/技能/素养/思政 四类文本（兼容旧版字符串）。"""
    if isinstance(lg, dict):
        labels = [("knowledge", "知识目标"), ("skill", "技能目标"),
                  ("literacy", "素养目标"), ("ideology", "思政目标")]
        parts = []
        for key, label in labels:
            v = lg.get(key, "")
            if v and str(v).strip():
                parts.append("【%s】\n%s" % (label, str(v).strip()))
        return "\n\n".join(parts)
    return str(lg or "")


def main():
    if not ensure_deps():
        print("依赖未就绪，请先运行 scripts/_bootstrap_deps.py 或请 AI 工具协助。")
        return 1

    from docx import Document  # 延迟导入，确保依赖先就绪

    if len(sys.argv) < 2:
        print("用法：python fill_design_plan.py <design_data.json> [输出目录]")
        return 1

    data = load_json(sys.argv[1])
    out_dir = sys.argv[2] if len(sys.argv) > 2 else os.path.dirname(os.path.abspath(sys.argv[1]))
    os.makedirs(out_dir, exist_ok=True)

    doc = open_template("混合式教学设计方案模版.docx")
    t = doc.tables[0]

    # ---- 1. 基本信息（表头） ----
    bi = data.get("basic_info", {}) or {}
    basic_cells = {  # JSON 字段 -> (行, 列)，依据模版合并单元格实测网格
        "course_name": (1, 1),
        "major": (1, 3),
        "content": (2, 0),
        "section": (2, 2),
        "hours": (2, 5),
        "students": (3, 1),
    }
    for key, (r, c) in basic_cells.items():
        val = bi.get(key)
        if val and str(val).strip():
            set_cell_text(t.cell(r, c), str(val))

    # ---- 2. 整门课程 + 一课次内容区（按占位符文本匹配） ----
    wc = data.get("whole_course", {}) or {}
    sl = data.get("single_lesson", {}) or {}
    eval_full = sl.get("evaluation") or ""
    quiz_block = build_quiz_text(data)
    if quiz_block:
        if str(eval_full).strip():
            eval_full = (str(eval_full)
                         + "\n\n—— 本课次面向学生的测试题（前测/中测/后测，题目与参考答案） ——\n"
                         + quiz_block)
        else:
            eval_full = quiz_block
    placeholders = [
        ("（整门课程的混合式教学设计思路）", wc.get("design_idea")),
        ("（整门课程的创新举措）", wc.get("innovation")),
        ("（整门课程的效果达成情况）", wc.get("effect")),
        ("（本课次学生的学情分析）", sl.get("learning_analysis")),
        ("（本课次学生能够掌握和展现的具体知识、技能和能力等，对课程目标的贡献等情况）", build_goals_text(sl.get("learning_goals"))),
        ("（本课次内容与资源的选取、制作、使用情况）", sl.get("content_resources")),
        ("（本课次混合教学的实施过程与方法）", sl.get("process_method")),
        ("（本课次学习评价与反馈方式）", eval_full),
        ("（本课次教学效果与特色）", sl.get("effect")),
    ]
    for ph, val in placeholders:
        if not val or not str(val).strip():
            continue
        replaced = False
        for row in t.rows:
            for cell in row.cells:
                if ph in cell.text:
                    set_cell_text(cell, str(val))
                    replaced = True
                    break
            if replaced:
                break
        if not replaced:
            print("提示：未找到占位符 %s，已跳过。" % ph)

    # ---- 3. 页脚：空 5 行后插入页脚说明 ----
    # 全表统一左对齐，符合阅读习惯
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    for row in t.rows:
        for c in row.cells:
            for para in c.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_word_footer(doc)

    out = os.path.join(out_dir, "混合式教学设计方案.docx")
    doc.save(out)
    print("已生成：%s" % out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
