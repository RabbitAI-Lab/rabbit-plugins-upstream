#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
gen_excel.py —— 由 gen_board.py 的同一数据层生成 Excel（口径永不与 HTML 分叉）

原理：读取 gen_board.py，切出 "DATA LAYER END" 之前的部分 exec 复用。
因此改数据只需改 gen_board.py 上半部一处，HTML 与 Excel 同时更新。

用法：python gen_excel.py   → 生成 board.xlsx
"""
import os
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "gen_board.py")
MARK = "# ============================ DATA LAYER END ============================"

_src = open(SRC, encoding="utf-8").read()
if MARK not in _src:
    raise SystemExit("data-layer marker not found in gen_board.py")
_data_src = _src.split(MARK)[0]
ns = {}
exec(compile(_data_src, "gen_board_data", "exec"), ns)

PROJECTS = ns["PROJECTS"]; SUPPLIERS = ns["SUPPLIERS"]; QUOTES = ns["QUOTES"]
TERMS = ns["TERMS"]; CLAR = ns["CLAR"]; TRANS_NOTE = ns["TRANS_NOTE"]
LEAD_START = ns["LEAD_START"]; LEAD_DEADLINE = ns["LEAD_DEADLINE"]
LEAD_SPAN = ns["LEAD_SPAN"]; LEAD_ABS_END = ns["LEAD_ABS_END"]; LEAD_BASIS = ns["LEAD_BASIS"]
LEAD_BASIS_PID = ns.get("LEAD_BASIS_PID", {})
BOARD_TITLE = ns["BOARD_TITLE"]; TENDER_DATE = ns["TENDER_DATE"]; INVITED_COUNT = ns["INVITED_COUNT"]

RECV_KEYS = [s["key"] for s in SUPPLIERS if s["status"] == "received"]
SUPMAP = {s["key"]: s for s in SUPPLIERS}

def lead_result(k, pid):
    est = LEAD_ABS_END[k][pid] if (k in LEAD_ABS_END and pid in LEAD_ABS_END[k]) else LEAD_START + dt.timedelta(days=LEAD_SPAN[k][pid])
    bzh, ben = LEAD_BASIS_PID.get((k, pid), LEAD_BASIS[k])
    return {"days": LEAD_SPAN[k][pid], "est": est, "ok": est <= LEAD_DEADLINE, "basis_zh": bzh}

TITLE = Font(bold=True, size=14, color="23282E")
HEADB = Font(bold=True, size=11, color="FFFFFF")
BORD = Border(*[Side(style="thin", color="D9DDD8")] * 4)
GREEN = PatternFill("solid", fgColor="E8F3EA")
RED = PatternFill("solid", fgColor="FDEAEA")
AMBER = PatternFill("solid", fgColor="FDF3E0")
GREY = PatternFill("solid", fgColor="F2F3F1")
WRAP = Alignment(wrap_text=True, vertical="top")
EURF = "#,##0.00"
wb = Workbook()

def style_header(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADB; cell.fill = PatternFill("solid", fgColor="3E6E9E"); cell.border = BORD

def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w

# ---------- 01 回标状态（含报价 + 工期，无跨项目合计） ----------
ws = wb.active; ws.title = "01 回标状态"
ws["A1"] = BOARD_TITLE[0]; ws["A1"].font = TITLE
ws["A2"] = "发标 %s / 邀请 %d 家 / 已回标 %d 家" % (TENDER_DATE, INVITED_COUNT, len(RECV_KEYS))
ws.append([])
hdr = ["供应商", "状态", "回标日期"]
for p in PROJECTS:
    hdr += [p["name"] + " 报价", p["name"] + " 所报工期→完成日"]
ws.append(hdr); style_header(ws, ws.max_row, len(hdr))
for s in SUPPLIERS:
    k = s["key"]
    if s["status"] != "received":
        ws.append([s["name"], "待回标", "—"] + ["—"] * (len(hdr) - 3)); continue
    row = [s["name"], "已回标", s["date"]]
    for p in PROJECTS:
        lr = lead_result(k, p["id"])
        row += [QUOTES[k][p["id"]]["total"], "%d 天 → %s（%s）" % (lr["days"], lr["est"].isoformat(), "满足" if lr["ok"] else "超期")]
    ws.append(row)
    rr = ws.max_row
    for c in range(1, len(hdr) + 1):
        ws.cell(row=rr, column=c).border = BORD
    for i, p in enumerate(PROJECTS):
        col = 4 + i * 2
        ws.cell(row=rr, column=col).number_format = EURF
        lo = min(QUOTES[kk][p["id"]]["total"] for kk in RECV_KEYS)
        if QUOTES[k][p["id"]]["total"] == lo:
            ws.cell(row=rr, column=col).fill = GREEN; ws.cell(row=rr, column=col).font = Font(bold=True, color="1F9D55")
        lr = lead_result(k, p["id"])
        ws.cell(row=rr, column=col + 1).fill = GREEN if lr["ok"] else RED
widths(ws, {"A": 16, "B": 10, "C": 12, "D": 16, "E": 26, "F": 16, "G": 26})

# ---------- 02 明细对比 ----------
ws = wb.create_sheet("02 明细对比")
ws.append(["项目", "供应商", "生产", "运输", "安装", "Total", "vs 最低", "运输口径备注"])
style_header(ws, ws.max_row, 8)
for p in PROJECTS:
    pid = p["id"]; lo = min(QUOTES[k][pid]["total"] for k in RECV_KEYS)
    for k in RECV_KEYS:
        d = QUOTES[k][pid]
        ws.append([p["name"], SUPMAP[k]["name"], d["prod"], d["trans"] or "未报价待澄清", d["inst"] or "未报价待澄清",
                   d["total"], "基准（最低）" if d["total"] == lo else "贵 %.2f" % (d["total"] - lo),
                   TRANS_NOTE.get((k, pid), "") or "—"])
        rr = ws.max_row
        for c in range(1, 9):
            ws.cell(row=rr, column=c).border = BORD
        for c in (3, 4, 5, 6):
            if isinstance(ws.cell(row=rr, column=c).value, (int, float)):
                ws.cell(row=rr, column=c).number_format = EURF
        if d["total"] == lo:
            ws.cell(row=rr, column=6).fill = GREEN; ws.cell(row=rr, column=6).font = Font(bold=True, color="1F9D55")
widths(ws, {"A": 26, "B": 14, "C": 12, "D": 14, "E": 14, "F": 12, "G": 12, "H": 30})

# ---------- 03 工期对比 ----------
ws = wb.create_sheet("03 工期对比")
ws.append(["供应商", "项目", "所报工期(天)", "完成日", "是否满足", "依据原文"])
style_header(ws, ws.max_row, 6)
for k in RECV_KEYS:
    for p in PROJECTS:
        lr = lead_result(k, p["id"])
        ws.append([SUPMAP[k]["name"], p["name"], lr["days"], lr["est"], "满足" if lr["ok"] else "需澄清", lr["basis_zh"]])
        rr = ws.max_row
        for c in range(1, 7):
            ws.cell(row=rr, column=c).border = BORD
        ws.cell(row=rr, column=4).number_format = "yyyy-mm-dd"
        ws.cell(row=rr, column=6).alignment = WRAP
        ws.cell(row=rr, column=5).fill = GREEN if lr["ok"] else RED
widths(ws, {"A": 16, "B": 26, "C": 14, "D": 12, "E": 10, "F": 46})

# ---------- 04 条款对比 ----------
ws = wb.create_sheet("04 条款对比")
ws.append(["条款"] + [s["name"] for s in SUPPLIERS]); style_header(ws, ws.max_row, 1 + len(SUPPLIERS))
for zh, en, *vals in TERMS:
    ws.append([zh] + list(vals))
    for c in range(1, 2 + len(SUPPLIERS)):
        ws.cell(row=ws.max_row, column=c).border = BORD; ws.cell(row=ws.max_row, column=c).alignment = WRAP
widths(ws, {"A": 16, "B": 30, "C": 30, "D": 30, "E": 30})

# ---------- 05 待澄清清单 ----------
ws = wb.create_sheet("05 待澄清清单")
ws.append(["级别", "事项（中）", "Item (EN)", "状态"]); style_header(ws, ws.max_row, 4)
for pr, zh, en in sorted(CLAR, key=lambda c: {"P0": 0, "P1": 1, "P2": 2}[c[0]]):
    ws.append([pr, zh, en, "待回复"])
    rr = ws.max_row
    for c in range(1, 5):
        ws.cell(row=rr, column=c).border = BORD; ws.cell(row=rr, column=c).alignment = WRAP
    ws.cell(row=rr, column=1).font = Font(bold=True, color="D6363A" if pr == "P0" else "B7791F")
    ws.cell(row=rr, column=1).fill = RED if pr == "P0" else AMBER
widths(ws, {"A": 8, "B": 60, "C": 60, "D": 12})

# ---------- 06 定标建议 ----------
ws = wb.create_sheet("06 定标建议")
ws["A1"] = "定标建议（满足工期前提下，价格最低优先）"; ws["A1"].font = TITLE
ws["A2"] = "规则：① 起算日 %s；② %s 前完成安装 → 满足；③ 满足工期者中报价最低优先。" % (LEAD_START.isoformat(), LEAD_DEADLINE.isoformat())
ws["A2"].font = Font(size=10, color="B7791F"); ws["A2"].alignment = WRAP
ws.append([])
ws.append(["项目", "推荐供应商", "报价", "完成日", "备选"]); style_header(ws, ws.max_row, 5)
pick_row = {}
for p in PROJECTS:
    pid = p["id"]
    cand = sorted([k for k in RECV_KEYS if lead_result(k, pid)["ok"]], key=lambda k: QUOTES[k][pid]["total"]) or sorted(RECV_KEYS, key=lambda k: QUOTES[k][pid]["total"])
    pick_row[pid] = cand[0]
    alt = [k for k in cand if k != cand[0]]
    ws.append([p["name"], SUPMAP[cand[0]]["name"], QUOTES[cand[0]][pid]["total"], lead_result(cand[0], pid)["est"],
               ("%s %.2f" % (SUPMAP[alt[0]]["name"], QUOTES[alt[0]][pid]["total"])) if alt else "—"])
    rr = ws.max_row
    for c in range(1, 6):
        ws.cell(row=rr, column=c).border = BORD
    ws.cell(row=rr, column=3).number_format = EURF
    ws.cell(row=rr, column=3).fill = GREEN; ws.cell(row=rr, column=3).font = Font(bold=True, color="1F9D55")
    ws.cell(row=rr, column=4).number_format = "yyyy-mm-dd"
ws.append([])
ws.append(["说明：分项目独立定标，不做跨项目合计。"])
ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color="B7791F")
widths(ws, {"A": 28, "B": 18, "C": 14, "D": 12, "E": 26})

_out = "board.xlsx"
wb.save(_out)
print("WROTE %s  bytes=%d" % (_out, os.path.getsize(_out)))
